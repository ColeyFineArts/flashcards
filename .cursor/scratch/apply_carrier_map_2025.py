#!/usr/bin/env python3
"""Lock carrier → property from user 2026-09-20.

User: “Lemomade is for 216 N Grove I believe. The Geico and Travelers is for 524 Ferdinand.”

Does NOT move dollars (216 insurance and 524 INSURANCE already match this map).
Does NOT re-run apply_checking_answers_2025.py.
Does NOT put Geico on the 524 home INSURANCE line (cars; State Farm HOME share stays).
Does NOT put Lemonade on 827 N Grove (closed 12/18 — different house).

Not tax advice.
"""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path("/workspace/.cursor/scratch")
sys.path.insert(0, str(ROOT))
import apply_cc_answers_2025 as ans  # noqa: E402

XLSX = ans.XLSX
CSV_DIR = ans.CSV_DIR
PACKET = ROOT / "cpa_packet_v1.json"
GREEN = PatternFill("solid", fgColor="C6EFCE")
CG = "Century Gothic"

USER = (
    "Lemomade is for 216 N Grove I believe. The Geico and Travelers is for 524 Ferdinand"
)
LOCK = (
    "LOCKED 2026-09-20: Lemonade = 216 N. Oak Park Avenue #1Z (LTR). "
    "Packet has no “216 N Grove” — that is this Oak Park unit, not 827 N Grove "
    "(Grove closed 12/18, different house). Travelers + Geico = 524 Ferdinand "
    "(Travelers = home, Geico = cars). Neither carrier is 216. Lemonade is not 524. "
    "Grove Collaborative ≠ 827 N Grove the house."
)


def patch_ask() -> None:
    path = CSV_DIR / "ASK.csv"
    rows = list(csv.reader(path.open(encoding="utf-8")))
    carrier_row = [
        "Lemonade = 216 Oak Park / Travelers+Geico = 524 Ferdinand",
        f'LOCKED 2026-09-20. User: “{USER}.” {LOCK} '
        "216 RENTER'S INSURANCE stays Lemonade $42.84/mo. 524 INSURANCE stays State Farm HOME share "
        "+ Travelers leftover $0. Geico $540.89 *AUTO credit remains EXCLUDED personal — not 216, not EPGC, "
        "not the 524 home INSURANCE line.",
    ]
    out = []
    inserted = False
    for row in rows:
        if row and row[0].startswith("Lemonade 216 LTR insurance"):
            out.append(carrier_row)
            out.append(row)
            inserted = True
            continue
        out.append(row)
    if not inserted:
        out.append(carrier_row)
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(out)


def write_lock_csv() -> None:
    path = CSV_DIR / "LOCK_carriers_2025.csv"
    rows = [
        ["Carrier", "Property", "Status", "Where the dollars are", "Not"],
        [
            "Lemonade",
            "216 N. Oak Park Avenue #1Z (LTR)",
            "LOCKED 2026-09-20",
            "216 RENTER'S INSURANCE $42.84/mo × 12 = $514.08 (cash $514.00 10/3 Megan Chase)",
            "Not 524. Not 827 N Grove. Not Grove Collaborative.",
        ],
        [
            "Travelers",
            "524 Ferdinand (home)",
            "LOCKED 2026-09-20",
            "524 INSURANCE leftover Jan 2025 $0.00 on cards/Monarch. Units 1 & 2 keep 50/50 home-share.",
            "Not 216. Not Geico (cars). Not Lemonade.",
        ],
        [
            "Geico",
            "524 Ferdinand household cars",
            "LOCKED 2026-09-20",
            "Personal auto. 2025-02-11 *AUTO credit $540.89 EXCLUDED. Not on 524 home INSURANCE line.",
            "Not 216. Not EPGC. Not Lemonade.",
        ],
        [
            "State Farm 2025 bundle",
            "524 Ferdinand car+home (replaced Travelers/Geico)",
            "LOCKED 2026-09-18 / 2026-09-20",
            "HOME share $1,105.46 on 524 INSURANCE 50/50 U1/U2. AUTO $1,019.33 personal.",
            "Not 216.",
        ],
        [
            "827 N Grove the house",
            "Closed 2025-12-18",
            "NOT Lemonade",
            "CapEx $53,660 locked. HD 1/3 $604.12 pre-close materials.",
            "User said “216 N Grove” — that is Oak Park #1Z, not this purchase.",
        ],
    ]
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def write_start_here() -> None:
    (CSV_DIR / "START_HERE_carriers.txt").write_text(
        "START HERE — insurance carriers (2026-09-20)\n"
        "\n"
        "Not tax advice. Organizational packet only.\n"
        "\n"
        f'User: “{USER}.”\n'
        "\n"
        "LOCKED (no dollar move — already booked this way)\n"
        "1. Lemonade = 216 N. Oak Park Avenue #1Z (LTR).\n"
        "   There is no 216 N Grove on this packet. 827 N Grove closed 12/18 is a different house.\n"
        "   2023/2024 Excel already amortized Lemonade on the 216 tab ($26.75 / $35.08).\n"
        "   2025 is $42.84/mo × 12 = $514.08 (cash $514.00 10/3 Megan Chase).\n"
        "2. Travelers = 524 Ferdinand home. Leftover Jan cash $0. 50/50 U1/U2 home-share.\n"
        "3. Geico = 524 Ferdinand cars. Credit $540.89 EXCLUDED personal.\n"
        "   Not on 216. Not on EPGC. Not on the 524 home INSURANCE line\n"
        "   (that line is State Farm HOME share + Travelers leftover $0).\n"
        "4. Grove Collaborative ≠ 827 N Grove the house.\n"
        "\n"
        "Do not re-run apply_checking_answers_2025.py.\n",
        encoding="utf-8",
    )


def patch_start_here_insurance() -> None:
    path = CSV_DIR / "START_HERE_insurance.txt"
    text = path.read_text(encoding="utf-8")
    extra = (
        "\n"
        f'User also: “{USER}.”\n'
        f"{LOCK}\n"
    )
    if "216 N Grove" not in text:
        text = text.replace(
            'User: “Traveler\'s was home insurance before for Ferdinand and Geico was for the cars.”\n',
            'User: “Traveler\'s was home insurance before for Ferdinand and Geico was for the cars.”\n'
            + extra,
        )
        path.write_text(text, encoding="utf-8")


def patch_start_here_oak() -> None:
    path = CSV_DIR / "START_HERE_oak_park_costs.txt"
    text = path.read_text(encoding="utf-8")
    add = (
        "   User: Lemonade is the 216 policy (said “216 N Grove” — this Oak Park unit, not 827 Grove). "
        "Travelers and Geico are 524 Ferdinand, not 216.\n"
    )
    needle = "   Travelers/Geico/State Farm are not 216.\n"
    if needle in text and "said “216 N Grove”" not in text:
        text = text.replace(needle, needle + add)
        path.write_text(text, encoding="utf-8")


def patch_ltr_lock() -> None:
    path = CSV_DIR / "LOCK_LTR_Lemonade.csv"
    if not path.exists():
        return
    rows = list(csv.reader(path.open(encoding="utf-8")))
    if rows:
        rows[1] = [
            "Lemonade = 216 N. Oak Park Ave #1Z (LTR)",
            "LOCKED 2026-09-20 (user: 216, not 524)",
            "216 RENTER'S INSURANCE $42.84/mo amortized (last-year pattern)",
            "$514.08 sheet / cash $514.00 10/3 Megan Chase. Not 827 Grove. Not 524.",
        ]
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def patch_travelers_lock() -> None:
    path = CSV_DIR / "LOCK_Travelers_Geico.csv"
    rows = list(csv.reader(path.open(encoding="utf-8")))
    extra = [
        "Lemonade = 216 Oak Park (not 524)",
        "LOCKED 2026-09-20",
        "216 RENTER'S INSURANCE only. User: Lemonade is the 216 policy; Travelers+Geico are Ferdinand.",
        "$42.84/mo × 12. Not 827 N Grove.",
    ]
    if not any(r and r[0].startswith("Lemonade") for r in rows):
        rows.append(extra)
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def patch_xlsx_note() -> None:
    wb = load_workbook(XLSX)
    oak = wb["216 N. Oak Park Ave"]
    oak["B64"] = (
        "Insurance — Lemonade = 216 N. Oak Park #1Z (user: 216, not 524). "
        "$42.84/mo × 12 = $514.08 (cash $514.00 10/3 Megan Chase). "
        "Travelers + Geico = 524 Ferdinand (home / cars), not 216. "
        "Not 827 N Grove. Not tax advice."
    )
    oak["B64"].alignment = Alignment(wrap_text=True)
    oak["B64"].font = Font(name=CG, size=10)
    oak["B64"].fill = GREEN
    u1 = wb["524 Ferdinand Ave, Unit 1"]
    u2 = wb["524 Ferdinand Ave, Unit 2"]
    add = (
        " User 2026-09-20: Travelers + Geico are 524 Ferdinand; Lemonade is 216 Oak Park (not Grove, not 524)."
    )
    for ws, cell in ((u1, "A45"), (u2, "A39")):
        v = ws[cell].value or ""
        if "Lemonade is 216" not in str(v):
            ws[cell] = str(v).rstrip() + add
            ws[cell].alignment = Alignment(wrap_text=True)
    wb.save(XLSX)


def patch_turbo_csv_note() -> None:
    path = CSV_DIR / "216_Oak_Park.csv"
    rows = list(csv.reader(path.open(encoding="utf-8")))
    prefix = (
        "CARRIER LOCK 2026-09-20: Lemonade = 216 N. Oak Park #1Z (user said “216 N Grove” — this unit, "
        "not 827 Grove). Travelers + Geico = 524 Ferdinand, not 216. "
    )
    for row in rows:
        if len(row) > 1 and isinstance(row[1], str) and "RENTER'S INSURANCE $42.84" in row[1]:
            if "CARRIER LOCK" not in row[1]:
                row[1] = prefix + row[1]
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def write_packet() -> None:
    pkt = json.loads(PACKET.read_text(encoding="utf-8")) if PACKET.exists() else {}
    pkt["updated"] = "2026-09-20"
    pkt["carrier_map"] = {
        "user": USER,
        "lemonade": "216 N. Oak Park Avenue #1Z",
        "not_lemonade": ["524 Ferdinand", "827 N Grove", "Grove Collaborative"],
        "travelers": "524 Ferdinand home",
        "geico": "524 Ferdinand cars (personal auto, not 216)",
        "note": "User said 216 N Grove; packet maps that to 216 Oak Park LTR, not 827 Grove.",
    }
    pkt["note"] = (
        "Lemonade = 216 Oak Park. Travelers = 524 home. Geico = 524 cars. "
        "Grove Collaborative ≠ 827 N Grove the house."
    )
    PACKET.write_text(json.dumps(pkt, indent=2), encoding="utf-8")


def main() -> None:
    patch_ask()
    write_lock_csv()
    write_start_here()
    patch_start_here_insurance()
    patch_start_here_oak()
    patch_ltr_lock()
    patch_travelers_lock()
    patch_xlsx_note()
    patch_turbo_csv_note()
    write_packet()
    print("carrier map locked: Lemonade=216 Oak Park; Travelers+Geico=524")


if __name__ == "__main__":
    main()

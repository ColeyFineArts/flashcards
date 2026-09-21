#!/usr/bin/env python3
"""Lock: Personal Income 2025 must be last year's workbook with 2025 numbers.

Fails if the file was rebuilt from CSVs (Calibri, no merges, missing W2/T4).
Run this before any Drive upload. Not tax advice.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from sanitize_xlsx_for_sheets import assert_formulas_are_real

ROOT = Path("/workspace/.cursor/scratch")
DEFAULT = ROOT / "tax_turbo_parts" / "Personal Income 2025.xlsx"
PRIOR = ROOT / "Personal_Income_prior.xlsx"

LAST_YEAR_TABS = [
    "Income",
    "524 Ferdinand Ave, Unit 2",
    "216 N. Oak Park Ave",
    "EPGC LLC",
    "Refrence Library",
    "Art Sales and Purchases",
    "GCM",
    "Hindman W2",
    "Megan T4",
    "Investments",
    "524 Ferdinand Ave, Unit 1",
]

# 216 2025 = columns AC–AN (29–40)
OAK_RENT = [1950, 1950, 1950, 1950, 1950, 1950, 1450, 2450, 0, 0, 0, 3900]
OAK_INS = [42.84] * 12
OAK_MORTGAGE = [1226.25] * 10 + [1177.52, 1177.52]
OAK_HOA = [421.53] * 12


def fail(msg: str) -> None:
    raise SystemExit(f"LOCK FAIL: {msg}")


def months(ws, row: int, start: int) -> list:
    return [ws.cell(row, c).value for c in range(start, start + 12)]


def assert_workbook(path: Path) -> None:
    if not path.exists():
        fail(f"missing {path}")
    wb = load_workbook(path, data_only=False)
    if wb.sheetnames[:11] != LAST_YEAR_TABS:
        fail(f"first 11 tabs {wb.sheetnames[:11]!r} != last year {LAST_YEAR_TABS!r}")
    extra_ok = {"2025 Data sources"}
    extras = [n for n in wb.sheetnames[11:] if n not in extra_ok]
    if extras:
        fail(f"extra tabs that are not last year: {extras}")
    if "2025 Data sources" not in wb.sheetnames:
        fail("missing 2025 Data sources (provenance lock)")

    prior = load_workbook(PRIOR, read_only=True)
    try:
        if list(prior.sheetnames) != LAST_YEAR_TABS:
            fail(f"prior model tabs changed: {list(prior.sheetnames)}")
    finally:
        prior.close()

    oak = wb["216 N. Oak Park Ave"]
    if oak["B14"].value != "RENTER'S INSURANCE":
        fail(f"216 B14 is {oak['B14'].value!r}")
    if oak["B14"].font.name != "Century Gothic":
        fail(f"216 B14 font {oak['B14'].font.name!r} (CSV rebuilds use Calibri)")
    if oak["B8"].value != "EXPENSES":
        fail(f"216 B8 is {oak['B8'].value!r}")
    merges = {str(r) for r in oak.merged_cells.ranges}
    if "C1:O1" not in merges:
        fail(f"216 lost last-year merge C1:O1; have {merges}")
    if "P1:AB1" not in merges:
        fail(f"216 missing 2024 year-block merge P1:AB1; have {merges}")
    if "AC1:AO1" not in merges:
        fail(f"216 missing 2025 year-block merge AC1:AO1; have {merges}")
    if oak["AC1"].value != oak["C1"].value:
        fail(f"216 AC1 {oak['AC1'].value!r} != C1 {oak['C1'].value!r}")
    if months(oak, 5, 29) != OAK_RENT:
        fail(f"216 2025 rent {months(oak, 5, 29)}")
    if months(oak, 14, 29) != OAK_INS:
        fail(f"216 2025 insurance {months(oak, 14, 29)} (need $42.84/mo)")
    if months(oak, 35, 29) != OAK_MORTGAGE:
        fail(f"216 2025 mortgage {months(oak, 35, 29)}")
    if months(oak, 36, 29) != OAK_HOA:
        fail(f"216 2025 HOA {months(oak, 36, 29)}")
    if oak.cell(15, 36).value != 150:
        fail("216 Aug INTERIOR MAINTENANCE != 150 Brennan")
    if oak.cell(15, 40).value != 11:
        fail("216 Dec INTERIOR MAINTENANCE != 11 Ace keys")
    if oak.cell(16, 40).value != 675:
        fail("216 Dec INTERIOR REPAIR != 675 Joan")
    if oak.cell(38, 38).value != 300:
        fail("216 Oct MOVE OUT FEE != 300 Imelda")

    inc = wb["Income"]
    if [inc["H1"].value, inc["I1"].value] != [2025, 2025]:
        fail(f"Income 2025 headers {inc['H1'].value!r} {inc['I1'].value!r}")
    if inc["I4"].value != 70618:
        fail(f"Income I4 {inc['I4'].value}")
    if inc["I5"].value != 19500:
        fail(f"Income I5 {inc['I5'].value}")
    if inc["I6"].value != 17086.94:
        fail(f"Income I6 {inc['I6'].value}")
    if inc["I7"].value != 21500:
        fail(f"Income I7 {inc['I7'].value}")
    if inc["I8"].value != 39055.06:
        fail(f"Income I8 {inc['I8'].value}")
    if inc["A4"].value != "Hindman ":
        fail(f"Income A4 {inc['A4'].value!r}")
    if inc["I10"].data_type != "f" or inc["I10"].value != "=SUM(I4:I9)":
        fail(f"Income I10 must be formula =SUM(I4:I9), got type={inc['I10'].data_type!r} value={inc['I10'].value!r}")
    if oak["AO5"].data_type != "f":
        fail(f"216 AO5 must be a formula, got type={oak['AO5'].data_type!r} value={oak['AO5'].value!r}")

    gcm = wb["GCM"]
    if [gcm["B24"].value, gcm["C24"].value, gcm["D24"].value, gcm["E24"].value] != [
        4100,
        9200,
        4100,
        4100,
    ]:
        fail(f"GCM 2025 {gcm['B24'].value},{gcm['C24'].value},{gcm['D24'].value},{gcm['E24'].value}")

    ep = wb["EPGC LLC"]
    if ep["A51"].value != 2025:
        fail(f"EPGC A51 {ep['A51'].value!r}")
    if ep["B53"].value != 14000 or ep["C53"].value != 136000:
        fail(f"EPGC Art Sales 2025 {ep['B53'].value}/{ep['C53'].value}")
    if ep["H53"].value != 76000:
        fail(f"EPGC Art Sales July {ep['H53'].value} (need Belt $50,000 + Venus $26,000)")
    if ep["K53"].value != 25000:
        fail(f"EPGC Art Sales October {ep['K53'].value} (need August Fortuna joint $25,000)")
    if ep["G54"].value != 13595:
        fail(f"EPGC Consultant June {ep['G54'].value}")
    if ep["H69"].value != 334.17 or ep["K69"].value != 658.63:
        fail(f"EPGC Consultant Fees {ep['H69'].value}/{ep['K69'].value}")

    art = wb["Art Sales and Purchases"]
    if art["F73"].value != 14000:
        fail("Art Sales seals F73 != 14000")
    if art["F74"].value != 105000:
        fail("Art Sales mosaics F74 != 105000")
    if art["F77"].value != 1000:
        fail("Art Sales Aquinas F77 != 1000")
    if art["A1"].value != "Art Sales and Purchases":
        fail(f"Art Sales A1 {art['A1'].value!r} (row 1 must exist or Google convert drops the tab)")
    if art["A78"].value != "Roman Gold Belt":
        fail(f"Art Sales A78 {art['A78'].value!r} (Belt SALE)")
    if art["C78"].value != 45000 or art["F78"].value != 50000:
        fail(f"Art Sales Belt Cost/Sale {art['C78'].value}/{art['F78'].value}")
    if "Venus" not in str(art["A79"].value or ""):
        fail(f"Art Sales A79 {art['A79'].value!r} (Venus SALE)")
    if art["C79"].value != 20000 or art["F79"].value != 26000:
        fail(f"Art Sales Venus Cost/proceeds {art['C79'].value}/{art['F79'].value}")
    if "August Fortuna" not in str(art["A80"].value or ""):
        fail(f"Art Sales A80 {art['A80'].value!r} (August Fortuna joint SALE)")
    if art["C80"].value != 20000 or art["F80"].value != 25000:
        fail(f"Art Sales August Cost/proceeds {art['C80'].value}/{art['F80'].value}")
    if str(art["A81"].value or "") != "Total":
        fail(f"Art Sales A81 {art['A81'].value!r} (Total after August)")
    parked = " ".join(str(art.cell(r, 1).value or "") for r in range(120, (art.max_row or 120) + 1))
    if "Roman Gold Belt — Fortuna inventory" in parked:
        fail("Belt still parked as unsold Fortuna inventory")
    if "likely Venus — Fortuna inventory" in parked:
        fail("Venus still parked as unsold Fortuna inventory")
    if "object still unnamed" in parked:
        fail("August Fortuna still parked as unsold unnamed inventory")

    inv = wb["Investments"]
    if "Coinbase" not in str(inv["A5"].value or ""):
        fail("Investments missing Coinbase row")

    ds = wb["2025 Data sources"]
    blob = " ".join(str(ds.cell(r, 1).value or "") for r in range(1, 8))
    if "HOW THIS FILE IS MADE" not in blob:
        fail("Data sources missing HOW THIS FILE IS MADE")

    print(f"LOCK OK {path} ({path.stat().st_size} bytes)")
    print("tabs", wb.sheetnames)
    assert_formulas_are_real(path)


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT
    assert_workbook(target)

#!/usr/bin/env python3
"""Third pass: recast State Farm as bundled home+auto (user 2026-09-18).

REPLACES 524 U1/U2 INSURANCE (does not add). Do not re-run apply_cc_2025.py or
apply_cc_answers_2025.py on the filled workbook.

User: bundled car and home at State Farm. Previously home=Travelers, car=Geico.
Only the home share hits 524 INSURANCE (50/50 U1/U2, yellow proxy). Auto share
and the Geico *AUTO credit are personal. Not tax advice.
"""
from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook

ROOT = Path("/workspace/.cursor/scratch")
sys.path.insert(0, str(ROOT))
import apply_cc_2025 as base  # noqa: E402
import apply_cc_answers_2025 as ans  # noqa: E402

XLSX = ans.XLSX
DELIVERABLE = ans.DELIVERABLE
PARTS = ans.PARTS
CSV_DIR = ans.CSV_DIR
JSON_PATH = ans.JSON_PATH

NONE_FILL = PatternFill(fill_type=None)
U2_NOTE = (
    "CC answers 2026-09-18: Nicor small / Prime ATT / Grove / Schauer / yard CONFIRMED. "
    f"March HD ${ans.MARCH_HD_TOTAL} → 524 EXTERIOR REPAIRS 50/50 with Unit 1. "
    "Remaining Prime HD/Lowe’s/IKEA 1/3 524 (50/50), 1/3 personal, 1/3 Grove. "
    "ELECTRIC yellow — REMIND JACOB: double-check ComEd with Megan (checking ACH Jacob $829.15; any extra check?). "
    "Home Depot March read as 524 (not 534). Extra Space $360 HOLD. "
    "State Farm: bundled home+auto (was Travelers home / Geico auto). Only HOME share on this INSURANCE line "
    "(50/50 with Unit 1, yellow = Feb $157.53 home / $145.25 auto proxy). Auto share and Geico $540.89 credit "
    "EXCLUDED personal. No Sep–Dec on Prime. Not tax advice."
)
U1_NOTE = (
    "§121 residence through sale 2025-12-18. Nicor large / water 50/50 CONFIRMED. "
    f"March HD + remaining HD/Lowe’s/IKEA 524-share → EXTERIOR REPAIRS (U1 ${sum(ans.u1_hd)}). "
    "ELECTRIC yellow — REMIND JACOB: double-check ComEd with Megan (checking ACH INDN Megan Gerrard $1,565.26; paid ACH and/or check?). "
    "Internet still $0 on this tab (Prime ATT on Unit 2). "
    "State Farm: bundled home+auto (was Travelers home / Geico auto). Only HOME share on this INSURANCE line "
    "(50/50 with Unit 2, yellow proxy). Auto + Geico credit personal. Not tax advice."
)


def replace_month_row(ws, row, values, colmap, fill):
    """Overwrite every month (including zeros). Used so insurance is not double-counted."""
    for i, m in enumerate(base.MONTHS):
        cell = ws.cell(row, colmap[m])
        cell.value = base.money(values[i])
        cell.number_format = "0.00"
        cell.fill = fill if values[i] != 0 else NONE_FILL


def write_sf_split(wb):
    if "SF_SPLIT" in wb.sheetnames:
        del wb["SF_SPLIT"]
    ws = wb.create_sheet("SF_SPLIT", 3)
    ws["A1"] = "State Farm 2025 — home vs auto split (user 2026-09-18) — not tax advice"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:H1")
    ws["A2"] = (
        "Bundled car + home at State Farm. Previously home=Travelers, car=Geico. "
        "Feb 11 Prime charged two unbundled policies: $157.53 (HOME — replaces Travelers) and $145.25 "
        "(AUTO — leftover Prime SF in 2026 after 524 sold is ~$147). Mar–Jul combined $302.75; Aug $308.26. "
        "Home = round(bill × 157.53/302.78); auto = remainder. Only HOME → 524 INSURANCE 50/50 U1/U2 (yellow). "
        "Geico *AUTO credit $540.89 (2025-02-11) is unused auto premium — personal, not 524. "
        "No 2025 Travelers on Prime/Sapphire/BoA/Monarch. No Sep–Dec SF on Prime."
    )
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:H2")
    ws.row_dimensions[2].height = 72
    headers = ["Month", "Prime billed $", "HOME → 524", "AUTO personal", "U2 50%", "U1 50%", "On 524 P&L?", "Note"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(4, i, h)
        cell.fill = base.HEADER_FILL
        cell.font = base.HEADER_FONT
        cell.border = base.THIN
    for i, rec in enumerate(base.sf_split_rows, 5):
        vals = [
            rec["month"],
            base.money(rec["billed"]),
            base.money(rec["home"]),
            base.money(rec["auto"]),
            base.money(rec["u2"]),
            base.money(rec["u1"]),
            "HOME only (yellow)",
            rec["note"],
        ]
        for c, val in enumerate(vals, 1):
            cell = ws.cell(i, c, val)
            cell.border = base.THIN
            if c in (2, 3, 4, 5, 6):
                cell.number_format = "0.00"
            if c == 3:
                cell.fill = base.YELLOW
            elif c == 4:
                cell.fill = base.GRAY
            elif c in (5, 6):
                cell.fill = base.YELLOW
    tot_r = 5 + len(base.sf_split_rows)
    ws.cell(tot_r, 1, "TOTAL").font = Font(bold=True)
    for col, val, fill in (
        (2, sum((r["billed"] for r in base.sf_split_rows), 0), base.HEADER_FILL),
        (3, sum(base.u2_ins) + sum(base.u1_ins), base.YELLOW),
        (4, sum(base.sf_auto), base.GRAY),
        (5, sum(base.u2_ins), base.YELLOW),
        (6, sum(base.u1_ins), base.YELLOW),
    ):
        cell = ws.cell(tot_r, col, base.money(val))
        cell.number_format = '"$"#,##0.00'
        cell.font = base.HEADER_FONT if col == 2 else Font(bold=True)
        cell.fill = fill
        cell.border = base.THIN
    ws.cell(tot_r + 2, 1, "Geico *AUTO credit 2025-02-11")
    g = ws.cell(tot_r + 2, 4, -base.money(base.GEICO_AUTO_CREDIT))
    g.number_format = '"$"#,##0.00'
    g.fill = base.GRAY
    ws.cell(tot_r + 2, 7, "EXCLUDED personal").fill = base.GRAY
    ws.cell(tot_r + 2, 8, "Unused Geico auto premium when switching to State Farm auto. Not 524 income.")
    ws.cell(tot_r + 4, 1, (
        "CPA: 2023–24 Personal Income booked all homeowners on Unit 1 ($0 on Unit 2). "
        "This packet 50/50’s the 2025 home share across both units because it is a building policy. "
        "Swap to 100% Unit 1 if that is the method you want to keep. Declarations pages beat this proxy."
    ))
    ws.cell(tot_r + 4, 1).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=tot_r + 4, start_column=1, end_row=tot_r + 5, end_column=8)
    ws.row_dimensions[tot_r + 4].height = 36
    for i, w in enumerate([10, 16, 16, 16, 12, 12, 22, 52], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    return ws


def main():
    assert XLSX.exists(), XLSX
    wb = load_workbook(XLSX)

    u2 = wb["524 Ferdinand Ave, Unit 2"]
    u2_map = ans.label_map(u2, 2)
    replace_month_row(u2, u2_map["INSURANCE"], base.u2_ins, base.PROP_COLS, base.YELLOW)
    u2["B30"] = U2_NOTE
    u2["B30"].alignment = Alignment(wrap_text=True)
    u2["B30"].fill = base.YELLOW
    u2.row_dimensions[30].height = 96

    u1 = wb["524 Ferdinand Ave, Unit 1"]
    u1_map = ans.label_map(u1, 2)
    replace_month_row(u1, u1_map["INSURANCE"], base.u1_ins, base.PROP_COLS, base.YELLOW)
    u1["B30"] = U1_NOTE
    u1["B30"].alignment = Alignment(wrap_text=True)
    u1["B30"].fill = base.YELLOW
    u1.row_dimensions[30].height = 96

    src = wb["_SOURCE_2025"]
    src["A20"] = "State Farm 2026-09-18"
    src["B20"] = (
        "Bundled home+auto (was Travelers home / Geico auto). "
        f"HOME share ${sum(base.u2_ins)+sum(base.u1_ins)} on 524 INSURANCE 50/50 U1/U2 (yellow proxy from Feb $157.53/$145.25). "
        f"AUTO share ${sum(base.sf_auto)} + Geico credit ${base.GEICO_AUTO_CREDIT} EXCLUDED personal. "
        "No 2025 Travelers found. No Sep–Dec SF on Prime. Not tax advice."
    )
    src["A20"].fill = base.YELLOW
    src["B20"].fill = base.YELLOW

    write_sf_split(wb)
    ans.rebuild_ledger(wb)
    ans.rebuild_ask(wb)

    wb.save(XLSX)
    shutil.copy2(XLSX, DELIVERABLE)

    export_names = [
        ("EPGC LLC", "EPGC_LLC.csv"),
        ("524 Ferdinand Ave, Unit 2", "524_Unit_2.csv"),
        ("524 Ferdinand Ave, Unit 1", "524_Unit_1.csv"),
        ("CC_LEDGER", "CC_LEDGER.csv"),
        ("ASK", "ASK.csv"),
        ("HOME_SPLIT", "HOME_SPLIT.csv"),
        ("SF_SPLIT", "SF_SPLIT.csv"),
        ("Art Sales and Purchases", "Art_Sales.csv"),
        ("827 Grove CapEx", "Grove_CapEx.csv"),
        ("_SOURCE_2025", "SOURCE_2025.csv"),
    ]
    for name, fname in export_names:
        ans.sheet_to_csv(wb[name], CSV_DIR / fname)
        safe = name.replace(",", "").replace(".", "")[:40]
        ans.save_sheet_part(wb[name], PARTS / f"{safe}.xlsx")

    payload = json.loads(JSON_PATH.read_text(encoding="utf-8")) if JSON_PATH.exists() else {}
    payload.update(
        {
            "updated": "2026-09-18",
            "state_farm": {
                "user": "bundled car and home; previously home=Travelers, car=Geico",
                "method": "Feb unbundled $157.53 home / $145.25 auto as ratio of later combined bills",
                "prime_billed_feb_aug": float(sum((r["billed"] for r in base.sf_split_rows), 0)),
                "home_524_total": float(sum(base.u2_ins) + sum(base.u1_ins)),
                "home_u2": float(sum(base.u2_ins)),
                "home_u1": float(sum(base.u1_ins)),
                "auto_excluded": float(sum(base.sf_auto)),
                "geico_credit_excluded": float(base.GEICO_AUTO_CREDIT),
                "travelers_2025_found": 0.0,
                "sep_dec_on_prime": 0.0,
                "unit_split": "50/50 home share (2023-24 Excel had all homeowners on U1)",
            },
            "confirmed": [a[0] for a in ans.CONFIRMED],
            "still_open": [a[0] for a in ans.STILL_OPEN],
        }
    )
    JSON_PATH.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    print("home U2", sum(base.u2_ins), "U1", sum(base.u1_ins), "total", sum(base.u2_ins) + sum(base.u1_ins))
    print("auto excluded", sum(base.sf_auto), "geico credit", base.GEICO_AUTO_CREDIT)
    print("billed", sum((r["billed"] for r in base.sf_split_rows), 0))
    for rec in base.sf_split_rows:
        print(rec["month"], rec["billed"], "home", rec["home"], "auto", rec["auto"], "u2", rec["u2"], "u1", rec["u1"])
    print("saved", XLSX, "bytes", XLSX.stat().st_size)


if __name__ == "__main__":
    main()

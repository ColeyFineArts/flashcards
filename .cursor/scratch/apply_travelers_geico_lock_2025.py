#!/usr/bin/env python3
"""Lock Travelers = 524 home / Geico = personal auto (user 2026-09-20).

Does NOT recast 524 INSURANCE dollars (already home-share only; Geico not on
524; Travelers leftover Jan $0 on cards/Monarch). Does NOT re-run
apply_cc_answers_2025.py or apply_checking_answers_2025.py (would double HD /
wipe Consultant). Not tax advice.
"""
from __future__ import annotations

import csv
import json
import shutil
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

ROOT = Path("/workspace/.cursor/scratch")
sys.path.insert(0, str(ROOT))
import apply_cc_2025 as base  # noqa: E402
import apply_cc_answers_2025 as ans  # noqa: E402
import apply_checking_answers_2025 as chk  # noqa: E402
import apply_state_farm_2025 as sf  # noqa: E402

XLSX = ans.XLSX
DELIVERABLE = ans.DELIVERABLE
CSV_DIR = ans.CSV_DIR
PACKET = ROOT / "cpa_packet_v1.json"
START_CC = CSV_DIR / "START_HERE_checking.txt"
LOCK_CSV = CSV_DIR / "LOCK_Travelers_Geico.csv"
START_INS = CSV_DIR / "START_HERE_insurance.txt"

U2_NOTE_LIKE = (
    "2025: platform STR net on RENTAL INCOME (Jun–Nov). TruGreen/Alsip on YARD SERVICE (same as 2024). "
    "Grove Collaborative on SUPPLIES (≠ 827 N Grove the house). "
    "INSURANCE LOCK 2026-09-20: Travelers was prior 524 home; Geico was the cars. "
    "Only State Farm HOME share on this INSURANCE line (50/50 U1/U2). Auto + Geico $540.89 credit personal, not EPGC. "
    "Travelers leftover Jan 2025 is 524 home ($0 on cards/Monarch). ELECTRIC yellow — confirm ComEd with Megan. Not tax advice."
)
U1_NOTE_LIKE = (
    "2025 columns P–AB. Residence through sale 2025-12-18. MORTGAGE = US Bank loan 09422 monthly P+I from joint 0203 "
    "(not $50k extra principal on 9922 Jul/Sep — those are not on this row). HOA n/a. "
    "INSURANCE LOCK 2026-09-20: Travelers was prior 524 home; Geico was the cars. "
    "Only State Farm HOME share on this INSURANCE line (50/50 with Unit 2). Travelers leftover Jan is 524 home "
    "($0 on cards/Monarch). Auto + Geico credit personal — not EPGC. Not tax advice."
)

TURBO_U2_NOTE = sf.U2_NOTE
TURBO_U1_NOTE = sf.U1_NOTE

STATE_FARM = {
    "user": "Traveler's was home insurance before for Ferdinand and Geico was for the cars",
    "locked": "2026-09-20",
    "method": "Feb unbundled $157.53 home / $145.25 auto as ratio of later combined bills",
    "prime_billed_feb_aug": float(sum((r["billed"] for r in base.sf_split_rows), 0)),
    "home_524_total": float(sum(base.u2_ins) + sum(base.u1_ins)),
    "home_u2": float(sum(base.u2_ins)),
    "home_u1": float(sum(base.u1_ins)),
    "auto_excluded": float(sum(base.sf_auto)),
    "geico_credit_excluded": float(base.GEICO_AUTO_CREDIT),
    "geico_on_524": False,
    "geico_on_epgc": False,
    "travelers_is_524_home": True,
    "travelers_2025_found": 0.0,
    "travelers_leftover_jan_524_home": 0.0,
    "sep_dec_on_prime": 0.0,
    "unit_split": "50/50 home share (2023-24 Excel had all homeowners on U1) — do not invent a new split",
}


def patch_csv_note(path: Path, old_substr: str, new_note: str) -> None:
    rows = list(csv.reader(path.open(encoding="utf-8")))
    found = False
    for row in rows:
        for i, cell in enumerate(row):
            if old_substr in cell and ("State Farm" in cell or "INSURANCE LOCK" in cell or "Travelers" in cell or "§121" in cell or "CC answers" in cell):
                row[i] = new_note
                found = True
                break
    if not found:
        # last non-empty row in col B for turbo 524
        for row in reversed(rows):
            if len(row) > 1 and row[1] and ("State Farm" in row[1] or "CC answers" in row[1] or "§121" in row[1]):
                row[1] = new_note
                found = True
                break
    if not found:
        raise SystemExit(f"could not patch note in {path}")
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def write_lock_csv() -> None:
    rows = [
        ["Lock", "Status", "Where it hits", "Dollars"],
        [
            "Travelers = prior 524 Ferdinand home",
            "LOCKED 2026-09-20",
            "524 INSURANCE Units 1 and 2 (existing 50/50 home-share mapping — do not invent a new split)",
            "Leftover Jan 2025 $0.00 on Prime/Sapphire/BoA/Monarch",
        ],
        [
            "Geico = cars / personal auto",
            "LOCKED 2026-09-20",
            "Personal auto only. Not 524 INSURANCE. Not EPGC.",
            f"${base.GEICO_AUTO_CREDIT} *AUTO credit 2025-02-11 EXCLUDED personal",
        ],
        [
            "2025 State Farm bundled home+auto",
            "CONFIRMED 2026-09-18 / locked 2026-09-20",
            "HOME share → 524 INSURANCE 50/50 U1/U2 (yellow proxy). AUTO share personal.",
            f"Home ${sum(base.u2_ins)+sum(base.u1_ins)} / Auto ${sum(base.sf_auto)} / billed ${sum((r['billed'] for r in base.sf_split_rows), 0)}",
        ],
        [
            "Grove Collaborative ≠ 827 N Grove the house",
            "LOCKED 2026-09-18",
            "Grove Collaborative → 524 Unit 2 SUPPLIES. HD 1/3 → 827 N Grove the house CapEx.",
            f"Grove Collaborative ${sum(base.u2_sup)} / house HD 1/3 ${sum(ans.grove_hd)}",
        ],
    ]
    with LOCK_CSV.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def write_start_here() -> None:
    START_INS.write_text(
        "START HERE — 524 insurance lock (2026-09-20)\n"
        "\n"
        "Not tax advice. Organizational packet only.\n"
        "\n"
        "User: “Traveler's was home insurance before for Ferdinand and Geico was for the cars.”\n"
        "\n"
        "LOCKED\n"
        "1. Travelers (Traveler’s) = prior homeowners for 524 Ferdinand.\n"
        "   Units 1 and 2 keep the existing 50/50 home-share mapping — do not invent a new split.\n"
        "   Any leftover Travelers 2025 (Jan leftover previously noted) is 524 INSURANCE / Sch E, not auto.\n"
        "   Cash on Prime / Sapphire / BoA / Monarch: $0.00.\n"
        "2. Geico = cars / personal auto. The 2025-02-11 Geico *AUTO credit $540.89 is already EXCLUDED personal.\n"
        "   Do not put Geico on 524 INSURANCE or EPGC.\n"
        "3. 2025 State Farm is the bundled car+home policy that replaced them.\n"
        "   Only the HOME share ($1,105.46; U2 $552.76 / U1 $552.70) hits 524 INSURANCE.\n"
        "   Auto share $1,019.33 is personal — not EPGC, not 524.\n"
        "4. Grove Collaborative ≠ 827 N Grove the house (already locked).\n"
        "\n"
        "ASK no longer asks which carrier was home vs cars.\n"
        "Still open (not carrier identity): State Farm declarations pages for the $ split, and Sep–Dec State Farm not on Prime.\n"
        "\n"
        "Mercury locks unchanged: Coinbase investment; Newstar COGS; Art Sales $150,000;\n"
        "David Aaron Consultant June $13,595; Wise expertise $992.80 Consultant Fees;\n"
        "Koziol/Ariadne $150k pass-through; Aquinas books $1,000 Feb Cost TBD;\n"
        "EOEB remainder $130,270 and L5 $87,396.32 mixed unallocated (not dumped on P&L).\n",
        encoding="utf-8",
    )


def patch_start_here_checking() -> None:
    text = START_CC.read_text(encoding="utf-8")
    old = (
        "Still true from last pass\n"
        "- Grove Collaborative $791.21 on 524 Unit 2 SUPPLIES ≠ 827 N Grove the house.\n"
        "- State Farm home share only on 524 INSURANCE.\n"
    )
    new = (
        "Still true / LOCKED 2026-09-20\n"
        "- Grove Collaborative $791.21 on 524 Unit 2 SUPPLIES ≠ 827 N Grove the house.\n"
        "- Travelers was prior 524 home insurance; Geico was the cars.\n"
        "- Only State Farm HOME share on 524 INSURANCE (50/50 U1/U2). Auto + Geico credit personal, not EPGC.\n"
        "- Travelers leftover Jan 2025 is 524 home ($0 on cards/Monarch).\n"
    )
    if old not in text:
        if "Travelers was prior 524" in text:
            return
        raise SystemExit("START_HERE_checking.txt lock block not found")
    START_CC.write_text(text.replace(old, new), encoding="utf-8")


def patch_source_csv() -> None:
    path = CSV_DIR / "SOURCE_2025.csv"
    rows = list(csv.reader(path.open(encoding="utf-8")))
    found = False
    note = (
        "LOCKED 2026-09-20: Travelers was prior 524 home; Geico was the cars. "
        f"HOME share ${sum(base.u2_ins)+sum(base.u1_ins)} on 524 INSURANCE 50/50 U1/U2 "
        "(yellow proxy from Feb $157.53/$145.25). "
        f"AUTO share ${sum(base.sf_auto)} + Geico credit ${base.GEICO_AUTO_CREDIT} EXCLUDED personal "
        "(not EPGC, not 524). Travelers leftover Jan 2025 is 524 home — $0 on Prime/Sapphire/BoA/Monarch. "
        "No Sep–Dec SF on Prime. Not tax advice."
    )
    for row in rows:
        if row and str(row[0]).startswith("State Farm"):
            row[0] = "Travelers / Geico / State Farm 2026-09-20"
            if len(row) < 2:
                row.append(note)
            else:
                row[1] = note
            found = True
    if not found:
        rows.append(["Travelers / Geico / State Farm 2026-09-20", note])
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def patch_xlsx_notes() -> None:
    wb = load_workbook(XLSX)
    u2 = wb["524 Ferdinand Ave, Unit 2"]
    u2["A39"] = U2_NOTE_LIKE
    u2["A39"].alignment = Alignment(wrap_text=True)
    u2["A39"].fill = base.YELLOW
    u2.row_dimensions[39].height = 72

    u1 = wb["524 Ferdinand Ave, Unit 1"]
    u1["A45"] = U1_NOTE_LIKE
    u1["A45"].alignment = Alignment(wrap_text=True)
    u1["A45"].fill = base.YELLOW
    u1.row_dimensions[45].height = 72

    # Insurance dollars already home-share only — do not rewrite months.
    wb.save(XLSX)
    if DELIVERABLE.parent.exists():
        shutil.copy2(XLSX, DELIVERABLE)


def export_helper_sheets() -> None:
    """Build ASK / SF_SPLIT / CC_LEDGER in a throwaway book so the like-last-year xlsx stays intact."""
    tw = Workbook()
    tw.remove(tw.active)
    # checking-answer ASK overlay (same as apply_checking_answers, without touching P&L)
    comed_open = []
    for item, detail in ans.STILL_OPEN:
        if item.startswith("REMIND JACOB"):
            detail = (
                "Joint 0203 statements show ComEd ACH (e.g. 1/21 Jacob $59.04 + Megan $132.94). "
                + detail
            )
        comed_open.append((item, detail))
    confirmed = list(ans.CONFIRMED) + list(chk.ANSWERS_CONFIRMED)
    still = comed_open + list(chk.ANSWERS_STILL_OPEN)
    orig_c, orig_s = list(ans.CONFIRMED), list(ans.STILL_OPEN)
    ans.CONFIRMED[:] = confirmed
    ans.STILL_OPEN[:] = still
    try:
        ans.rebuild_ask(tw)
        chk.append_ask_inflow_totals(tw["ASK"])
        ans.rebuild_ledger(tw)
        sf.write_sf_split(tw)
        ans.sheet_to_csv(tw["ASK"], CSV_DIR / "ASK.csv")
        ans.sheet_to_csv(tw["CC_LEDGER"], CSV_DIR / "CC_LEDGER.csv")
        ans.sheet_to_csv(tw["SF_SPLIT"], CSV_DIR / "SF_SPLIT.csv")
        ans.save_sheet_part(tw["ASK"], ans.PARTS / "ASK.xlsx")
        ans.save_sheet_part(tw["CC_LEDGER"], ans.PARTS / "CC_LEDGER.xlsx")
        ans.save_sheet_part(tw["SF_SPLIT"], ans.PARTS / "SF_SPLIT.xlsx")
    finally:
        ans.CONFIRMED[:] = orig_c
        ans.STILL_OPEN[:] = orig_s


def export_524_like_csv() -> None:
    wb = load_workbook(XLSX)
    ans.sheet_to_csv(wb["524 Ferdinand Ave, Unit 2"], CSV_DIR / "524_Unit_2_like_last_year.csv")
    ans.sheet_to_csv(wb["524 Ferdinand Ave, Unit 1"], CSV_DIR / "524_Unit_1_like_last_year.csv")
    ans.save_sheet_part(wb["524 Ferdinand Ave, Unit 2"], ans.PARTS / "524 Ferdinand Ave Unit 2.xlsx")
    ans.save_sheet_part(wb["524 Ferdinand Ave, Unit 1"], ans.PARTS / "524 Ferdinand Ave Unit 1.xlsx")
    wb.close()


def write_json() -> None:
    payload = json.loads(ans.JSON_PATH.read_text(encoding="utf-8")) if ans.JSON_PATH.exists() else {}
    payload["updated"] = "2026-09-20"
    payload["state_farm"] = STATE_FARM
    payload["confirmed"] = [a[0] for a in ans.CONFIRMED] + [a[0] for a in chk.ANSWERS_CONFIRMED]
    payload["still_open"] = [a[0] for a in ans.STILL_OPEN] + [a[0] for a in chk.ANSWERS_STILL_OPEN]
    ans.JSON_PATH.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    pkt = json.loads(PACKET.read_text(encoding="utf-8")) if PACKET.exists() else {}
    pkt["updated"] = "2026-09-20"
    pkt["insurance_lock"] = STATE_FARM
    pkt["note"] = (
        "Grove Collaborative ≠ 827 N Grove the house. "
        "Travelers = prior 524 home; Geico = personal auto. Only State Farm home share on 524 INSURANCE."
    )
    PACKET.write_text(json.dumps(pkt, indent=2), encoding="utf-8")


def main() -> None:
    assert XLSX.exists(), XLSX
    # Geico must not be on 524 insurance months
    assert all(v >= 0 for v in base.u2_ins + base.u1_ins)
    assert float(sum(base.u2_ins) + sum(base.u1_ins)) == 1105.46
    assert float(sum(base.sf_auto)) == 1019.33
    assert base.u2_ins[0] == 0 and base.u1_ins[0] == 0  # Jan leftover Travelers $0

    patch_xlsx_notes()
    export_helper_sheets()
    patch_csv_note(CSV_DIR / "524_Unit_2.csv", "State Farm", TURBO_U2_NOTE)
    patch_csv_note(CSV_DIR / "524_Unit_1.csv", "State Farm", TURBO_U1_NOTE)
    patch_source_csv()
    write_lock_csv()
    write_start_here()
    patch_start_here_checking()
    export_524_like_csv()
    write_json()

    print("home U2", sum(base.u2_ins), "U1", sum(base.u1_ins))
    print("auto excluded", sum(base.sf_auto), "geico", base.GEICO_AUTO_CREDIT)
    print("travelers leftover", 0)
    print("geico on 524", False)
    print("saved", XLSX, "bytes", XLSX.stat().st_size)


if __name__ == "__main__":
    main()

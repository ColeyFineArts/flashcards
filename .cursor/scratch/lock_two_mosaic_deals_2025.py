#!/usr/bin/env python3
"""Unbundle Plutus/Antiquarium from Jamal mosaics Cost.

User 2026-09-21 follow-up: the 40/50 split was wrong. Plutus / DBA Antiquarium
is a different mosaic deal, also involving EOEB. EOEB may have paid for both
at the same time (2/21 $105,000 MAKE A PAYMENT). Amounts can look the same;
the transfer mix is different.

LOCKED cash (do not recast without saying so)
  EOEB 2/21 +$105,000 stays February Art Sales gross (pooled proceeds).
  Plutus/Antiquarium 2/24 $10 + 2/25 $49,990 = $50,000 Cost of Deal B.
  Combined profit $15,000 stays WORKING in I8 (user; not un-said).
  Combined Cost $90,000 = $105k − $15k WORKING, not “Jamal 50+40”.

UNLOCKED
  Plutus is NOT Jamal DBA paying Deal A Cost.
  $40k BoA wire to Jamal is not on Monarch 9922. Do not invent it.
  2/24 $90k EPGC→9922 is an owner transfer. Not extra Cost. Not a locked
  $40k Cost-funding split. BoA 9922 draws back to $261,172.

Do NOT dump EOEB remainder $130,270. Do not touch Belt / Venus / August.
Not tax advice. Do not re-run apply_checking_answers_2025.py.
"""
from __future__ import annotations

import csv
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

import sys

sys.path.insert(0, str(Path("/workspace/.cursor/scratch")))
from sanitize_xlsx_for_sheets import sanitize_xlsx

ROOT = Path("/workspace/.cursor/scratch")
PARTS = ROOT / "tax_turbo_parts"
CSV_DIR = ROOT / "cc_fill_csv"
XLSX = PARTS / "Personal Income 2025.xlsx"
DELIVERABLE = ROOT / "tax_turbo_deliverable" / "Personal Income 2025.xlsx"
PACKET = ROOT / "cpa_packet_v1.json"
MERCURY_JSON = ROOT / "mercury_2025.json"

GREEN = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
CG = Font(name="Century Gothic", size=10)
WRAP = Alignment(wrap_text=True, vertical="top")

OPEN_XLSX = "https://docs.google.com/spreadsheets/d/1uAipm4kgyiZiEPE2Eme6_RWKNXfHQ7XqBEIyFTLzAV4/edit"
OPEN_07 = "https://docs.google.com/spreadsheets/d/1t2dRKpWmN4cT-bqmJxBtIovzvYmoJFrNr9P1bXaCeUg/edit"
OPEN_ROOT = "https://docs.google.com/spreadsheets/d/175x6fMYC0zVads6l4QyuY4Zkfn-xOeNm7SgKe-yxmJM/edit"
LABELED_ART = "https://docs.google.com/spreadsheets/d/1OfHYP14zIuJKGwVZHr1HmaAY-eTI-sJJN1ghqZp-wP0/edit"
CLASSIFIER = "https://docs.google.com/spreadsheets/d/1UglEtTDOb1Mu2G9HWqcggTLmRAKNqwhOMG-SdXINt0k/edit"
ASK = "https://docs.google.com/spreadsheets/d/1wUQF1PjXhbnFzNu1t03iB0KmhztCWSl3rSVUE_oM61w/edit"
FORTUNA = "https://docs.google.com/spreadsheets/d/1TK93q6iwXuoWodkikxwFNKgGkfDiGgyk1RpvMEKmPPw/edit"
START_HERE = "https://docs.google.com/document/d/1RCiZrR-aE3qsza89JshNEQoqnzcth8hMsZr77Psxx_Q/edit"

MOSAIC_NOTE = (
    "UNLOCKED 2026-09-21 40/50 mix. TWO mosaic deals, both EOEB. "
    "Deal A: Three Ancient Mosaics / Jamal Rifai — Cost paid from personal BoA (amount ASK). "
    "Deal B: Plutus & Mnemosyne / DBA Antiquarium — Cost $50,000 on Mercury 8291 (2/24 $10 + 2/25 $49,990). "
    "EOEB 2/21 $105,000 MAKE A PAYMENT is pooled proceeds (user: perhaps paid for both at once). "
    "Do NOT treat Plutus as Jamal DBA. Combined Sale $105,000 LOCKED as February gross. "
    "Combined profit $15,000 WORKING (user; I8). Combined Cost $90,000 WORKING ($105k−$15k), "
    "not a locked Jamal 50+40. 2/24 $90k EPGC→9922 is owner transfer, not extra Cost. "
    "Monarch 9922 has no labeled $40k/$90k OUT to Jamal in Feb (only $90k IN from EPGC). "
    "Split of the $105k and Jamal BoA Cost still ASK. Not tax advice."
)
ART_FOOTER = (
    "Mercury 8291 MATCHED/LOCKED: Berk $14,000 (1/10 seals) + $30,000 (2/7 lots). "
    "EOEB 2/21 $105,000 mosaics = TWO deals pooled (Jamal Deal A + Plutus/Antiquarium Deal B). "
    "Plutus $50,000 is Deal B Cost, not Jamal. Combined profit $15,000 WORKING in I8. "
    "Fortuna $13,000 (1/8 seals Cost). "
    "Roman Gold Belt SALE to L5: Cost $45,000 / Sale $50,000 / net $5,000. "
    "Venus joint SALE EPGC share: Cost $20,000 / proceeds $26,000 / net $6,000. "
    "August Fortuna joint SALE EPGC share: Cost $20,000 / proceeds $25,000 / net $5,000 "
    "(Sale Price = remittance, not full hammer; object unnamed). "
    "Newstar $23,581 jewelry fabrication COGS parked below. "
    "David Aaron $13,595 LOCKED consultant (EPGC Consultant June — not a sale). "
    "Aquinas Hobor $1,000 LOCKED book sale (Feb cash; Cost TBD, not in I8). "
    "Do not dump EOEB remainder $130,270 (see Deal Classifier per invoice) or other "
    "L5 $37,396.32. Not tax advice."
)
INCOME_NOTE = (
    "2025 Actual: I4 Monarch paycheck cash ≠ W-2 Box 1. I5 216 LTR cash $19,500. "
    "I6 524 #2 STR platform net LOCKED. I7 GCM 1099-NEC $21,500. "
    "I8 Art net $39,055.06 = Berk lots + seals $1,000 + mosaics combined $15,000 WORKING "
    "+ Belt net $5,000 + Venus EPGC-share net $6,000 + August Fortuna joint net $5,000. "
    "Mosaics: TWO EOEB deals pooled in 2/21 $105,000. Plutus/Antiquarium $50k is Deal B Cost "
    "(not Jamal). Jamal Deal A Cost from personal BoA ASK. Combined profit $15k WORKING. "
    "Excludes Sale 6428 $11,000 until Cost. Excludes Aquinas books $1,000 until Cost. "
    "David Aaron $13,595 is EPGC Consultant June (not I8). "
    "EOEB remainder $130,270 stays UNALLOCATED (per-invoice Deal Classifier; proposed "
    "FEE, no matching Cost — not dumped). Other L5 $37,396.32 / Erdal 11/7 $6,000 "
    "/ Aysel $50,000 still ASK. Not tax advice."
)
EPGC_NOTE = (
    "2025 Mercury 2026-09-21: Art Sales gross cash by month = Jan $14,000 + Feb "
    "$136,000 MATCHED + July $76,000 (Belt L5 $50,000 + Venus remittance $26,000) "
    "+ October $25,000 (August Fortuna joint remittance) = $251,000. "
    "Feb mosaics $105,000 is TWO EOEB deals pooled (Jamal + Plutus/Antiquarium). "
    "Plutus $50k on 8291 is Deal B Cost, not Jamal. Combined profit $15,000 WORKING in I8. "
    "Consultant June $13,595 = David Aaron Limited LOCKED (not a sale). "
    "Consultant Fees July $334.17 + October $658.63 = Wise expertise write-ups LOCKED ($992.80). "
    "EOEB remainder $130,270 is per-invoice UNALLOCATED on the Deal Classifier "
    "(proposed FEE) — not dumped. Other L5 $37,396.32 MIXED — not dumped. "
    "Erdal 11/7 $6,000 FORTUNA PAYMENT still ASK. "
    "Aysel Monarch +$50k vs native Failed −$50k still ASK (not the belt). "
    "Koziol/Ariadne $150,000 LOCKED pass-through — not P&L. Coinbase is Investments. "
    "Newstar $23,581 is Art Sales COGS. Dec Wise $565.75 reimbursed. Aquinas books "
    "Cost TBD. Joint Sale Price = EPGC proceeds, not full hammer. GCM 1099 is "
    "personal. Not tax advice."
)
LEDGER_BANNER = (
    "Mercury Checking 8291 — EPGC LLC 2025. 64 unique cash txns. "
    "Deal Classifier: SALE vs FEE vs PASS vs UNALLOCATED. "
    "TWO mosaic deals (user 2026-09-21): Plutus/DBA Antiquarium is NOT Jamal. "
    "EOEB 2/21 $105,000 pooled SALE. Plutus 2/24–25 $50,000 = Deal B Cost. "
    "Jamal Deal A Cost from personal BoA ASK. Combined profit $15,000 WORKING. "
    "2/24 $90k EPGC→9922 = owner transfer, not extra Cost (draws $261,172). "
    "Roman Gold Belt = SALE. Venus = joint SALE. August Fortuna = joint SALE. "
    "Remaining EOEB $130,270 per-invoice UNALLOCATED (proposed FEE — not dumped). "
    "Other L5 $37,396.32 UNALLOCATED. David Aaron = FEE. Koziol/Ariadne = PASS. "
    "11/7 $6,000 FORTUNA PAYMENT still ASK. Not tax advice."
)
DS_NOTE = (
    "Mosaics 2026-09-21: TWO EOEB deals. Plutus/Antiquarium $50,000 on 8291 is Deal B Cost, "
    "not Jamal DBA. EOEB 2/21 $105,000 pooled SALE. Combined profit $15,000 WORKING in I8. "
    "Jamal BoA Cost amount ASK. 2/24 $90k EPGC→9922 is transfer, not extra Cost. "
    "Do not dump EOEB remainder $130,270. Not tax advice."
)
START_HERE_MERCURY = f"""START HERE — Mercury Bank 8291 (two mosaic deals)

Not tax advice. Same Personal Income.xlsx tabs as last year. Organizational packet only.

Mercury Checking 8291 is the EPGC LLC operating account (Choice Financial).
64 unique 2025 cash transactions from Monarch IDs.

OPEN THIS 11-tab workbook (numbers; Google convert dropped Art object labels):
{OPEN_XLSX}
07_income copy:
{OPEN_07}
root copy:
{OPEN_ROOT}

Art Sales NAMES (use this for object labels — Belt / Venus / August / two mosaic deals):
{LABELED_ART}

Deal Classifier (per-invoice EOEB; remainder $130,270 UNALLOCATED — do not dump):
{CLASSIFIER}

ASK (Q10 mosaics mix REOPENED — two deals):
{ASK}

Fortuna inventory LOCKED SOLD 2025 (Belt, Venus, August — not parked unsold):
{FORTUNA}

LOCKED this pass
- TWO mosaic deals, both EOEB. User 2026-09-21: 40/50 was wrong. Plutus / DBA Antiquarium is a different mosaic deal, not Jamal.
  Deal A — Three Ancient Mosaics / Jamal Rifai. Cost from personal BoA (amount ASK). $10 Mercury tests to Jamal 2/21 and 2/24 only.
  Deal B — Plutus & Mnemosyne / DBA Antiquarium. Cost $50,000 on Mercury 8291 (native COGS 2/24 $10 + 2/25 $49,990).
  EOEB 2/21 $105,000 MAKE A PAYMENT is pooled proceeds (perhaps both at once). February Art Sales gross LOCKED.
  Combined profit $15,000 WORKING in I8 (user; not un-said). Combined Cost $90,000 WORKING ($105k−$15k), not a locked Jamal 50+40.
  2/24 $90k EPGC→9922 is owner transfer, not extra Cost. BoA 9922 draws $261,172.
- Roman Gold Belt SALE: Fortuna Cost $45,000 / L5 Sale $50,000 / net $5,000 → EPGC Art Sales July + I8.
- Venus joint SALE: Fortuna Cost $20,000 / Erdal proceeds $26,000 / net $6,000 → EPGC Art Sales July + I8.
- August Fortuna joint SALE: Fortuna Cost $20,000 / Erdal proceeds $25,000 / net $5,000 → EPGC Art Sales October + I8.
- David Aaron Limited $13,595 (6/24) → EPGC Consultant June. FEE, not a sale.
- Koziol $150,000 (10/20) / Ariadne $150,000 (10/21) → PASS. Not P&L.
- Wise expertise write-ups $992.80 → EPGC Consultant Fees.
- EOEB 12/2 $565.75 reimbursed — not income.

EPGC 2025 Art Sales LOCKED cash $251,000 (Jan $14,000 / Feb $136,000 / July $76,000 / Oct $25,000).
EPGC Consultant June is $13,595 (David Aaron only).
Income I8 Art net is $39,055.06 (includes mosaics combined profit $15,000 WORKING). I10 EBITDA =SUM(I4:I9) → $167,760.

ASK remaining — do not dump
1. EOEB remainder $130,270.00 — per-invoice UNALLOCATED (proposed FEE, no matching Cost). Jan 16 $1,000 Canosan-horse OPEN. 5/1 $21,000 is NOT Venus. Fortuna outs are NOT EOEB Cost.
2. Other L5 $37,396.32 after removing $50k belt.
3. Erdal 11/7 $6,000 FORTUNA PAYMENT — not August, not Venus.
4. Aysel Dere — Monarch 7/18 +$50,000 IN vs native Failed 7/17 −$50,000 OUT. Not the belt.
10. REOPENED — split of EOEB $105k between Deal A (Jamal) and Deal B (Plutus/Antiquarium). Jamal BoA Cost amount (9922 statement). Whether $15k profit is combined or Deal A only.

Sale 6428 $11,000 and Aquinas books $1,000 wait on Cost (not in I8).
"""
START_HERE_ALL = f"""START HERE — last year's spreadsheet, 2025 numbers filled
Not tax advice. Organizational packet only.

OPEN THIS ONE spreadsheet (same 11 tab names as last year, 2023 | 2024 | 2025 year blocks, 2025 numbers, formulas that calculate; I8 $39,055.06; EPGC July $76,000 + Oct $25,000; TWO mosaic deals pooled in EOEB $105k). Google convert kept Art numbers but dropped object labels:
{OPEN_XLSX}

07_income copy:
{OPEN_07}

2025 Taxes root copy:
{OPEN_ROOT}

Art Sales NAMES (2022–2024 lots + two mosaic deals + Belt / Venus / August):
{LABELED_ART}

Deal Classifier (per-invoice EOEB — SALE/FEE/PASS/UNALLOCATED; remainder $130,270 NOT dumped):
{CLASSIFIER}

Fortuna inventory LOCKED SOLD 2025 (Belt, Venus, August — not parked unsold):
{FORTUNA}

Tabs in last year's order (all inside that one file)
1. Income — 2022–2025. 2025 Actual: Hindman $70,618 / 216 $19,500 / 524 #2 $17,086.94 / GCM $21,500 / Art $39,055.06. I10 EBITDA =SUM(I4:I9) → $167,760.
2. 524 Ferdinand Ave, Unit 2
3. 216 N. Oak Park Ave
4. EPGC LLC — 2025 Art Sales $14,000 + $136,000 + July $76,000 + October $25,000 = $251,000; Consultant June $13,595 (David Aaron only).
5. Refrence Library
6. Art Sales and Purchases — 2022–2024 lots, then Sales 2025 including TWO mosaic deals (Jamal Deal A + Plutus/Antiquarium Deal B, EOEB pooled $105k), Roman Gold Belt, Venus joint, August Fortuna joint. Do not dump EOEB $130,270.
7. GCM
8. Hindman W2
9. Megan T4
10. Investments — Coinbase ACH net -$51,000.
11. 524 Ferdinand Ave, Unit 1

Mosaics (user 2026-09-21): 40/50 mix UNLOCKED. Plutus / DBA Antiquarium is a different mosaic deal, also EOEB. Perhaps EOEB paid for both at once. Combined Sale $105,000 LOCKED. Combined profit $15,000 WORKING. Plutus $50k on Mercury is Deal B Cost, not Jamal. Jamal Cost from personal BoA ASK. Do not double-count the 2/24 $90k EPGC→9922 transfer as extra Cost. BoA 9922 draws $261,172.

Deal Classifier: SALE vs FEE vs PASS vs UNALLOCATED.
Belt = SALE. Venus = SALE (joint). August Fortuna = SALE (joint). Mosaics = two SALE deals, pooled cash. EOEB remainder = UNALLOCATED per invoice (proposed FEE). David Aaron = FEE.

START HERE Doc:
{START_HERE}

Not tax advice.
"""


def paint(cell, value=None, fill=None, num=None):
    if value is not None:
        cell.value = value
    if fill is not None:
        cell.fill = fill
    if num is not None:
        cell.number_format = num
    cell.font = CG


def patch_xlsx() -> None:
    wb = load_workbook(XLSX)
    art = wb["Art Sales and Purchases"]
    if art["A74"].value is None or "Mosaic" not in str(art["A74"].value):
        raise SystemExit(f"Art A74 is {art['A74'].value!r}")
    paint(art["A74"], "Two mosaic deals (EOEB 2/21 pooled)")
    paint(art["B74"], "Jamal Rifai (Deal A) + Plutus/Antiquarium (Deal B)")
    paint(art["C74"], 90000, fill=YELLOW)
    paint(art["F74"], 105000, fill=GREEN)
    paint(art["I74"], MOSAIC_NOTE, fill=YELLOW)
    art["I74"].alignment = WRAP
    art.row_dimensions[74].height = 96
    for r in range(78, 95):
        a = str(art.cell(r, 1).value or "")
        if a.startswith("Mercury 8291 MATCHED") or "Plutus" in a or "mosaics" in a.lower():
            paint(art.cell(r, 1), ART_FOOTER, fill=YELLOW)
            art.cell(r, 1).alignment = WRAP
            art.row_dimensions[r].height = 84
            break

    inc = wb["Income"]
    paint(inc["A12"], INCOME_NOTE)
    inc["A12"].alignment = WRAP

    ep = wb["EPGC LLC"]
    for r in range(74, 95):
        v = str(ep.cell(r, 1).value or "")
        if "2025 Mercury" in v or "Art Sales gross" in v:
            paint(ep.cell(r, 1), EPGC_NOTE)
            ep.cell(r, 1).alignment = WRAP
            break

    ds = wb["2025 Data sources"]
    wrote = False
    for r in range(1, 50):
        v = str(ds.cell(r, 1).value or "")
        if v.startswith("Mosaics Cost") or "40,000 of Cost" in v or "Plutus $50,000 on Mercury" in v:
            paint(ds.cell(r, 1), DS_NOTE, fill=YELLOW)
            wrote = True
    if not wrote:
        ds.cell((ds.max_row or 1) + 1, 1, DS_NOTE)

    try:
        wb.defined_names.clear()
    except Exception:
        pass
    wb._external_links = []
    wb.save(XLSX)
    sanitize_xlsx(XLSX)
    DELIVERABLE.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(XLSX, DELIVERABLE)

    def sheet_to_csv(ws, path: Path) -> None:
        max_r, max_c = ws.max_row or 1, ws.max_column or 1
        rows = []
        for r in range(1, max_r + 1):
            rows.append([ws.cell(r, c).value for c in range(1, max_c + 1)])
        with path.open("w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerows(rows)

    sheet_to_csv(art, CSV_DIR / "Art_Sales.csv")
    sheet_to_csv(inc, CSV_DIR / "Income.csv")
    sheet_to_csv(ep, CSV_DIR / "EPGC_LLC.csv")


def patch_ledger() -> None:
    path = CSV_DIR / "MERCURY_LEDGER.csv"
    rows = list(csv.reader(path.open(encoding="utf-8")))
    rows[0] = [LEDGER_BANNER] + [""] * max(0, len(rows[0]) - 1)
    header = rows[2]
    idx = {name: i for i, name in enumerate(header)}
    date_i, who_i, amt_i = idx["Date"], idx["Counterparty"], idx["Amount"]
    notes_i = idx["Notes"]
    deal_i = idx["Deal / object"]
    bkt_i = idx["Bucket"]
    st_i = idx["Status"]
    clf_i = idx.get("Classifier")
    for r in rows[3:]:
        if len(r) <= notes_i or not r[date_i]:
            continue
        date, who, amt = r[date_i], r[who_i], r[amt_i]
        if date == "2025-02-24" and "9922" in who and amt.startswith("-90000"):
            r[bkt_i] = "TRANSFER"
            r[st_i] = "LOCKED"
            r[deal_i] = "Owner transfer to BoA 9922 (not extra mosaics Cost)"
            r[notes_i] = (
                "UNLOCKED 40/50 Cost-funding split. User: Plutus is a different mosaic deal. "
                "This $90,000 EPGC→9922 is an owner transfer of personal, NOT extra Cost and "
                "NOT a locked $40k Jamal funding slice. Jamal Deal A Cost from 9922 is ASK "
                "(no labeled $40k/$90k OUT to Jamal on Monarch 9922 in Feb). BoA draws $261,172. "
                "Not tax advice."
            )
        if date == "2025-02-25" and "Plutus" in who:
            r[bkt_i] = "ART_PURCHASE"
            if clf_i is not None:
                r[clf_i] = "SALE"
            r[st_i] = "LOCKED Cost / ASK sale split"
            r[deal_i] = "Mosaic Deal B — Plutus/Antiquarium Cost $50,000 (NOT Jamal Deal A)"
            r[notes_i] = (
                "Native Mercury: DBA Antiquarium (Plutus & Mnemosyne, LLC) COGS. "
                "With the 2/24 $10 test this is $50,000 Cost of a different mosaic deal "
                "(user 2026-09-21). Do NOT pair as Jamal Cost of Three Ancient Mosaics. "
                "Matching Sale is inside EOEB 2/21 $105,000 pooled cash, split ASK. Not tax advice."
            )
        if date == "2025-02-24" and "Plutus" in who:
            r[deal_i] = "Mosaic Deal B — $10 test of Plutus/Antiquarium $50,000 Cost"
            r[notes_i] = (
                "Native: DBA Antiquarium (Plutus & Mnemosyne, LLC). First $10 of Deal B Cost. "
                "Not Jamal. $49,990 follows 2/25."
            )
        if date == "2025-02-21" and who == "EOEB LLC" and amt.startswith("105000"):
            r[deal_i] = "Two mosaic deals — pooled Sale Price (Jamal Deal A + Plutus Deal B)"
            r[st_i] = "LOCKED cash / ASK split"
            r[notes_i] = (
                "EOEB LLC MAKE A PAYMENT $105,000. User: perhaps paid for both mosaic deals "
                "at once. Hits EPGC Art Sales February. Split of $105k between Deal A (Jamal / "
                "Yantis mosaics) and Deal B (Plutus/Antiquarium) ASK. Combined profit $15,000 "
                "WORKING. Do not dump. Not tax advice."
            )
        if date in ("2025-02-21", "2025-02-24") and "Jamal" in who:
            r[deal_i] = "$10 test wire to Jamal M. Rifai (Deal A counterparty)"
            r[notes_i] = (
                "Not P&L. Mercury could not send the Jamal Cost here — user paid Deal A from "
                "personal BoA. Distinct from Plutus/Antiquarium Deal B."
            )
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def patch_classifier() -> None:
    path = CSV_DIR / "DEAL_CLASSIFIER.csv"
    rows = list(csv.reader(path.open(encoding="utf-8")))
    out = []
    inserted_plutus = False
    has_plutus = any(
        len(r) >= 3 and r[0] == "2025-02-25" and "Plutus" in (r[2] or "") for r in rows
    )
    for r in rows:
        if len(r) >= 7 and r[0] == "2025-02-21" and "Mosaic" in r[2]:
            r[2] = "Two mosaic deals / EOEB MAKE A PAYMENT (pooled)"
            r[3] = (
                "Deal B Cost LOCKED: Plutus/Antiquarium 2/24–25 $50,000 on 8291. "
                "Deal A Cost ASK: Jamal from personal BoA (amount unknown; 40/50 UNLOCKED). "
                "2/24 $90k EPGC→9922 is transfer, not extra Cost. Combined profit $15,000 WORKING."
            )
            r[4] = "SALE"
            r[5] = (
                "User 2026-09-21: 40/50 was wrong. Plutus/DBA Antiquarium is a different mosaic "
                "deal, also EOEB. Perhaps EOEB paid for both at the same time. $105k stays "
                "February Art Sales gross. Split ASK. Not dumped."
            )
            r[6] = "LOCKED cash / ASK per-deal split"
            out.append(r)
            if not inserted_plutus and not has_plutus:
                out.append(
                    [
                        "2025-02-25",
                        "-50000.00",
                        "Mosaic Deal B / Plutus & Mnemosyne DBA Antiquarium",
                        "this row is Deal B Cost ($10 on 2/24 + $49,990 on 2/25)",
                        "SALE",
                        "Native Mercury COGS. User: different mosaic deal from Jamal. Matching Sale is inside EOEB 2/21 $105k pooled, split ASK. Not tax advice.",
                        "LOCKED Cost / ASK sale split",
                    ]
                )
                inserted_plutus = True
            continue
        out.append(r)
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(out)


def patch_ask() -> None:
    path = CSV_DIR / "ASK_Mercury.csv"
    rows = list(csv.reader(path.open(encoding="utf-8")))
    out = []
    for r in rows:
        line = ",".join(r)
        if r and (
            str(r[0]).startswith("10. Mosaics") or str(r[0]).startswith("10. TWO mosaic")
        ):
            out.append(
                [
                    "10. TWO mosaic deals — split REOPENED (Plutus ≠ Jamal)",
                    "User 2026-09-21: 40/50 was wrong. Plutus/DBA Antiquarium is a different mosaic deal, also EOEB. Perhaps EOEB paid for both at once in 2/21 $105,000. Need: (a) split of $105k Deal A vs Deal B; (b) Jamal BoA Cost amount on 9922 statement; (c) whether $15k profit is combined or Deal A only. Plutus $50k Cost on 8291 LOCKED as Deal B. Combined $105k/$15k stays WORKING on the P&L until the split. Do not dump. Monarch 9922 has no labeled $40k OUT to Jamal in Feb.",
                    "SALE (two deals; split ASK)",
                ]
            )
            continue
        if r and str(r[0]).startswith("Mosaics"):
            out.append(
                [
                    "Mosaics (two deals)",
                    "EOEB $105,000 pooled 2/21. Deal B Plutus/Antiquarium Cost $50,000 on 8291. Deal A Jamal Cost from BoA ASK. Combined profit $15,000 WORKING.",
                    "SALE (split ASK)",
                ]
            )
            continue
        if r and str(r[0]).startswith("ASK — remaining"):
            r[0] = (
                "ASK — remaining Mercury 8291 questions after Belt + Venus + August Fortuna SALE lock. "
                "TWO mosaic deals (Plutus ≠ Jamal). Q10 mix REOPENED. EOEB remainder per-invoice "
                "UNALLOCATED (proposed FEE) — not dumped. Q6–Q9 LOCKED. Not tax advice."
            )
        out.append(r)
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(out)


def patch_lock_q() -> None:
    path = CSV_DIR / "LOCK_Mercury_Q6_Q9.csv"
    rows = list(csv.reader(path.open(encoding="utf-8")))
    out = []
    for r in rows:
        if r and r[0] == "EOEB mosaics":
            out.append(
                [
                    "EOEB mosaics cash",
                    "LOCKED cash / ASK split",
                    "EOEB 2/21 $105,000 pooled TWO mosaic deals",
                    "EPGC Art Sales February. Classifier: SALE (split ASK).",
                ]
            )
            continue
        if r and str(r[0]).startswith("10 mosaics"):
            out.append(
                [
                    "10 mosaics mix",
                    "REOPENED",
                    "Plutus/Antiquarium is Deal B Cost $50k, not Jamal. Jamal BoA Cost ASK. Combined profit $15k WORKING.",
                    "Do not treat 40/50 as one Jamal payment. 2/24 $90k EPGC→9922 is transfer, not extra Cost. Draws $261,172.",
                ]
            )
            continue
        if r and r[0] == "LOCKED totals":
            r[2] = (
                'Art Sales gross $251,000 (Jan 14k + Feb 136k + July 76k + Oct 25k); '
                "I8 $39,055.06 (mosaics combined $15k WORKING); Consultant June $13,595"
            )
        out.append(r)
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(out)


def patch_json() -> None:
    packet = json.loads(PACKET.read_text(encoding="utf-8"))
    packet["updated"] = "2026-09-21"
    note = (
        "2026-09-21: TWO mosaic deals. Plutus/DBA Antiquarium ≠ Jamal. "
        "EOEB 2/21 $105,000 pooled SALE LOCKED as February gross. Plutus $50k Deal B Cost LOCKED. "
        "Jamal BoA Cost ASK. Combined profit $15,000 WORKING in I8 $39,055.06. "
        "2/24 $90k EPGC→9922 is transfer (draws $261,172), not extra Cost. "
        "EOEB remainder $130,270 UNALLOCATED. OPEN THIS numbers 1uAipm4k; labeled Art names "
        "1OfHYP14. ASK 1wUQF1Pj. Classifier 1UglEtTD. START HERE 1RCiZrR. Not tax advice."
    )
    packet["xlsx_note"] = note
    m = packet.setdefault("mercury_8291", {})
    m["updated"] = "2026-09-21"
    locked = m.setdefault("locked", {})
    locked["mosaics_sale"] = 105000.0
    locked["mosaics_cost"] = 90000.0
    locked["mosaics_profit"] = 15000.0
    locked["mosaics_plutus_8291"] = 50000.0
    locked.pop("mosaics_boa_9922", None)
    locked["mosaics_two_deals"] = True
    locked["mosaics_cost_note"] = (
        "Combined Cost $90,000 WORKING ($105k sale − $15k profit). "
        "Deal B Plutus/Antiquarium $50,000 LOCKED on 8291. Deal A Jamal BoA ASK. "
        "Not a locked 50+40 Jamal payment."
    )
    locked["boa_9922_draws"] = 261172.0
    locked["boa_9922_draws_note"] = (
        "Restored $261,172. The $40k pull assumed Plutus $50k was Jamal Cost. "
        "2/24 $90k EPGC→9922 is an owner transfer. Jamal funding slice ASK."
    )
    locked["art_net_locked"] = 39055.06
    ask = m.setdefault("ask", {})
    ask.pop("mosaics_boa_40k_wire_label", None)
    ask["mosaics_two_deal_split"] = 105000.0
    ask["mosaics_jamal_boa_cost"] = True
    ask["mosaics_two_deal_note"] = (
        "Split of EOEB $105k Deal A (Jamal) vs Deal B (Plutus/Antiquarium). "
        "Jamal Cost amount on 9922 statement. Whether $15k profit is combined or Deal A only. "
        "Do not dump. Not tax advice."
    )
    dc = packet.setdefault("deal_classifier", {})
    dc["eoeb_mosaics"] = "SALE (two deals pooled in $105k; split ASK)"
    dc["plutus_antiquarium"] = "SALE Deal B Cost $50k LOCKED"
    dc["jamal_mosaics"] = "SALE Deal A; Cost from personal BoA ASK"
    PACKET.write_text(json.dumps(packet, indent=2) + "\n", encoding="utf-8")

    mercury = json.loads(MERCURY_JSON.read_text(encoding="utf-8"))
    mercury["updated"] = "2026-09-21"
    mercury.setdefault("locked", {}).update(locked)
    mercury["locked"].pop("mosaics_boa_9922", None)
    mercury.setdefault("ask", {}).update(ask)
    mercury["ask"].pop("mosaics_boa_40k_wire_label", None)
    mercury.setdefault("deal_classifier", {}).update(
        {
            "eoeb_mosaics": dc["eoeb_mosaics"],
            "plutus_antiquarium": dc["plutus_antiquarium"],
            "jamal_mosaics": dc["jamal_mosaics"],
        }
    )
    MERCURY_JSON.write_text(json.dumps(mercury, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    patch_xlsx()
    patch_ledger()
    patch_classifier()
    patch_ask()
    patch_lock_q()
    (CSV_DIR / "START_HERE_mercury.txt").write_text(START_HERE_MERCURY, encoding="utf-8")
    (CSV_DIR / "START_HERE_all_tabs.txt").write_text(START_HERE_ALL, encoding="utf-8")
    patch_json()
    print("unbundled Plutus Deal B from Jamal Deal A")
    print("I8 unchanged 39055.06 (combined profit WORKING)")
    print("boa draws restored 261172")


if __name__ == "__main__":
    main()

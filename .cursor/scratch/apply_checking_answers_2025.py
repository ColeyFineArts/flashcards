#!/usr/bin/env python3
"""Apply 2026-09-18 checking-inflow answers onto Personal Income 2025 Tax Turbo.

User answers:
  1  Ava is Delach co-tenant — keep $500/mo on 216 Oak Park rent
  2  Vacant October and half of November
  3  Joan paid half of November (cash not on 0203). 12/1 Zelle OUT $675 to Joan
     is a repair → 216 INTERIOR MAINTENANCE. 12/1 $1,950 IN stays Dec rent;
     12/31 $1,950 stays Jan 2026 prepaid cash in TY2025.
  4  Remove GCM from EPGC Consultant (personal 1099 / GCM tab + Income E7 only)
  5  May 22 Gates $609.02 is reimbursement
  6  Bear Creek $38k is gift to Jake from parents; UTMA Thomas $5k each
  7  Freeman's / CIBC $11,000 is art-sale settlement Sale 6428; check to Megan
  8  Megan asked to send Checking 8507 statements

Do not re-run apply_checking_inflows_2025.py (would put GCM back on Consultant)
or apply_cc_answers_2025.py (would double HD). Not tax advice.
"""
from __future__ import annotations

import json
import shutil
import sys
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path("/workspace/.cursor/scratch")
sys.path.insert(0, str(ROOT))
import apply_cc_2025 as base  # noqa: E402
import apply_cc_answers_2025 as ans  # noqa: E402
import apply_checking_inflows_2025 as inf  # noqa: E402

XLSX = ans.XLSX
DELIVERABLE = ans.DELIVERABLE
JSON_PATH = ROOT / "checking_inflows_2025.json"

GREEN = base.GREEN
YELLOW = base.YELLOW
BLUE = base.BLUE
GRAY = base.GRAY
HEADER_FILL = base.HEADER_FILL
HEADER_FONT = base.HEADER_FONT
THIN = base.THIN
ORANGE = ans.ORANGE
MONTHS = base.MONTHS
PROP_COLS = base.PROP_COLS
EPGC_COLS = base.EPGC_COLS
D = base.D
money = base.money
zeros = base.zeros
add = base.add

JOAN_REPAIR = D("675.00")
FREEMANS_SALE = D("11000.00")
NO_FILL = PatternFill(fill_type=None)


def write_month_row(ws, row: int, values, colmap: dict, fill, zero_fill=NO_FILL) -> None:
    """Replace a monthly row and clear leftover fills on $0 months."""
    for i, m in enumerate(MONTHS):
        cell = ws.cell(row, colmap[m])
        cell.value = money(values[i])
        cell.number_format = "0.00"
        cell.fill = fill if values[i] != 0 else zero_fill


ANSWERS_CONFIRMED = [
    (
        "Ava L Hernandez $500/mo rent",
        "CONFIRMED Delach co-tenant. Keep eight $500 Zelle credits on 216 Oak Park STR / RENTAL INCOME. "
        "Delach $1,450 + Ava $500 = $1,950 unit rent (same as Joan later).",
    ),
    (
        "Oak Park vacant Oct + half Nov",
        "CONFIRMED vacant October and the first half of November (held for rent). "
        "Occupancy: Delach/Ava through September; Joan from mid-November. "
        "Cash-basis months stay as deposited (Sep cash is $0 because September rent hit 8/29). "
        "Jan 2025 occupancy cash hit 12/31/2024 — still out of TY2025.",
    ),
    (
        "Joan 12/1 Zelle OUT $675 = repair",
        "CONFIRMED. Checking 8507 Zelle to Joan 12/1 $675 → 216 INTERIOR MAINTENANCE (December). "
        "CPA may reclass interior vs exterior. 12/1 $1,950 IN on joint 0203 stays December occupancy cash; "
        "12/31 $1,950 “January rent” stays Jan 2026 prepaid cash in TY2025.",
    ),
    (
        "GCM off EPGC Consultant",
        "CONFIRMED. BOM PAYMEN $21,500 stays on the GCM tab + Income E7 (personal 1099-NEC, Jacob’s name). "
        "EPGC Consultant zeroed. JAKE C / SEFOF reimbursements $5,214.69 still excluded.",
    ),
    (
        "May 22 Gates $609.02",
        "CONFIRMED reimbursement (JAKE C, not BOM PAYMEN). Excluded from income.",
    ),
    (
        "Bear Creek $38,000 to Jake from parents",
        "CONFIRMED annual gift to Jacob from parents — not income. "
        "User: UTMA was not ~$10k each; Thomas wires $5,000 each (Emma CMA 2T48 + Phoebe CMA 4729) = $10,000 excluded.",
    ),
    (
        "Freeman’s / CIBC $11,000 art-sale settlement",
        "CONFIRMED Sale 6428 Contract 303468. Check went to Megan (joint 0203 2025-12-24). "
        "Booked Art Sales Sale Price $11,000; Cost and object name yellow TBD. Not Mercury, not W-2 payroll.",
    ),
    (
        "Checking 8507 statements",
        "WAIT-MEGAN. Jacob asked Megan to send 2025 NOW 8507 statements.",
    ),
]

ANSWERS_STILL_OPEN = [
    (
        "WAIT MEGAN — Checking 8507 (NOW) 2025 statements",
        "Jacob asked Megan to send them. Needed to (a) book Joan’s half-November rent if the deposit is there, "
        "(b) keep reconciling Airbnb/VRBO vs platform, (c) keep the three SEFOF reimbursements tied. "
        "11/30 Zelle FROM MEGAN $1,950 on 8507 is a Megan transfer, not Joan rent.",
    ),
    (
        "Joan half-November rent $975",
        "User: Joan paid half of November. Not on joint 0203. Nov cash on 216 Oak Park stays $0 until 8507 "
        "(or another account) shows the deposit. Do not invent $975.",
    ),
    (
        "Sale 6428 object / Cost",
        "Freeman’s $11,000 is Art Sales Sale Price (check to Megan 12/24). Cost and object blank yellow — "
        "Income E8 net stays $0 until Cost is filled. 2026-01-08 Hindman SETTLEMENT SALE 6428 $937.83 is 2026, not TY2025.",
    ),
]


def rebuild_gcm_tab(wb) -> None:
    if "GCM" in wb.sheetnames:
        idx = wb.sheetnames.index("GCM")
        del wb["GCM"]
        ws = wb.create_sheet("GCM", idx)
    else:
        ws = wb.create_sheet("GCM")
    ws["A1"] = "GCM / Gates Capital 2025 — personal 1099 (NOT EPGC Consultant) — not tax advice"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:N1")
    ws["A2"] = (
        f"BOM PAYMEN ${inf.GCM_FEE_TOTAL} = Form 1099-NEC $21,500 → this tab + Income E7 only. "
        "User 2026-09-18: remove GCM from the EPGC Consultant line (1099 is in Jacob’s name). "
        f"JAKE C / SEFOF reimbursements ${inf.REIMB_TOTAL} EXCLUDED (May 22 $609.02 CONFIRMED reimbursement)."
    )
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:N2")
    ws.row_dimensions[2].height = 48

    headers = ["LINE", *MONTHS, "YR TOTAL"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(4, i, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN
    ws.cell(5, 1, "Board fee (BOM PAYMEN → personal 1099 / GCM tab)")
    for i, _m in enumerate(MONTHS):
        cell = ws.cell(5, i + 2, money(inf.gcm_fee[i]))
        cell.number_format = "0.00"
        cell.border = THIN
        if inf.gcm_fee[i]:
            cell.fill = GREEN
    tot = ws.cell(5, 14, "=SUM(B5:M5)")
    tot.number_format = "0.00"
    tot.fill = GREEN
    tot.border = THIN

    ws.cell(6, 1, "Reimbursements (EXCLUDED from income)")
    reimb_m = zeros()
    for dt, amt, *_ in inf.GCM_REIMB_9922 + inf.GCM_REIMB_8507:
        add(reimb_m, inf.month_idx(dt), amt)
    for i, _m in enumerate(MONTHS):
        cell = ws.cell(6, i + 2, money(reimb_m[i]))
        cell.number_format = "0.00"
        cell.border = THIN
        cell.fill = GRAY
    cell = ws.cell(6, 14, "=SUM(B6:M6)")
    cell.number_format = "0.00"
    cell.fill = GRAY
    cell.border = THIN

    ws.cell(8, 1, "Transaction register")
    ws.cell(8, 1).font = Font(bold=True)
    hdr = ["Date", "Account", "Amount", "Bucket", "Status", "Memo"]
    for i, h in enumerate(hdr, 1):
        cell = ws.cell(9, i, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN
    rows = [
        ("2025-02-12", "9922", D("4100"), "Personal 1099 / GCM tab", "APPLIED", "BOM PAYMEN — not EPGC Consultant"),
        ("2025-04-18", "9922", D("9200"), "Personal 1099 / GCM tab", "APPLIED", "BOM PAYMEN — not EPGC Consultant"),
        ("2025-09-13", "9922", D("4100"), "Personal 1099 / GCM tab", "APPLIED", "BOM PAYMEN — not EPGC Consultant"),
        ("2025-11-13", "9922", D("4100"), "Personal 1099 / GCM tab", "APPLIED", "BOM PAYMEN — not EPGC Consultant"),
        ("2025-02-18", "9922", D("1150.99"), "Reimbursement", "EXCLUDED", "PMT INFO: BOARD MEETING EXPENSES"),
        ("2025-05-22", "9922", D("609.02"), "Reimbursement", "EXCLUDED", "CONFIRMED reimbursement (JAKE C, not BOM)"),
        ("2025-11-05", "9922", D("1273.82"), "Reimbursement", "EXCLUDED", "Monarch tag Reimburse; statement JAKE C (not BOM)"),
        ("2025-07-25", "8507", D("462.48"), "Reimbursement", "EXCLUDED", "SEFOF WHARTON EXPENSES — WAIT 8507 statements"),
        ("2025-07-31", "8507", D("1254.93"), "Reimbursement", "EXCLUDED", "SEFOF ANNUAL CONFERENCE EXPENSES — WAIT 8507 statements"),
        ("2025-11-05", "8507", D("463.45"), "Reimbursement", "EXCLUDED", "SEFOF EDUCATION COMMITTEE — WAIT 8507 statements"),
    ]
    fills = {"APPLIED": GREEN, "EXCLUDED": GRAY}
    for i, row in enumerate(rows, 10):
        for c, val in enumerate(row, 1):
            cell = ws.cell(i, c, money(val) if isinstance(val, Decimal) else val)
            cell.border = THIN
            cell.fill = fills[row[4]]
            if c == 3:
                cell.number_format = "0.00"
    last = 9 + len(rows)
    ws.cell(
        last + 2,
        1,
        "1099-NEC Gates Capital Management LLC TIN 84-1331261 box 1 = $21,500. "
        "BOM sum ties. Reimbursements are not on the 1099. EPGC Consultant is $0.",
    )
    ws.cell(last + 2, 1).fill = BLUE
    ws.merge_cells(start_row=last + 2, start_column=1, end_row=last + 2, end_column=6)
    widths = [22, 14, 14, 28, 14, 80]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for i in range(7, 15):
        ws.column_dimensions[get_column_letter(i)].width = 12


def rebuild_inflow_ledger(wb) -> None:
    if "INFLOW_LEDGER" in wb.sheetnames:
        del wb["INFLOW_LEDGER"]
    idx = wb.sheetnames.index("ASK") + 1 if "ASK" in wb.sheetnames else 2
    led = wb.create_sheet("INFLOW_LEDGER", idx)
    led["A1"] = "2025 checking inflows after 2026-09-18 answers — JSC 9922 + joint 0203 — not tax advice"
    led["A1"].font = Font(bold=True, size=14)
    led.merge_cells("A1:H1")
    led["A2"] = (
        "Cash-basis by deposit date. GCM BOM → personal 1099 / GCM tab (NOT EPGC Consultant). "
        "JAKE C / SEFOF → reimbursement EXCLUDED (May 22 CONFIRMED). "
        "Delach + Ava (co-tenant CONFIRMED) + Joan → 216 Oak Park STR / RENTAL INCOME. "
        "Bear Creek = gift to Jake from parents. UTMA Thomas $5k each excluded. "
        "Freeman’s Sale 6428 $11,000 → Art Sales (check to Megan). "
        "8507 Zelle OUT Joan $675 → 216 INTERIOR MAINTENANCE."
    )
    led["A2"].alignment = Alignment(wrap_text=True)
    led.merge_cells("A2:H2")
    led.row_dimensions[2].height = 52
    headers = ["Account", "Date", "Counterparty", "Amount", "Bucket", "Tab", "Status", "Note"]
    for i, h in enumerate(headers, 1):
        cell = led.cell(4, i, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN
    status_fill = {
        "APPLIED": GREEN,
        "APPLIED-YELLOW": YELLOW,
        "EXCLUDED": GRAY,
        "HOLD": YELLOW,
        "WAIT-MEGAN": ORANGE,
        "TRANSFER": BLUE,
    }
    rent_notes = {
        ("2025-12-01", "JOAN FRIEDBERGER"): (
            "APPLIED",
            "rent. December occupancy cash. (Half-Nov $975 not on 0203 — WAIT 8507.)",
        ),
        ("2025-12-31", "JOAN FRIEDBERGER"): (
            "APPLIED",
            "January rent. Jan 2026 prepaid cash in TY2025.",
        ),
    }
    rows = []
    for dt, who, amt, memo, _status, note in inf.RENT_TXNS:
        if who == "AVA L HERNANDEZ":
            status, note2 = "APPLIED", f"{memo}. Delach co-tenant CONFIRMED."
        elif (dt, who) in rent_notes:
            status, note2 = rent_notes[(dt, who)]
        else:
            status, note2 = "APPLIED", f"{memo}. LTR 216 Oak Park"
        rows.append(("0203 joint", dt, who, float(amt), "STR / RENTAL INCOME", "216 N. Oak Park Ave", status, note2))
    for dt, amt, acct, memo in [
        ("2025-02-12", D("4100"), "9922 JSC", "BOM PAYMEN — personal 1099, not EPGC"),
        ("2025-04-18", D("9200"), "9922 JSC", "BOM PAYMEN — personal 1099, not EPGC"),
        ("2025-09-13", D("4100"), "9922 JSC", "BOM PAYMEN — personal 1099, not EPGC"),
        ("2025-11-13", D("4100"), "9922 JSC", "BOM PAYMEN — personal 1099, not EPGC"),
    ]:
        rows.append((acct, dt, "Gates Capital", float(amt), "Personal 1099", "GCM tab + Income E7", "APPLIED", memo))
    for dt, amt, acct, memo in inf.GCM_REIMB_9922 + inf.GCM_REIMB_8507:
        label = "9922 JSC" if acct == "9922" else "8507 NOW (WAIT-MEGAN statements)"
        if dt == "2025-05-22":
            memo = "CONFIRMED reimbursement (JAKE C, not BOM)"
        rows.append((label, dt, "Gates Capital", float(amt), "Reimbursement", "—", "EXCLUDED", memo))
    for dt, amt, memo in inf.BEAR_CREEK:
        rows.append(
            (
                "9922 JSC",
                dt,
                "Bear Creek Inc.",
                float(amt),
                "Annual gift to Jake from parents",
                "—",
                "EXCLUDED",
                memo,
            )
        )
    for dt, amt, acct, memo in inf.UTMA_GIFTS:
        rows.append((acct, dt, "THOMAS (wire)", float(amt), "UTMA gift", "—", "EXCLUDED", memo))
    rows.append(
        (
            "0203 joint",
            "2025-12-24",
            "Freeman's LLC / CIBC",
            11000.00,
            "Art-sale settlement Sale 6428",
            "Art Sales",
            "APPLIED",
            "Check to Megan. Sale Price $11,000; Cost / object TBD.",
        )
    )
    rows.append(
        (
            "8507 NOW",
            "2025-12-01",
            "JOAN FRIEDBERGER",
            -675.00,
            "Repair (Zelle OUT)",
            "216 INTERIOR MAINTENANCE",
            "APPLIED",
            "User: towards a repair. Booked December INTERIOR. Interior vs exterior not specified.",
        )
    )
    rows.append(
        (
            "8507 NOW",
            "pending",
            "JOAN FRIEDBERGER",
            975.00,
            "Half-November rent",
            "216 N. Oak Park Ave",
            "WAIT-MEGAN",
            "User: Joan paid half of November. Not on 0203. Do not book until 8507 shows the deposit.",
        )
    )
    for i, row in enumerate(rows, 5):
        for c, val in enumerate(row, 1):
            cell = led.cell(i, c, val)
            cell.border = THIN
            cell.fill = status_fill.get(row[6], BLUE)
            if c == 4:
                cell.number_format = "0.00"
    last = 4 + len(rows)
    led.cell(
        last + 2,
        1,
        "STATUS: APPLIED = in the workbook. EXCLUDED = not income. WAIT-MEGAN = parked until 8507 statements.",
    )
    led.merge_cells(start_row=last + 2, start_column=1, end_row=last + 2, end_column=8)
    widths = [28, 14, 28, 12, 36, 26, 16, 78]
    for i, w in enumerate(widths, 1):
        led.column_dimensions[get_column_letter(i)].width = w
    led.auto_filter.ref = f"A4:H{last}"
    led.freeze_panes = "A5"


def append_ask_inflow_totals(ask: Worksheet) -> None:
    r = ask.max_row + 2
    ask.cell(r, 1, "Checking inflows after answers 2026-09-18").font = Font(bold=True, color="FFFFFF")
    ask.cell(r, 1).fill = HEADER_FILL
    ask.cell(r, 2).fill = HEADER_FILL
    summary = [
        ("216 Oak Park STR / RENTAL INCOME (cash, Ava CONFIRMED co-tenant)", inf.OAK_RENT_TOTAL, GREEN),
        ("  Delach 8 × $1,450", D("11600"), GREEN),
        ("  Ava Hernandez 8 × $500 (co-tenant CONFIRMED)", D("4000"), GREEN),
        ("  Joan Friedberger 2 × $1,950 (Dec cash + Jan 2026 prepaid)", D("3900"), GREEN),
        ("  Joan half-Nov $975 NOT booked (WAIT 8507)", D("975"), ORANGE),
        ("216 INTERIOR MAINTENANCE Dec (Joan Zelle OUT repair)", JOAN_REPAIR, GREEN),
        ("EPGC Consultant (GCM removed)", D("0"), GRAY),
        ("GCM personal 1099-NEC (tab + Income E7)", inf.GCM_FEE_TOTAL, GREEN),
        ("GCM reimbursements EXCLUDED (9922 + 8507; May 22 CONFIRMED)", inf.REIMB_TOTAL, GRAY),
        ("Bear Creek gift to Jake from parents EXCLUDED", inf.BEAR_CREEK_TOTAL, GRAY),
        ("UTMA Thomas wires EXCLUDED (2 × $5,000, not ~$10k each)", inf.UTMA_GIFT_TOTAL, GRAY),
        ("Art Sales Sale 6428 Sale Price (Cost TBD)", FREEMANS_SALE, YELLOW),
    ]
    for i, (lab, val, fill) in enumerate(summary):
        ask.cell(r + 1 + i, 1, lab).fill = fill
        cell = ask.cell(r + 1 + i, 2, money(val))
        cell.number_format = '"$"#,##0.00'
        cell.fill = fill


def apply_art_sales(art) -> None:
    """Replace leftover row-23 placeholder with Freeman's Sale 6428."""
    art["A23"] = "Sale 6428 Contract 303468 (object TBD)"
    art["B23"] = "Cost TBD"
    art["C23"] = None
    art["D23"] = None
    art["E23"] = "Freeman's LLC — check to Megan (joint 0203)"
    art["F23"] = money(FREEMANS_SALE)
    art["G23"] = datetime(2025, 12, 24)
    art["H23"] = '=IF(OR(F23="",C23=""),"",F23-C23)'
    art["I23"] = '=IF(OR(H23="",C23=0),"",H23/C23)'
    art["A23"].fill = YELLOW
    art["B23"].fill = YELLOW
    art["C23"].fill = YELLOW
    art["E23"].fill = GREEN
    art["F23"].fill = GREEN
    art["F23"].number_format = "0.00"
    art["G23"].fill = GREEN
    art["G23"].number_format = "YYYY-MM-DD"
    art["A26"] = (
        "Yellow Sale Price = still need Berk buyer invoices. "
        "Row 23 Freeman’s Sale 6428 $11,000 CONFIRMED (check to Megan 12/24); Cost / object yellow. "
        "Green Cost = COGS already entered. Not tax advice."
    )
    art["A26"].fill = YELLOW
    art.merge_cells("A26:I26")
    art["A35"] = (
        "User 2026-09-18: Freeman’s / CIBC $11,000 on joint 0203 is an art-sale settlement; the check went to Megan. "
        "Do not add Cost into the total until the object is named. 2026-01-08 Hindman Sale 6428 $937.83 is 2026."
    )
    art["A35"].fill = YELLOW
    art["A35"].alignment = Alignment(wrap_text=True)
    art.merge_cells("A35:I35")
    art.row_dimensions[35].height = 36


def main() -> None:
    assert XLSX.exists(), XLSX
    wb = ans.load_workbook(XLSX)

    oak = wb["216 N. Oak Park Ave"]
    oak_map = ans.label_map(oak, 2)
    write_month_row(oak, oak_map["STR / RENTAL INCOME"], inf.oak_rent, PROP_COLS, GREEN)
    interior = zeros()
    add(interior, 11, JOAN_REPAIR)
    write_month_row(oak, oak_map["INTERIOR MAINTENANCE"], interior, PROP_COLS, GREEN)
    oak["B30"] = (
        "LTR after answers 2026-09-18: Ava CONFIRMED Delach co-tenant — keep $500/mo. "
        f"TY2025 cash ${inf.OAK_RENT_TOTAL} (Delach $11,600 + Ava $4,000 + Joan $3,900). "
        "Occupancy vacant October and half of November (held for rent). "
        "Joan half-Nov $975 NOT booked — not on 0203, WAIT 8507. "
        "12/1 $1,950 IN = December occupancy cash; 12/31 $1,950 = Jan 2026 prepaid cash in TY2025. "
        "12/1 8507 Zelle OUT $675 to Joan → INTERIOR MAINTENANCE (repair). "
        "Jan 2025 occupancy was paid 12/31/2024 — not in this year. "
        "HOA + Rocket payment cash from Monarch (green). Deduct interest from Rocket 1098 only — not full P+I. "
        "Not tax advice."
    )
    oak["B30"].alignment = Alignment(wrap_text=True)
    oak["B30"].fill = YELLOW
    oak.row_dimensions[30].height = 108

    epgc = wb["EPGC LLC"]
    epgc_map = {}
    for r in range(1, 40):
        v = epgc.cell(r, 1).value
        if v:
            epgc_map[str(v).strip()] = r
    write_month_row(epgc, epgc_map["Consultant"], zeros(), EPGC_COLS, GREEN, zero_fill=GRAY)
    note_r = 43
    epgc.cell(
        note_r,
        1,
        (
            "Checking answers 2026-09-18: GCM BOM PAYMEN $21,500 REMOVED from EPGC Consultant "
            "(personal 1099 / GCM tab + Income E7 only). Consultant row is $0. "
            f"JAKE C / SEFOF reimbursements ${inf.REIMB_TOTAL} still EXCLUDED. "
            "Bear Creek / UTMA not on this tab. Not tax advice."
        ),
    )
    epgc.cell(note_r, 1).alignment = Alignment(wrap_text=True)
    epgc.cell(note_r, 1).fill = GRAY
    epgc.merge_cells(start_row=note_r, start_column=1, end_row=note_r + 1, end_column=14)

    inc = wb["Income"]
    inc["E5"] = money(inf.OAK_RENT_TOTAL)
    inc["E5"].number_format = "0.00"
    inc["E5"].fill = GREEN
    inc["E7"] = money(inf.GCM_FEE_TOTAL)
    inc["E7"].number_format = "0.00"
    inc["E7"].fill = GREEN
    inc["E8"] = 0
    inc["E8"].number_format = "0.00"
    inc["E8"].fill = YELLOW
    inc["A12"] = (
        f"Notes: E4 Monarch paycheck cash ≠ W-2 Box 1. E5 216 Oak Park LTR cash ${inf.OAK_RENT_TOTAL} "
        "(Delach + Ava co-tenant CONFIRMED + Joan; vacant Oct + half Nov; half-Nov $975 WAIT 8507). "
        "E6 PLATFORM STR net LOCKED. E7 GCM 1099-NEC $21,500 on GCM tab only — NOT EPGC Consultant "
        "(reimbursements excluded; May 22 CONFIRMED). "
        "E8 Art net $0 until Sale 6428 Cost is filled (Sale Price $11,000 on Art Sales, check to Megan). "
        "Not tax advice."
    )
    inc["A12"].alignment = Alignment(wrap_text=True)
    inc.merge_cells("A12:E13")
    inc.row_dimensions[12].height = 64

    apply_art_sales(wb["Art Sales and Purchases"])
    rebuild_gcm_tab(wb)

    src = wb["_SOURCE_2025"]
    src["A22"] = "Checking inflows 2026-09-18"
    src["B22"] = (
        "JSC 9922 + joint 0203 full-year eStmts. "
        f"GCM BOM ${inf.GCM_FEE_TOTAL} → GCM tab + Income E7 only (NOT EPGC Consultant). "
        f"Reimburse ${inf.REIMB_TOTAL} excluded (May 22 CONFIRMED). "
        f"216 Oak Park rent cash ${inf.OAK_RENT_TOTAL} (Ava co-tenant CONFIRMED). "
        "Vacant Oct + half Nov. Joan repair $675 → INTERIOR Dec. "
        f"Bear Creek ${inf.BEAR_CREEK_TOTAL} gift to Jake from parents. "
        f"UTMA Thomas ${inf.UTMA_GIFT_TOTAL} not income. "
        "Freeman’s Sale 6428 $11,000 Art Sales (check to Megan). 8507 WAIT-MEGAN. Not tax advice."
    )
    src["A22"].fill = GREEN
    src["B22"].fill = GREEN
    src["B22"].alignment = Alignment(wrap_text=True)
    src["A23"] = "Checking answers 2026-09-18"
    src["B23"] = (
        "Q1 Ava co-tenant keep. Q2 vacant Oct + half Nov. Q3 Joan half-Nov WAIT 8507; "
        "12/1 Zelle OUT $675 repair. Q4 GCM off EPGC Consultant. Q5 May 22 reimbursement. "
        "Q6 Bear Creek $38k to Jake from parents. Q7 Freeman’s $11k art settlement to Megan. "
        "Q8 Megan sending 8507 statements."
    )
    src["A23"].fill = GREEN
    src["B23"].fill = GREEN
    src["B23"].alignment = Alignment(wrap_text=True)

    comed_open = []
    for item, detail in ans.STILL_OPEN:
        if item.startswith("REMIND JACOB"):
            detail = (
                "Joint 0203 statements show ComEd ACH (e.g. 1/21 Jacob $59.04 + Megan $132.94). "
                + detail
            )
        comed_open.append((item, detail))
    ans.CONFIRMED.extend(ANSWERS_CONFIRMED)
    ans.STILL_OPEN[:] = comed_open + ANSWERS_STILL_OPEN
    ans.rebuild_ask(wb)
    append_ask_inflow_totals(wb["ASK"])
    rebuild_inflow_ledger(wb)

    wb.save(XLSX)
    shutil.copy2(XLSX, DELIVERABLE)

    export_names = [
        ("EPGC LLC", "EPGC_LLC.csv"),
        ("216 N. Oak Park Ave", "216_Oak_Park.csv"),
        ("GCM", "GCM.csv"),
        ("Income", "Income.csv"),
        ("INFLOW_LEDGER", "INFLOW_LEDGER.csv"),
        ("ASK", "ASK.csv"),
        ("_SOURCE_2025", "SOURCE_2025.csv"),
        ("Art Sales and Purchases", "Art_Sales.csv"),
        ("524 Ferdinand Ave, Unit 2", "524_Unit_2.csv"),
        ("524 Ferdinand Ave, Unit 1", "524_Unit_1.csv"),
    ]
    for name, fname in export_names:
        ans.sheet_to_csv(wb[name], ans.CSV_DIR / fname)
        safe = name.replace(",", "").replace(".", "")[:40]
        ans.save_sheet_part(wb[name], ans.PARTS / f"{safe}.xlsx")

    payload = {
        "updated": "2026-09-18",
        "disclaimer": "Not tax advice. Organizational packet only.",
        "pass": "checking-answers",
        "accounts": inf.JSON_PATH.exists()
        and json.loads(inf.JSON_PATH.read_text(encoding="utf-8")).get("accounts")
        or {},
        "gcm_board_1099_personal": float(inf.GCM_FEE_TOTAL),
        "epgc_consultant": 0.0,
        "gcm_reimburse_excluded": float(inf.REIMB_TOTAL),
        "oak_park_ltr_cash": float(inf.OAK_RENT_TOTAL),
        "oak_park_by_month": [float(x) for x in inf.oak_rent],
        "oak_interior_repair_dec": float(JOAN_REPAIR),
        "joan_half_nov_unbooked": 975.0,
        "vacancy": "October + first half of November",
        "bear_creek_gift_to_jake_excluded": float(inf.BEAR_CREEK_TOTAL),
        "utma_thomas_excluded": float(inf.UTMA_GIFT_TOTAL),
        "art_sale_6428_price": float(FREEMANS_SALE),
        "art_sale_6428_cost": None,
        "checking_8507": "WAIT-MEGAN",
        "confirmed": [a[0] for a in ANSWERS_CONFIRMED],
        "still_open": [a[0] for a in ANSWERS_STILL_OPEN],
    }
    JSON_PATH.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    pkt = json.loads(ans.JSON_PATH.read_text(encoding="utf-8")) if ans.JSON_PATH.exists() else {}
    pkt["checking_inflows"] = payload
    ans.JSON_PATH.write_text(json.dumps(pkt, indent=2), encoding="utf-8")

    print("Oak Park rent", inf.OAK_RENT_TOTAL)
    print("Joan repair Dec", JOAN_REPAIR)
    print("EPGC Consultant", 0)
    print("GCM 1099 personal", inf.GCM_FEE_TOTAL)
    print("Freeman Sale Price", FREEMANS_SALE)
    print("saved", XLSX, "bytes", XLSX.stat().st_size)


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""Ultra-sparse OOXML of last year's 11 tabs + 2025.

xlsx cell refs (AO5) survive Google conversion; SpreadsheetML ss:Index does not.
Keep the file small enough for Drive MCP base64 (~<20KB zip).

A1 formulas (not R1C1). No theme, no external links. Not tax advice.
"""
from __future__ import annotations

import html
import io
import sys
import zipfile
from pathlib import Path

from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from export_ssml_personal_income_2025 import (
    LAST_YEAR_TABS,
    PROPERTY_TABS,
    YEAR_TOTAL_COLS,
    compute_formulas,
    keep_formula,
    keep_row,
    keep_value,
    max_col_for,
    merge_map,
    rgb,
    sheets_formula,
    PEACH,
    GREEN,
    YELLOW,
)

ROOT = Path("/workspace/.cursor/scratch")
SRC = ROOT / "tax_turbo_parts" / "Personal Income 2025.xlsx"
OUT = ROOT / "tax_turbo_parts" / "Personal Income 2025.tiny.xlsx"

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
WS_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"
SS_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings"
ST_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles"


def num_str(v: float) -> str:
    if abs(v - round(v)) < 1e-9:
        return str(int(round(v)))
    return f"{v:.12g}"


def xf_of(cell, write_f: bool) -> int:
    bg = None
    if cell.fill and cell.fill.fill_type not in (None, "none"):
        bg = rgb(getattr(cell.fill, "fgColor", None))
    if bg == YELLOW:
        return 3
    if bg == GREEN:
        return 2
    if cell.font and cell.font.bold:
        return 1
    if write_f:
        return 4
    return 0


def write_book(src: Path, dest: Path) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(src, data_only=False)
    strings: list[str] = []
    sindex: dict[str, int] = {}

    def sid(text: str) -> int:
        if text not in sindex:
            sindex[text] = len(strings)
            strings.append(text)
        return sindex[text]

    sheets_xml: list[str] = []
    for name in LAST_YEAR_TABS:
        ws = wb[name]
        cached = compute_formulas(ws)
        merges = merge_map(ws)
        rows_xml: list[str] = []
        max_r = ws.max_row or 1
        max_c = max_col_for(name, ws)
        for r in range(1, max_r + 1):
            if not keep_row(name, r, ws):
                continue
            cells_xml: list[str] = []
            for c in range(1, max_c + 1):
                cell = ws.cell(r, c)
                is_f = cell.data_type == "f" and isinstance(cell.value, str)
                raw_f = ""
                if is_f:
                    raw_f = sheets_formula(cell.value) or ""
                    if "[" in (cell.value or ""):
                        raw_f = ""
                        is_f = False
                write_f = bool(is_f and raw_f and keep_formula(ws, r, c, raw_f[1:] if raw_f else ""))
                val = cell.value
                if name == "Income" and r == 1 and isinstance(val, (int, float)) and not isinstance(val, bool):
                    val = str(int(val))
                if is_f and not write_f:
                    val = cached.get((r, c), 0.0)
                    if not keep_value(val, cell, ws) and (r, c) not in merges:
                        continue
                elif not is_f:
                    if not keep_value(val, cell, ws) and (r, c) not in merges:
                        continue
                if (
                    isinstance(val, (int, float))
                    and not isinstance(val, bool)
                    and val == 0
                    and (r, c) not in merges
                    and c not in YEAR_TOTAL_COLS
                    and name != "Income"
                ):
                    if not write_f:
                        continue
                ref = f"{get_column_letter(c)}{r}"
                xf = xf_of(cell, write_f)
                s_attr = f' s="{xf}"' if xf else ""
                if write_f:
                    fml = raw_f[1:] if raw_f.startswith("=") else raw_f
                    cached_v = cached.get((r, c), 0.0)
                    cells_xml.append(
                        f'<c r="{ref}"{s_attr}><f>{html.escape(fml)}</f><v>{num_str(cached_v)}</v></c>'
                    )
                elif isinstance(val, (int, float)) and not isinstance(val, bool):
                    cells_xml.append(f'<c r="{ref}"{s_attr}><v>{num_str(float(val))}</v></c>')
                else:
                    text = str(val) if val is not None else ""
                    if r == 2 and c in YEAR_TOTAL_COLS and name in PROPERTY_TABS and "YR TOTAL" in text:
                        if c == 28 and name == "524 Ferdinand Ave, Unit 1":
                            text = "2025 YR TOTAL"
                        elif c == 15:
                            text = "2023 YR TOTAL"
                        elif c == 28:
                            text = "2024 YR TOTAL"
                        elif c == 41:
                            text = "2025 YR TOTAL"
                    if len(text) > 36:
                        text = text[:33] + "..."
                    if not text.strip() and (r, c) not in merges:
                        continue
                    cells_xml.append(f'<c r="{ref}"{s_attr} t="s"><v>{sid(text)}</v></c>')
            if cells_xml:
                rows_xml.append(f'<row r="{r}">{"".join(cells_xml)}</row>')
        kept_merges = []
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= (ws.max_row or 1) and keep_row(name, rng.min_row, ws):
                if rng.max_col <= max_c:
                    kept_merges.append(str(rng))
        merge_xml = ""
        if kept_merges:
            parts = "".join(f'<mergeCell ref="{m}"/>' for m in kept_merges)
            merge_xml = f'<mergeCells count="{len(kept_merges)}">{parts}</mergeCells>'
        sheets_xml.append(
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<worksheet xmlns="{NS}">'
            f'<sheetData>{"".join(rows_xml)}</sheetData>{merge_xml}</worksheet>'
        )

    styles = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<styleSheet xmlns="{NS}">'
        '<numFmts count="1"><numFmt numFmtId="164" formatCode="$#,##0.00"/></numFmts>'
        '<fonts count="2">'
        '<font><sz val="10"/><name val="Century Gothic"/></font>'
        '<font><b/><sz val="11"/><name val="Century Gothic"/></font>'
        "</fonts>"
        '<fills count="4">'
        '<fill><patternFill patternType="none"/></fill>'
        '<fill><patternFill patternType="gray125"/></fill>'
        f'<fill><patternFill patternType="solid"><fgColor rgb="FF{GREEN}"/></patternFill></fill>'
        f'<fill><patternFill patternType="solid"><fgColor rgb="FF{YELLOW}"/></patternFill></fill>'
        "</fills>"
        '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
        '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
        '<cellXfs count="5">'
        '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
        '<xf numFmtId="0" fontId="1" fillId="0" borderId="0" xfId="0" applyFont="1"/>'
        '<xf numFmtId="164" fontId="0" fillId="2" borderId="0" xfId="0" applyFill="1" applyNumberFormat="1"/>'
        '<xf numFmtId="164" fontId="0" fillId="3" borderId="0" xfId="0" applyFill="1" applyNumberFormat="1"/>'
        '<xf numFmtId="164" fontId="1" fillId="0" borderId="0" xfId="0" applyFont="1" applyNumberFormat="1"/>'
        "</cellXfs>"
        '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>'
        "</styleSheet>"
    )
    sst_items = "".join(
        f'<si><t xml:space="preserve">{html.escape(t)}</t></si>' for t in strings
    )
    sst = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<sst xmlns="{NS}" count="{len(strings)}" uniqueCount="{len(strings)}">{sst_items}</sst>'
    )

    sheet_entries, rels, overrides = [], [], []
    for i, name in enumerate(LAST_YEAR_TABS, start=1):
        rid = f"rId{i}"
        sheet_entries.append(f'<sheet name="{html.escape(name)}" sheetId="{i}" r:id="{rid}"/>')
        rels.append(f'<Relationship Id="{rid}" Type="{WS_TYPE}" Target="worksheets/sheet{i}.xml"/>')
        overrides.append(
            f'<Override PartName="/xl/worksheets/sheet{i}.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        )
    n = len(LAST_YEAR_TABS)
    rels.append(f'<Relationship Id="rId{n+1}" Type="{ST_TYPE}" Target="styles.xml"/>')
    rels.append(f'<Relationship Id="rId{n+2}" Type="{SS_TYPE}" Target="sharedStrings.xml"/>')
    workbook = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<workbook xmlns="{NS}" xmlns:r="{REL_NS}">'
        f'<sheets>{"".join(sheet_entries)}</sheets></workbook>'
    )
    workbook_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<Relationships xmlns="{PKG_REL}">{"".join(rels)}</Relationships>'
    )
    root_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<Relationships xmlns="{PKG_REL}">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
        "</Relationships>"
    )
    ct = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
        f'{"".join(overrides)}</Types>'
    )

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", root_rels)
        z.writestr("xl/workbook.xml", workbook)
        z.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        z.writestr("xl/styles.xml", styles)
        z.writestr("xl/sharedStrings.xml", sst)
        for i, xml in enumerate(sheets_xml, start=1):
            z.writestr(f"xl/worksheets/sheet{i}.xml", xml)
    dest.write_bytes(buf.getvalue())
    print(f"tiny ooxml {dest} {dest.stat().st_size} bytes strings={len(strings)}")


if __name__ == "__main__":
    write_book(SRC, OUT)
    import base64
    b64 = base64.b64encode(OUT.read_bytes())
    print("base64", len(b64))
    from openpyxl import load_workbook
    wb = load_workbook(OUT, data_only=False)
    print("tabs", wb.sheetnames)
    inc = wb["Income"]
    print("I10", inc["I10"].value, inc["I10"].data_type, "I4", inc["I4"].value)
    oak = wb["216 N. Oak Park Ave"]
    print("merges", list(oak.merged_cells.ranges))
    print("O2", oak["O2"].value, "AB2", oak["AB2"].value, "AO2", oak["AO2"].value)
    print("AO5", oak["AO5"].value, oak["AO5"].data_type, "AC5", oak["AC5"].value, "C5", oak["C5"].value)
    print("B14", oak["B14"].value, "AC14", oak["AC14"].value, "AO14", oak["AO14"].value)

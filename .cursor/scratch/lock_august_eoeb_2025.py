#!/usr/bin/env python3
"""Lock August Fortuna joint SALE + visible per-invoice EOEB Deal Classifier.

Assumes Belt (A78) and Venus (A79) are already locked by lock_belt_venus_2025.py.
Does NOT revert those rows. Does NOT dump EOEB remainder onto Consultant or Art Sales.

User 2026-09-21:
  August Fortuna is another joint investment with Erdal. He sold; Jacob got
  capital + earnings. Fortuna 8/15 (native 8/18) −$20,000 unnamed object.
  Working settlement: Erdal 10/9 +$25,000 = $20k capital + $5k earnings.
  EPGC Art Sales October. 7/24 cannot be August (too early — that is Venus).
  11/7 $6,000 FORTUNA PAYMENT still ASK. Not consultant. Not unsold inventory.

  EOEB: lock mosaics 2/21 $105,000 SALE and 12/2 $565.75 Wise reimburse
  (not income). Remainder $130,270 per invoice: proposed FEE (no matching
  dealer Cost on 8291; Fortuna joints are NOT EOEB Cost) but lock
  UNALLOCATED. 1/16 Canosan-horse OPEN. 5/1 $21,000 is NOT Venus.

Not tax advice. Organizational packet only.
"""
from __future__ import annotations

import csv
import json
import shutil
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import sys

sys.path.insert(0, str(Path("/workspace/.cursor/scratch")))
from sanitize_xlsx_for_sheets import sanitize_xlsx

ROOT = Path("/workspace/.cursor/scratch")
PARTS = ROOT / "tax_turbo_parts"
CSV_DIR = ROOT / "cc_fill_csv"
XLSX = PARTS / "Personal Income 2025.xlsx"
DELIVERABLE = ROOT / "tax_turbo_deliverable" / "Personal Income 2025.xlsx"
PACKET = ROOT / "cpa_packet_v1.json"
MERCURY_JSON = ROOT / "mercury_2025.json"

GREEN = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
PEACH = PatternFill("solid", fgColor="F7CAAC")
BLUE = PatternFill("solid", fgColor="DDEBF7")
CG = Font(name="Century Gothic", size=10)
CG_B = Font(name="Century Gothic", size=10, bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
ACCT = '_("$"* #,##0.00_);_("$"* \\(#,##0.00\\);_("$"* "-"??_);_(@_)'

I8 = 39055.06  # 34055.06 + August joint net 5000
EPGC_JULY_GROSS = 76000.0
EPGC_OCT_GROSS = 25000.0
ART_GROSS = 251000.0  # 226000 + 25000
L5_REMAINING = 37396.32
ERDAL_REMAINING = 6000.0  # only 11/7 FORTUNA PAYMENT
EOEB_REMAINING = 130270.0
I10_EXPECTED = 70618 + 19500 + 17086.94 + 21500 + I8  # 167760.00

AUG_NOTE = (
    "LOCKED 2026-09-21 joint SALE (EPGC share). Another joint investment with "
    "Erdal (user 2026-09-21). He sold; Jacob got capital + earnings. EPGC Cost: "
    "Fortuna 8/15 −$20,000 (Mercury native 8/18, same cash; object unnamed). "
    "Working remittance: Erdal Dere IN 10/9 +$25,000 = $20,000 capital + $5,000 "
    "earnings. Sale Price on this register is EPGC proceeds ($25,000), not the "
    "object’s full hammer. Hits EPGC Art Sales October. 7/24 is Venus, not August. "
    "11/7 $6,000 FORTUNA PAYMENT still ASK. Classifier: SALE (joint, EPGC share). "
    "Not consultant. Not unsold inventory. Not tax advice."
)
ART_FOOTER = (
    "Mercury 8291 MATCHED/LOCKED: Berk $14,000 (1/10 seals) + $30,000 (2/7 lots). "
    "EOEB $105,000 (2/21 mosaics). Fortuna $13,000 (1/8 seals Cost). "
    "Plutus $50,000 of mosaics $90,000 Cost — $40,000 of Cost not on 8291 (ASK). "
    "Roman Gold Belt SALE to L5: Cost $45,000 / Sale $50,000 / net $5,000. "
    "Venus joint SALE EPGC share: Cost $20,000 / proceeds $26,000 / net $6,000. "
    "August Fortuna joint SALE EPGC share: Cost $20,000 / proceeds $25,000 / net $5,000 "
    "(Sale Price = remittance, not full hammer; object unnamed). "
    "Newstar $23,581 jewelry fabrication COGS parked below. "
    "David Aaron $13,595 LOCKED consultant (EPGC Consultant June — not a sale). "
    "Aquinas Hobor $1,000 LOCKED book sale (Feb cash; Cost TBD, not in I8). "
    "Do not dump EOEB remainder $130,270 (see Deal Classifier per invoice) or other "
    "L5 $37,396.32. Not tax advice."
)
INCOME_NOTE = (
    "2025 Actual: I4 Monarch paycheck cash ≠ W-2 Box 1. I5 216 LTR cash $19,500. "
    "I6 524 #2 STR platform net LOCKED. I7 GCM 1099-NEC $21,500. "
    "I8 Art net $39,055.06 = prior MATCHED nets $23,055.06 (Berk lots + seals "
    "$1,000 + mosaics $15,000) + Belt net $5,000 + Venus EPGC-share net $6,000 "
    "+ August Fortuna joint net $5,000. "
    "Excludes Sale 6428 $11,000 until Cost. Excludes Aquinas books $1,000 until Cost. "
    "David Aaron $13,595 is EPGC Consultant June (not I8). "
    "EOEB remainder $130,270 stays UNALLOCATED (per-invoice Deal Classifier; proposed "
    "FEE, no matching Cost — not dumped). Other L5 $37,396.32 / Erdal 11/7 $6,000 "
    "/ Aysel $50,000 still ASK. Not tax advice."
)
EPGC_NOTE = (
    "2025 Mercury 2026-09-21: Art Sales gross cash by month = Jan $14,000 + Feb "
    "$136,000 MATCHED + July $76,000 (Belt L5 $50,000 + Venus remittance $26,000) "
    "+ October $25,000 (August Fortuna joint remittance) = $251,000. "
    "Consultant June $13,595 = David Aaron Limited LOCKED (not a sale). "
    "Consultant Fees July $334.17 + October $658.63 = Wise expertise write-ups "
    "LOCKED ($992.80). EOEB remainder $130,270 is per-invoice UNALLOCATED on the "
    "Deal Classifier (proposed FEE, no matching dealer Cost; Fortuna joints are "
    "NOT EOEB Cost) — not dumped onto Consultant or Art Sales. Other L5 $37,396.32 "
    "MIXED — not dumped. Erdal 11/7 $6,000 FORTUNA PAYMENT still ASK. Aysel Monarch "
    "+$50k vs native Failed −$50k still ASK (not the belt). Koziol/Ariadne $150,000 "
    "LOCKED pass-through — not P&L. Coinbase is Investments. Newstar $23,581 is Art "
    "Sales COGS. Dec Wise $565.75 reimbursed. Aquinas books Cost TBD — cash on Art "
    "Sales, not in I8. Joint Sale Price = EPGC proceeds, not full hammer. GCM 1099 "
    "is personal. Not tax advice."
)

FEE_HYPOTHESIS = (
    "Working hypothesis: leftover EOEB after mosaics has no matching dealer Cost "
    "on Mercury 8291 (Fortuna Venus/Belt/August joints are NOT EOEB Cost) and looks "
    "like a sourcing/consultant FEE. Do NOT silently post to EPGC Consultant or Art "
    "Sales. Lock FEE only with strong per-invoice evidence; otherwise UNALLOCATED. "
    "Not tax advice."
)


def paint(cell, value, fill=GREEN, num=None) -> None:
    cell.value = value
    cell.fill = fill
    cell.font = CG
    if num:
        cell.number_format = num


def patch_art(ws) -> None:
    ws["A1"] = "Art Sales and Purchases"
    ws["A1"].font = CG_B
    for header_row in (3, 16, 39, 50):
        if str(ws.cell(header_row, 1).value or "") == "Object":
            ws.cell(header_row, 11, "Classifier").font = CG_B

    if str(ws["A78"].value or "") != "Roman Gold Belt":
        raise SystemExit("LOCK FAIL: Belt missing at A78 — do not run this before lock_belt_venus_2025.py")
    if "Venus" not in str(ws["A79"].value or ""):
        raise SystemExit("LOCK FAIL: Venus missing at A79 — do not revert Belt/Venus")

    a80 = str(ws["A80"].value or "")
    if a80 == "Total" or a80 == "":
        to_unmerge = [str(rng) for rng in ws.merged_cells.ranges if rng.min_row in (80, 81)]
        for rng in to_unmerge:
            try:
                ws.unmerge_cells(rng)
            except Exception:
                pass
        if a80 == "Total":
            ws.insert_rows(80, 1)

    paint(ws["A80"], "August Fortuna joint — unnamed object (EPGC share)")
    paint(ws["B80"], "Fortuna / Erdal Dere")
    paint(ws["C80"], 20000, num=ACCT)
    paint(ws["D80"], datetime(2025, 8, 15))
    ws["D80"].number_format = "YYYY-MM-DD"
    paint(ws["E80"], "Erdal Dere")
    paint(ws["F80"], 25000, num=ACCT)
    paint(ws["G80"], datetime(2025, 10, 9))
    ws["G80"].number_format = "YYYY-MM-DD"
    paint(ws["H80"], "=F80-C80")
    ws["H80"].number_format = ACCT
    paint(ws["I80"], AUG_NOTE)
    ws["I80"].alignment = WRAP
    paint(ws["K80"], "SALE (joint, EPGC share)")
    ws.row_dimensions[80].height = 64
    for col in range(1, 12):
        ws.cell(80, col).font = CG
        if ws.cell(80, col).fill.fill_type in (None, "none"):
            ws.cell(80, col).fill = GREEN

    total_row = None
    for r in range(80, 92):
        if str(ws.cell(r, 1).value or "") == "Total":
            total_row = r
            break
    if total_row is None:
        total_row = 81
        ws.cell(81, 1, "Total")
    paint(ws.cell(total_row, 1), "Total", fill=PEACH)
    ws.cell(total_row, 8).value = "=SUM(H51:H75)+H78+H79+H80"
    ws.cell(total_row, 8).number_format = ACCT
    ws.cell(total_row, 8).fill = PEACH
    ws.cell(total_row, 8).font = CG_B

    note_row = total_row + 1
    for rng in [str(r) for r in ws.merged_cells.ranges if r.min_row == note_row]:
        try:
            ws.unmerge_cells(rng)
        except Exception:
            pass
    ws.cell(note_row, 1, ART_FOOTER)
    ws.cell(note_row, 1).fill = GREEN
    ws.cell(note_row, 1).font = CG
    ws.cell(note_row, 1).alignment = WRAP
    try:
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=9)
    except Exception:
        pass
    ws.row_dimensions[note_row].height = 72

    for r in range(1, (ws.max_row or 1) + 1):
        a = str(ws.cell(r, 1).value or "")
        if a.startswith("2025 Mercury Fortuna"):
            paint(
                ws.cell(r, 1),
                "2025 Mercury Fortuna — Venus + Belt + August joint SOLD 2025 "
                "(see Sales 2025 rows 78–80); no unnamed Fortuna Cost left on 8291",
                fill=GREEN,
            )
            ws.cell(r, 1).font = CG_B
            ws.cell(r, 1).alignment = WRAP
        if "object still unnamed" in a or a.startswith("Fortuna / Erdal Dere — object still unnamed"):
            paint(ws.cell(r, 1), "SOLD 2025 — August Fortuna joint (see Sales 2025 row 80). Not unsold inventory.")
            paint(ws.cell(r, 3), None)
            paint(
                ws.cell(r, 9),
                "Moved to sold lots. Erdal 10/9 $25,000 / Cost $20,000 / net $5,000. "
                "Object unnamed. 11/7 $6,000 still ASK.",
            )
        if a.startswith("Fortuna unnamed remaining 2025"):
            paint(
                ws.cell(r, 1),
                "Fortuna unnamed remaining 2025 (excl. seals $13,000 sold + Venus $20,000 "
                "sold + Belt $45,000 sold + August joint $20,000 sold)",
                fill=GREEN,
            )
            paint(ws.cell(r, 3), 0, fill=GREEN, num=ACCT)
        if a == "Newstar 2025 total":
            ws.cell(r, 3).value = f"=SUM(C{r-3}:C{r-1})"
            ws.cell(r, 3).number_format = ACCT
            ws.cell(r, 3).fill = GREEN


def patch_income(ws) -> None:
    ws["I8"] = I8
    ws["I8"].number_format = ACCT
    ws["I8"].fill = GREEN
    ws["I8"].font = CG
    ws["A12"] = INCOME_NOTE
    ws["A12"].alignment = WRAP
    ws["A12"].font = CG


def patch_epgc(ws) -> None:
    ws["H53"] = EPGC_JULY_GROSS
    ws["H53"].number_format = ACCT
    ws["H53"].fill = GREEN
    ws["H53"].font = CG
    ws["K53"] = EPGC_OCT_GROSS
    ws["K53"].number_format = ACCT
    ws["K53"].fill = GREEN
    ws["K53"].font = CG
    ws["B53"].fill = PEACH
    ws["C53"].fill = PEACH
    note_row = None
    for r in range(70, 82):
        v = str(ws.cell(r, 1).value or "")
        if v.startswith("2025 Mercury"):
            note_row = r
            break
    if note_row is None:
        note_row = 75
    ws.cell(note_row, 1, EPGC_NOTE)
    ws.cell(note_row, 1).alignment = WRAP
    ws.cell(note_row, 1).font = CG
    ws.row_dimensions[note_row].height = 84


def patch_data_sources(ws) -> None:
    for r in range(1, (ws.max_row or 1) + 1):
        a = str(ws.cell(r, 1).value or "")
        b = str(ws.cell(r, 2).value or "")
        if a == "Income" and "I8" in b:
            ws.cell(r, 3, "39,055.06")
            ws.cell(
                r,
                4,
                "Prior MATCHED $23,055.06 + Belt net $5,000 + Venus EPGC-share net $6,000 "
                "+ August Fortuna joint net $5,000. Sale 6428 and Aquinas wait on Cost.",
            )
            ws.cell(r, 5, "LOCKED green")
            for c in range(1, 6):
                ws.cell(r, c).fill = GREEN
        if a == "EPGC LLC" and b == "Art Sales":
            ws.cell(r, 3, "14,000 Jan + 136,000 Feb + 76,000 July + 25,000 Oct = 251,000")
            ws.cell(
                r,
                4,
                "Gross sale cash. July = Belt L5 $50,000 + Venus $26,000. "
                "October = August Fortuna joint remittance $25,000. "
                "Do not dump EOEB $130,270 or other L5 $37,396.32.",
            )
            ws.cell(r, 5, "LOCKED")
            for c in range(1, 6):
                ws.cell(r, c).fill = GREEN
        if a == "Art Sales tab":
            ws.cell(r, 2, "Seals / lots / mosaics / Belt / Venus / August joint")
            ws.cell(r, 3, "rows 70–80")
            ws.cell(
                r,
                4,
                "2022–2024 lots stay. 2025 adds Belt SALE, Venus joint SALE, August Fortuna joint SALE.",
            )
            ws.cell(r, 5, "LOCKED")
        if a.startswith("Dump EOEB"):
            ws.cell(r, 1, "Dump EOEB $130,270 or other L5 $37,396.32 onto P&L")
            ws.cell(
                r,
                2,
                "EOEB remainder is per-invoice UNALLOCATED (proposed FEE, no matching Cost). "
                "Not dumped. Belt $50k SALE; remaining L5 not attached.",
            )
        if a.startswith("EPGC Art Sales vs"):
            ws.cell(r, 1, "EPGC Art Sales vs $251,000")
            ws.cell(r, 3, 251000)

    # Replace any prior conceptual classifier block, then append.
    start = None
    for r in range(1, (ws.max_row or 1) + 1):
        if str(ws.cell(r, 1).value or "").startswith("DEAL CLASSIFIER"):
            start = r
            break
    if start is None:
        start = (ws.max_row or 1) + 2
    ws.cell(start, 1, "DEAL CLASSIFIER (conceptual)").font = CG_B
    ws.cell(start, 1).fill = BLUE
    headers = ["Classifier", "Meaning", "2025 locked example", "Still ASK / unallocated"]
    for i, h in enumerate(headers, 1):
        ws.cell(start + 1, i, h).font = CG_B
        ws.cell(start + 1, i).fill = BLUE
    rows = [
        ("SALE", "Art object sold; Art Sales register + EPGC Art Sales gross cash; I8 gets NET if Cost known", "Roman Gold Belt L5 $50,000 / Cost $45,000", "Remaining L5 $37,396.32"),
        ("SALE (joint, EPGC share)", "Joint purchase; register Sale Price = EPGC proceeds, not full hammer", "Venus $26,000 / Cost $20,000; August unnamed $25,000 / Cost $20,000", "11/7 $6,000 FORTUNA PAYMENT"),
        ("FEE", "Advisory / consultant income or expertise expense — never I8", "David Aaron June $13,595 Consultant; Wise write-ups $992.80 Fees", "EOEB remainder proposed FEE but not locked"),
        ("PASS", "In and out, not P&L", "Koziol $150,000 10/20 → Ariadne $150,000 10/21", ""),
        ("UNALLOCATED", "Show the invoice; do not dump onto Consultant or Art Sales", f"EOEB remainder ${EOEB_REMAINING:,.2f} per-invoice (see DEAL_CLASSIFIER.csv)", f"L5 remainder ${L5_REMAINING:,.2f}; Aysel $50,000"),
    ]
    for i, row in enumerate(rows):
        fill = GREEN if i < 4 else YELLOW
        for c, val in enumerate(row, 1):
            ws.cell(start + 2 + i, c, val).fill = fill
            ws.cell(start + 2 + i, c).font = CG
            ws.cell(start + 2 + i, c).alignment = WRAP
        ws.row_dimensions[start + 2 + i].height = 36


def sheet_to_csv(ws, path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        max_c = min(ws.max_column or 1, 14)
        for r in range(1, (ws.max_row or 1) + 1):
            w.writerow([("" if ws.cell(r, c).value is None else ws.cell(r, c).value) for c in range(1, max_c + 1)])


def classify_ledger_row(date, counterparty, amount, bucket, status, deal) -> tuple[str, str, str, str, str]:
    amt = float(amount)
    if date == "2025-08-15" and abs(amt + 20000) < 0.01:
        return (
            "SALE (joint, EPGC share)",
            "ART_PURCHASE",
            "LOCKED",
            "August Fortuna joint — unnamed object Cost (Erdal)",
            "User 2026-09-21: another joint with Erdal. Native 8/18 = Monarch 8/15. "
            "Paired with Erdal 10/9 +$25,000. Not unsold inventory. Not consultant. "
            "7/24 is Venus. Not tax advice.",
        )
    if date == "2025-10-09" and abs(amt - 25000) < 0.01:
        return (
            "SALE (joint, EPGC share)",
            "ART_SALE",
            "LOCKED",
            "August Fortuna joint — EPGC proceeds (Erdal remittance)",
            "User 2026-09-21: $20k capital + $5k earnings. Hits EPGC Art Sales October. "
            "11/7 $6,000 still ASK. 7/24 is Venus. Not tax advice.",
        )
    if date == "2025-11-07" and abs(amt - 6000) < 0.01:
        return (
            "UNALLOCATED",
            "ASK",
            "ASK",
            "Erdal Dere Fortuna PAYMENT $6,000 IN — still ASK",
            "Not the August joint (that is 10/9 $25k). Not Venus (7/24 $26k). Not tax advice.",
        )
    if date == "2025-01-16" and abs(amt - 1000) < 0.01 and "EOEB" in counterparty:
        return (
            "UNALLOCATED",
            "UNALLOCATED",
            "ASK",
            "EOEB $1,000 — Canosan horse vs advisory (STILL OPEN)",
            "OPEN. Art Sales register horse is to Erdal $1,000; this cash is EOEB. "
            "Do not lock as the horse sale. Proposed FEE (no matching Cost) but UNALLOCATED. "
            "Not tax advice.",
        )
    if date == "2025-02-21" and abs(amt - 105000) < 0.01:
        return ("SALE", "ART_SALE", "MATCHED", "Three Ancient Mosaics — Sale Price", "")
    if date == "2025-12-02" and abs(amt - 565.75) < 0.01 and "EOEB" in counterparty:
        return ("n/a", "REIMBURSE", "LOCKED", "EOEB reimburses Wise $565.75", "Not income. Not expense.")
    if date == "2025-05-01" and abs(amt - 21000) < 0.01:
        return (
            "UNALLOCATED",
            "UNALLOCATED",
            "ASK",
            "EOEB 5/1 $21,000 — UNALLOCATED (proposed FEE; NOT Venus)",
            "NOT Venus (Venus Cost is Fortuna 5/9 $20,000). No matching dealer Cost on 8291. "
            + FEE_HYPOTHESIS,
        )
    if "EOEB" in counterparty and status in {"ASK", "ADVISORY"} or (
        "EOEB" in counterparty and bucket in {"ADVISORY", "UNALLOCATED", "ASK"} and abs(amt) not in {105000.0, 565.75}
    ):
        if abs(amt - 105000) < 0.01 or abs(amt - 565.75) < 0.01:
            pass
        else:
            return (
                "UNALLOCATED",
                "UNALLOCATED",
                "ASK",
                f"EOEB {date[5:7].lstrip('0')}/{date[8:].lstrip('0')} ${amt:,.2f} — UNALLOCATED (proposed FEE, no matching Cost)",
                FEE_HYPOTHESIS,
            )
    if date == "2025-07-15" and abs(amt - 50000) < 0.01:
        return ("SALE", bucket, "LOCKED", deal, "")
    if date == "2025-07-21" and abs(amt + 45000) < 0.01:
        return ("SALE", bucket, "LOCKED", deal, "")
    if date == "2025-05-09" and abs(amt + 20000) < 0.01:
        return ("SALE (joint, EPGC share)", bucket, "LOCKED", deal, "")
    if date == "2025-07-24" and abs(amt - 26000) < 0.01:
        return ("SALE (joint, EPGC share)", bucket, "LOCKED", deal, "")
    if "David Aaron" in deal or (date == "2025-06-24" and abs(amt - 13595) < 0.01):
        return ("FEE", bucket, status, deal, "")
    if bucket == "PASS_THROUGH":
        return ("PASS", bucket, status, deal, "")
    if counterparty == "L5" and not (date == "2025-07-15" and abs(amt - 50000) < 0.01):
        return (
            "UNALLOCATED",
            "ADVISORY",
            "ASK",
            "L5 — mixed unallocated after Belt $50k SALE removed",
            f"Remaining L5 ${L5_REMAINING:,.2f} still ASK. Do not dump. Not tax advice.",
        )
    if "Aysel" in counterparty:
        return ("UNALLOCATED", bucket, "ASK", deal, "Not the belt. Monarch IN vs native Failed OUT still ASK.")
    if bucket in {"ART_SALE", "ART_PURCHASE"} and status in {"LOCKED", "MATCHED"}:
        return ("SALE" if "joint" not in deal.lower() and "Venus" not in deal and "August" not in deal else "SALE (joint, EPGC share)", bucket, status, deal, "")
    if bucket in {"TRANSFER", "TEST", "INVESTMENT", "COGS_JEWELRY", "REIMBURSE"}:
        return ("n/a", bucket, status, deal, "")
    if "expertise" in deal.lower() or (bucket == "EPGC_EXPENSE" and "Wise" in counterparty):
        return ("FEE", bucket, status, deal, "")
    return ("UNALLOCATED" if status == "ASK" else "n/a", bucket, status, deal, "")


def patch_mercury_ledger() -> None:
    path = CSV_DIR / "MERCURY_LEDGER.csv"
    rows = list(csv.reader(path.open(encoding="utf-8")))
    banner = (
        "Mercury Checking 8291 — EPGC LLC 2025. 64 unique cash txns. "
        "Deal Classifier: SALE vs FEE vs PASS vs UNALLOCATED. "
        "Roman Gold Belt = SALE. Venus = joint SALE. August Fortuna = joint SALE "
        "(Cost $20,000 / proceeds $25,000). Remaining EOEB $130,270 per-invoice "
        "UNALLOCATED (proposed FEE, no matching Cost — not dumped). "
        "Other L5 $37,396.32 = UNALLOCATED. David Aaron = FEE. Koziol/Ariadne = PASS. "
        "11/7 $6,000 FORTUNA PAYMENT still ASK. Not tax advice."
    )
    rows[0] = [banner] + [""] * 12
    header = rows[2]
    if "Classifier" not in header:
        header.insert(5, "Classifier")
        for r in rows[3:]:
            if len(r) >= 5:
                r.insert(5, "")
    idx = {name: i for i, name in enumerate(header)}
    cls_i, bkt_i, st_i = idx["Classifier"], idx["Bucket"], idx["Status"]
    deal_i, hits_i, notes_i = idx["Deal / object"], idx["Hits EPGC Art Sales?"], idx["Notes"]
    cons_i = idx.get("Hits EPGC Consultant?")
    date_i, who_i, amt_i = idx["Date"], idx["Counterparty"], idx["Amount"]
    for r in rows[3:]:
        if len(r) < notes_i + 1 or not r[date_i]:
            continue
        cls, bkt, st, deal, extra = classify_ledger_row(
            r[date_i], r[who_i], r[amt_i], r[bkt_i], r[st_i], r[deal_i]
        )
        r[cls_i] = cls
        r[bkt_i] = bkt
        r[st_i] = st
        r[deal_i] = deal
        if extra:
            r[notes_i] = extra
        if cls.startswith("SALE") and float(r[amt_i]) > 0 and bkt == "ART_SALE" and st in {"LOCKED", "MATCHED"}:
            r[hits_i] = "YES"
        if "August Fortuna joint — unnamed object Cost" in deal or "Venus — EPGC Cost" in deal or "Roman Gold Belt — Cost" in deal:
            r[hits_i] = "no (Cost of SALE)"
        if bkt == "UNALLOCATED" and "EOEB" in r[who_i] and cons_i is not None:
            r[cons_i] = "no — not dumped"
            r[hits_i] = "no — not dumped"
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


EOEB_INVOICES = [
    {
        "date": "2025-01-16",
        "amount": 1000.00,
        "object": "EOEB invoice — Canosan-horse OPEN (cash is EOEB; register horse is to Erdal)",
        "cost": "none on 8291 (Art Sales Canosan horse Cost is Hindman $900 sold TO Erdal — different deal)",
        "bucket": "UNALLOCATED",
        "evidence": "OPEN. Mixed early-year EOEB taking funds. Do not lock as the horse sale. Proposed FEE (no matching Cost) but UNALLOCATED until named.",
        "lock": "ASK / OPEN",
    },
    {
        "date": "2025-02-21",
        "amount": 105000.00,
        "object": "Three Ancient Mosaics / EOEB MAKE A PAYMENT",
        "cost": "Plutus & Mnemosyne 2/25 $49,990 + $10 tests = $50,000 of $90,000 Cost (NOT Fortuna). $40,000 Cost missing on 8291.",
        "bucket": "SALE",
        "evidence": "Matches Art Sales to Jonathan Yantis $105,000. LOCKED mosaics sale.",
        "lock": "LOCKED / MATCHED",
    },
    {
        "date": "2025-03-10",
        "amount": 14000.00,
        "object": "EOEB invoice 3/10 $14,000",
        "cost": "none on 8291 — Fortuna joints are NOT EOEB Cost",
        "bucket": "UNALLOCATED",
        "evidence": FEE_HYPOTHESIS,
        "lock": "UNALLOCATED (proposed FEE)",
    },
    {
        "date": "2025-05-01",
        "amount": 21000.00,
        "object": "EOEB invoice 5/1 $21,000 (NOT Venus)",
        "cost": "none on 8291. Venus Cost is Fortuna 5/9 −$20,000 — different object, different counterparty.",
        "bucket": "UNALLOCATED",
        "evidence": "User: 5/1 $21,000 is NOT Venus. " + FEE_HYPOTHESIS,
        "lock": "UNALLOCATED (proposed FEE)",
    },
    {
        "date": "2025-05-13",
        "amount": 16200.00,
        "object": "EOEB invoice 5/13 $16,200",
        "cost": "none on 8291 — Fortuna joints are NOT EOEB Cost",
        "bucket": "UNALLOCATED",
        "evidence": FEE_HYPOTHESIS,
        "lock": "UNALLOCATED (proposed FEE)",
    },
    {
        "date": "2025-06-12",
        "amount": 25000.00,
        "object": "EOEB invoice 6/12 $25,000",
        "cost": "none on 8291 — Fortuna joints are NOT EOEB Cost",
        "bucket": "UNALLOCATED",
        "evidence": FEE_HYPOTHESIS,
        "lock": "UNALLOCATED (proposed FEE)",
    },
    {
        "date": "2025-06-12",
        "amount": 3300.00,
        "object": "EOEB invoice 6/12 $3,300",
        "cost": "none on 8291 — Fortuna joints are NOT EOEB Cost",
        "bucket": "UNALLOCATED",
        "evidence": FEE_HYPOTHESIS,
        "lock": "UNALLOCATED (proposed FEE)",
    },
    {
        "date": "2025-07-08",
        "amount": 2000.00,
        "object": "EOEB invoice 7/8 $2,000",
        "cost": "none on 8291 — Fortuna joints are NOT EOEB Cost",
        "bucket": "UNALLOCATED",
        "evidence": "Later-year EOEB looks like sourcing/consultant fee rather than object sales. " + FEE_HYPOTHESIS,
        "lock": "UNALLOCATED (proposed FEE)",
    },
    {
        "date": "2025-08-29",
        "amount": 10000.00,
        "object": "EOEB invoice 8/29 $10,000",
        "cost": "none on 8291 — August Fortuna Cost is Erdal joint, NOT EOEB Cost",
        "bucket": "UNALLOCATED",
        "evidence": FEE_HYPOTHESIS,
        "lock": "UNALLOCATED (proposed FEE)",
    },
    {
        "date": "2025-09-19",
        "amount": 5000.00,
        "object": "EOEB invoice 9/19 $5,000",
        "cost": "none on 8291",
        "bucket": "UNALLOCATED",
        "evidence": FEE_HYPOTHESIS,
        "lock": "UNALLOCATED (proposed FEE)",
    },
    {
        "date": "2025-09-25",
        "amount": 3100.00,
        "object": "EOEB invoice 9/25 $3,100",
        "cost": "none on 8291",
        "bucket": "UNALLOCATED",
        "evidence": FEE_HYPOTHESIS,
        "lock": "UNALLOCATED (proposed FEE)",
    },
    {
        "date": "2025-10-01",
        "amount": 10000.00,
        "object": "EOEB invoice 10/1 $10,000",
        "cost": "none on 8291 — Erdal 10/9 $25k is August Fortuna proceeds, NOT EOEB Cost",
        "bucket": "UNALLOCATED",
        "evidence": FEE_HYPOTHESIS,
        "lock": "UNALLOCATED (proposed FEE)",
    },
    {
        "date": "2025-10-15",
        "amount": 16000.00,
        "object": "EOEB invoice 10/15 $16,000",
        "cost": "none on 8291",
        "bucket": "UNALLOCATED",
        "evidence": FEE_HYPOTHESIS,
        "lock": "UNALLOCATED (proposed FEE)",
    },
    {
        "date": "2025-10-24",
        "amount": 1500.00,
        "object": "EOEB invoice 10/24 $1,500",
        "cost": "none on 8291",
        "bucket": "UNALLOCATED",
        "evidence": FEE_HYPOTHESIS,
        "lock": "UNALLOCATED (proposed FEE)",
    },
    {
        "date": "2025-10-31",
        "amount": 2170.00,
        "object": "EOEB invoice 10/31 $2,170",
        "cost": "none on 8291",
        "bucket": "UNALLOCATED",
        "evidence": FEE_HYPOTHESIS,
        "lock": "UNALLOCATED (proposed FEE)",
    },
    {
        "date": "2025-12-02",
        "amount": 565.75,
        "object": "EOEB Wise reimburse (same-day Wise OUT $565.75)",
        "cost": "n/a — nets to $0",
        "bucket": "PASS",
        "evidence": "LOCKED not income. Monarch tag Reimburse. Distinct from Wise expertise write-ups.",
        "lock": "LOCKED not income",
    },
]


def write_deal_classifier() -> None:
    path = CSV_DIR / "DEAL_CLASSIFIER.csv"
    rows = [
        [
            "Deal Classifier — Mercury 8291 / EPGC 2025. "
            "Columns: date, amount, object/invoice, matching Cost out?, proposed bucket SALE/FEE/PASS/UNALLOCATED, evidence, lock status. "
            "Last-year model: Art Sales = gross sale cash by month; Income I8 = art NET for lots with Cost; Consultant never goes to I8. "
            "Do NOT dump EOEB remainder $130,270 as a lump onto Consultant or Art Sales. Not tax advice.",
            "",
            "",
            "",
            "",
            "",
            "",
        ],
        ["date", "amount", "object/invoice", "matching Cost out?", "proposed bucket", "evidence", "lock status"],
    ]
    for inv in EOEB_INVOICES:
        rows.append(
            [
                inv["date"],
                f"{inv['amount']:.2f}",
                inv["object"],
                inv["cost"],
                inv["bucket"],
                inv["evidence"],
                inv["lock"],
            ]
        )
    rows.append(
        [
            "remainder total",
            f"{EOEB_REMAINING:.2f}",
            "EOEB remainder after mosaics $105,000 and Wise reimburse $565.75",
            "none on 8291 as a group",
            "UNALLOCATED",
            "Sum of OPEN + proposed-FEE invoices above. Shown per invoice — not posted to EPGC Consultant or Art Sales.",
            "UNALLOCATED (do not dump)",
        ]
    )
    extra = [
        ["2025-07-15", "50000.00", "Roman Gold Belt / L5", "Fortuna 7/21 −$45,000 (NOT EOEB)", "SALE", "Only L5 payment attached to the belt. Remaining L5 not attached.", "LOCKED"],
        ["2025-07-21", "-45000.00", "Roman Gold Belt / Fortuna Cost", "this row is the Cost", "SALE", "Bought from Fortuna, sold to L5. Not unsold inventory. Not EOEB Cost.", "LOCKED"],
        ["2025-07-24", "26000.00", "Venus / Erdal remittance", "Fortuna 5/9 −$20,000 (NOT EOEB)", "SALE", "Joint. Sale Price = EPGC proceeds, not full hammer. 7/24 cannot be August.", "LOCKED"],
        ["2025-05-09", "-20000.00", "Venus / Fortuna Cost", "this row is the Cost", "SALE", "Joint with Erdal. Not EOEB Cost.", "LOCKED"],
        ["2025-10-09", "25000.00", "August Fortuna joint / Erdal remittance", "Fortuna 8/15 −$20,000 (NOT EOEB)", "SALE", "Joint. $20k capital + $5k earnings. EPGC Art Sales October.", "LOCKED"],
        ["2025-08-15", "-20000.00", "August Fortuna joint / Fortuna Cost (native 8/18)", "this row is the Cost", "SALE", "Unnamed object. Not unsold inventory. Not consultant. Not EOEB Cost.", "LOCKED"],
        ["2025-11-07", "6000.00", "Erdal Dere FORTUNA PAYMENT", "none paired", "UNALLOCATED", "Still ASK. Not August (10/9). Not Venus (7/24).", "ASK"],
        ["2025-01-10", "14000.00", "Berk seals", "Fortuna 1/8 $13,000", "SALE", "I8 net $1,000.", "LOCKED / MATCHED"],
        ["2025-02-07", "30000.00", "Berk lots lump", "rows 51–70 Cost", "SALE", "F71 lump. Hits EPGC Art Sales February.", "LOCKED / MATCHED"],
        ["2025-02-04", "1000.00", "Aquinas Hobor books", "Cost TBD", "SALE", "Hits EPGC February. Not in I8 until Cost.", "LOCKED"],
        ["2025-06-24", "13595.00", "David Aaron Limited", "n/a — fee not a sale", "FEE", "EPGC Consultant June only. Not Art Sales. Not I8.", "LOCKED"],
        ["2025-07-09 / 2025-10-30", "-992.80", "Wise expertise write-ups", "n/a — expense", "FEE", "EPGC Consultant Fees. Dec $565.75 is reimbursed, not this.", "LOCKED"],
        ["2025-10-20 / 2025-10-21", "150000.00", "Koziol → Ariadne", "n/a — pass-through", "PASS", "Not P&L.", "LOCKED"],
        ["various", f"{L5_REMAINING:.2f}", "L5 remainder after belt", "none attached to remaining invoices", "UNALLOCATED", "7/11 $8,534.79; 7/15 $2,500; 9/2 $5,042; 9/23 $9,119.27; 10/21 $7,500; 12/4 $4,700.26. Do not dump.", "UNALLOCATED"],
        ["2025-07-18", "50000.00", "Aysel Dere Monarch IN", "native Failed OUT 7/17 −$50,000", "UNALLOCATED", "Not the belt. Do not dump.", "ASK"],
    ]
    rows.extend(extra)
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def write_ask_mercury() -> None:
    path = CSV_DIR / "ASK_Mercury.csv"
    banner = (
        "ASK — remaining Mercury 8291 questions after Belt + Venus + August Fortuna SALE lock 2026-09-21. "
        "EOEB remainder is per-invoice UNALLOCATED (proposed FEE) — not dumped. "
        "Q6 consultant, Q7 Wise expertise, Q8 pass-through, Q9 Aquinas books LOCKED. Not tax advice."
    )
    rows = [
        [banner, "", ""],
        ["", "", ""],
        ["Question", "What I need", "Classifier"],
        [
            "1. EOEB LLC remainder after mosaics — per-invoice UNALLOCATED (proposed FEE, not dumped)",
            f"${EOEB_REMAINING:,.2f} of EOEB IN is not mosaics $105,000 and not Wise reimburse $565.75. "
            "See DEAL_CLASSIFIER.csv for every invoice. Hypothesis: no matching dealer Cost → looks like FEE. "
            "Do NOT dump this remainder onto Consultant or Art Sales. 1/16 $1,000 Canosan-horse still OPEN. "
            "5/1 $21,000 is NOT Venus. Fortuna outs are NOT EOEB Cost.",
            "UNALLOCATED",
        ],
        [
            "2. L5 remainder after Belt $50,000 SALE — MIXED PATTERN CONFIRMED, dollars unallocated",
            f"${L5_REMAINING:,.2f} remaining after locking 7/15 $50,000 as Roman Gold Belt SALE. "
            "Do NOT attach these invoices to the belt. Do NOT dump onto Consultant or Art Sales. "
            "7/11 $8,534.79; 7/15 $2,500; 9/2 $5,042; 9/23 $9,119.27; 10/21 $7,500; 12/4 $4,700.26.",
            "UNALLOCATED",
        ],
        [
            "3. August Fortuna joint — LOCKED SALE (object still unnamed)",
            "Fortuna 8/15 (native 8/18) −$20,000 Cost / Erdal 10/9 +$25,000 proceeds "
            "($20k capital + $5k earnings). Hits EPGC Art Sales October. Object name still unknown. "
            "Not consultant. Not unsold inventory.",
            "SALE (joint, EPGC share)",
        ],
        [
            "4. Erdal Dere 11/7 $6,000 FORTUNA PAYMENT — still ASK",
            f"${ERDAL_REMAINING:,.2f}. Not the August joint (10/9 $25,000). Not Venus (7/24 $26,000).",
            "UNALLOCATED",
        ],
        [
            "5. Aysel Dere $50,000 — Monarch IN vs native Failed OUT",
            "Monarch 7/18 +$50,000 IN is in the 64. Native Failed $50,000 OUT 7/17. Not the belt. Do not dump.",
            "UNALLOCATED",
        ],
        [
            "6. David Aaron Limited $13,595 IN (6/24) — LOCKED consultant fee",
            "CONFIRMED: consultant fee, not a sale. Booked EPGC Consultant June $13,595. Not Art Sales. Not I8.",
            "FEE",
        ],
        [
            "7. Wise $334.17 (7/9) and $658.63 (10/30) — LOCKED expertise write-ups",
            "EPGC Consultant Fees $992.80. Distinct from Dec Wise $565.75 reimbursed by EOEB.",
            "FEE",
        ],
        [
            "8. Jack Koziol & Tracy Hoffman $150,000 IN (10/20) / Ariadne Demirjian $150,000 OUT (10/21) — LOCKED pass-through",
            "CONFIRMED: pass-through, not P&L.",
            "PASS",
        ],
        [
            "9. Aquinas Hobor $1,000 IN (2/4) — LOCKED book sale",
            "Hits EPGC Art Sales February. Cost / titles TBD (not in I8 until Cost).",
            "SALE",
        ],
        [
            "10. Mosaics Cost $90,000 vs Mercury $50,000",
            "Plutus & Mnemosyne paid $50,000 in February. Where is the other $40,000 of Cost?",
            "ASK",
        ],
        ["", "", ""],
        ["LOCKED / MATCHED this pass (do not recast without saying so)", "", "Classifier"],
        ["Roman Gold Belt", "Fortuna 7/21 −$45,000 Cost / L5 7/15 +$50,000 Sale. Net $5,000. Remaining L5 not attached.", "SALE"],
        ["Venus (bronze head, joint)", "Fortuna 5/9 −$20,000 / Erdal 7/24 +$26,000. Sale Price = EPGC proceeds.", "SALE (joint, EPGC share)"],
        ["August Fortuna joint (unnamed)", "Fortuna 8/15 −$20,000 / Erdal 10/9 +$25,000. EPGC Art Sales October. Net $5,000.", "SALE (joint, EPGC share)"],
        ["Mosaics", "EOEB $105,000 sale (2/21) / Plutus $50,000 of $90,000 Cost", "SALE"],
        ["David Aaron Limited", "$13,595 (6/24) consultant fee → EPGC Consultant June. Not a sale.", "FEE"],
        ["Koziol / Ariadne", "$150,000 IN 10/20 + $150,000 OUT 10/21 LOCKED pass-through — not P&L", "PASS"],
        ["Wise/EOEB 12/2", "$565.75 reimbursed — not income, not expense", "n/a"],
        ["EOEB remainder", f"${EOEB_REMAINING:,.2f} per-invoice UNALLOCATED (proposed FEE). Not dumped.", "UNALLOCATED"],
    ]
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def write_fortuna() -> None:
    path = CSV_DIR / "FORTUNA_INVENTORY.csv"
    rows = [
        ["Object", "Source", "Cost", "Date", "Status", "Classifier", "Notes"],
        [
            "Bronze head of a goddess likely Venus",
            "Mercury 8291 Fortuna",
            "20000.00",
            "2025-05-09",
            "LOCKED SOLD 2025",
            "SALE (joint, EPGC share)",
            "Erdal remitted 7/24 +$26,000 = $20,000 + $6,000 earnings. Sale Price = EPGC proceeds.",
        ],
        [
            "Roman Gold Belt",
            "Mercury 8291 Fortuna",
            "45000.00",
            "2025-07-21",
            "LOCKED SOLD 2025",
            "SALE",
            "Sold to L5 7/15 +$50,000. Net $5,000. Remaining L5 not attached.",
        ],
        [
            "August Fortuna joint — unnamed object",
            "Mercury 8291 Fortuna",
            "20000.00",
            "2025-08-15",
            "LOCKED SOLD 2025",
            "SALE (joint, EPGC share)",
            "Monarch 8/15 = native 8/18. Erdal 10/9 +$25,000 = $20k capital + $5k earnings. Not unsold inventory. Not consultant. Object name still unknown.",
        ],
    ]
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def write_lock_q() -> None:
    path = CSV_DIR / "LOCK_Mercury_Q6_Q9.csv"
    rows = [
        ["Q", "Status", "What", "Where it hits"],
        ["6", "LOCKED", "David Aaron Limited $13,595 IN (6/24) consultant fee", "EPGC Consultant June $13,595. Not Art Sales. Not I8. Classifier: FEE."],
        ["7", "LOCKED", "Wise $334.17 (7/9) + $658.63 (10/30) expertise write-ups", "EPGC Consultant Fees $992.80. Classifier: FEE."],
        ["8", "LOCKED", "Koziol $150,000 IN (10/20) / Ariadne $150,000 OUT (10/21)", "Pass-through — not P&L. Classifier: PASS."],
        ["9", "LOCKED", "Aquinas Hobor $1,000 IN (2/4) book SALE", "EPGC Art Sales February. Cost TBD. Classifier: SALE."],
        ["Belt", "LOCKED", "Roman Gold Belt Fortuna $45,000 / L5 $50,000", "EPGC Art Sales July +$50,000. I8 net +$5,000. Classifier: SALE."],
        ["Venus", "LOCKED", "Venus Fortuna $20,000 / Erdal remittance $26,000", "EPGC Art Sales July +$26,000. I8 net +$6,000. Classifier: SALE (joint, EPGC share)."],
        ["August", "LOCKED", "August Fortuna joint Fortuna $20,000 / Erdal 10/9 $25,000", "EPGC Art Sales October +$25,000. I8 net +$5,000. Classifier: SALE (joint, EPGC share). 11/7 $6,000 still ASK."],
        ["EOEB mosaics", "LOCKED", "EOEB 2/21 $105,000 mosaics", "EPGC Art Sales February. Classifier: SALE."],
        ["EOEB 12/2", "LOCKED not income", "EOEB/Wise $565.75 reimburse", "Not P&L."],
        ["1", "UNALLOCATED per invoice (proposed FEE)", f"EOEB remainder ${EOEB_REMAINING:,.2f}", "DEAL_CLASSIFIER.csv. Do NOT dump onto Consultant or Art Sales."],
        ["2", "PATTERN CONFIRMED / dollars unallocated", f"L5 remainder ${L5_REMAINING:,.2f}", "Do NOT dump. Classifier: UNALLOCATED."],
        ["LOCKED totals", "LOCKED", f"Art Sales gross $251,000 (Jan 14k + Feb 136k + July 76k + Oct 25k); I8 $39,055.06; Consultant June $13,595", "EOEB remainder / other L5 not on P&L."],
    ]
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


START_HERE_MERCURY = """START HERE — Mercury Bank 8291 (2026-09-21 Belt + Venus + August Fortuna + EOEB classifier)

Not tax advice. Same Personal Income.xlsx tabs as last year. Organizational packet only.

Mercury Checking 8291 is the EPGC LLC operating account (Choice Financial).
64 unique 2025 cash transactions from Monarch IDs.

Deal Classifier (per-invoice EOEB — do not dump remainder):
https://docs.google.com/spreadsheets/d/CLASSIFIER_PLACEHOLDER/edit

One spreadsheet like last year (11 tabs + Belt, Venus, August joint on Art Sales):
https://docs.google.com/spreadsheets/d/XLSX_PLACEHOLDER/edit

LOCKED this pass
- Roman Gold Belt SALE: Fortuna Cost $45,000 / L5 Sale $50,000 / net $5,000 → EPGC Art Sales July + I8.
- Venus joint SALE: Fortuna Cost $20,000 / Erdal proceeds $26,000 / net $6,000 → EPGC Art Sales July + I8.
- August Fortuna joint SALE: Fortuna Cost $20,000 (8/15, native 8/18) / Erdal proceeds $25,000 (10/9) / net $5,000 → EPGC Art Sales October + I8.
- David Aaron Limited $13,595 (6/24) → EPGC Consultant June. FEE, not a sale.
- Koziol $150,000 (10/20) / Ariadne $150,000 (10/21) → PASS. Not P&L.
- Wise expertise write-ups $992.80 → EPGC Consultant Fees.
- EOEB mosaics $105,000 (2/21) SALE. EOEB 12/2 $565.75 reimbursed — not income.

EPGC 2025 Art Sales LOCKED cash $251,000 (Jan $14,000 / Feb $136,000 / July $76,000 / Oct $25,000).
EPGC Consultant June is $13,595 (David Aaron only).
Income I8 Art net is $39,055.06. I10 EBITDA =SUM(I4:I9) → $167,760.

ASK remaining — do not dump
1. EOEB remainder $130,270.00 — per-invoice UNALLOCATED (proposed FEE, no matching Cost). Jan 16 $1,000 Canosan-horse OPEN. 5/1 $21,000 is NOT Venus.
2. Other L5 $37,396.32 after removing $50k belt.
3. Erdal 11/7 $6,000 FORTUNA PAYMENT — not August, not Venus.
4. Aysel Dere — Monarch 7/18 +$50,000 IN vs native Failed 7/17 −$50,000 OUT. Not the belt.
10. Mosaics Cost $40,000 missing on Mercury 8291.

Sale 6428 $11,000 and Aquinas books $1,000 wait on Cost (not in I8).
"""

START_HERE_ALL = """START HERE — last year's spreadsheet, 2025 numbers filled
Not tax advice. Organizational packet only.

Open this ONE spreadsheet (same 11 tab names as last year, 2023 | 2024 | 2025 year blocks, 2025 numbers, formulas that calculate, Art Sales has 2022–2024 lots + 2025 Belt, Venus, August joint):
https://docs.google.com/spreadsheets/d/XLSX_PLACEHOLDER/edit

Deal Classifier (EOEB per invoice — proposed FEE, lock UNALLOCATED):
https://docs.google.com/spreadsheets/d/CLASSIFIER_PLACEHOLDER/edit

START HERE doc:
https://docs.google.com/document/d/DOC_PLACEHOLDER/edit

Tabs in last year's order (all inside that one file)
1. Income — 2022–2025. 2025 Actual: Hindman $70,618 / 216 $19,500 / 524 #2 $17,086.94 / GCM $21,500 / Art $39,055.06. I10 EBITDA =SUM(I4:I9) → $167,760.
2. 524 Ferdinand Ave, Unit 2
3. 216 N. Oak Park Ave
4. EPGC LLC — 2025 Art Sales $14,000 + $136,000 + July $76,000 + October $25,000 = $251,000; Consultant June $13,595 (David Aaron only).
5. Refrence Library
6. Art Sales and Purchases — 2022–2024 lots, then Sales 2025 including Roman Gold Belt, Venus joint, August Fortuna joint. Do not dump EOEB $130,270.
7. GCM
8. Hindman W2
9. Megan T4
10. Investments — Coinbase ACH net -$51,000.
11. 524 Ferdinand Ave, Unit 1

Deal Classifier: SALE vs FEE vs PASS vs UNALLOCATED.
Belt = SALE. Venus = SALE (joint). August Fortuna = SALE (joint). EOEB remainder = UNALLOCATED per invoice (proposed FEE). David Aaron = FEE.

How this stays correct
1. Copy last year's Personal Income.xlsx. Never rebuild from CSVs.
2. Fill the 2025 year block only. 2023 and 2024 stay as they were.
3. python3 .cursor/scratch/assert_personal_income_2025.py must print LOCK OK + SHEETS-SAFE OK.
4. Drive upload is a compact xlsx Google converts. Include Art Sales row 1 and 2022–2024 lots.
5. CSV builders (.cursor/scratch/build_drive_all_tabs.py) now refuse to run.

Not tax advice.
"""


def write_start_here() -> None:
    (CSV_DIR / "START_HERE_mercury.txt").write_text(START_HERE_MERCURY, encoding="utf-8")
    (CSV_DIR / "START_HERE_all_tabs.txt").write_text(START_HERE_ALL, encoding="utf-8")


def patch_json() -> None:
    packet = json.loads(PACKET.read_text(encoding="utf-8"))
    packet["updated"] = "2026-09-21"
    packet["xlsx_note"] = (
        "2026-09-21: Live file is ONE Google Sheet with last year's 11 tab names, "
        "2023|2024|2025 year blocks, 2025 numbers, calculating SUM formulas. "
        "Art Sales includes 2022–2024 lots plus Belt, Venus joint, and August Fortuna joint. "
        f"Income I8 $39,055.06. I10 $167,760. EPGC Art Sales July $76,000 + October $25,000 "
        "(year $251,000). EOEB remainder $130,270 per-invoice UNALLOCATED. Not tax advice."
    )
    m = packet.setdefault("mercury_8291", {})
    m["updated"] = "2026-09-21"
    locked = m.setdefault("locked", {})
    locked["art_sale_locked_gross"] = ART_GROSS
    locked["art_net_locked"] = I8
    locked["roman_gold_belt_sale"] = 50000.0
    locked["roman_gold_belt_cost"] = 45000.0
    locked["roman_gold_belt_net"] = 5000.0
    locked["venus_epgc_proceeds"] = 26000.0
    locked["venus_epgc_cost"] = 20000.0
    locked["venus_epgc_net"] = 6000.0
    locked["august_fortuna_proceeds"] = 25000.0
    locked["august_fortuna_cost"] = 20000.0
    locked["august_fortuna_net"] = 5000.0
    locked.pop("fortuna_venus_inventory", None)
    locked.pop("fortuna_roman_gold_belt", None)
    ask = m.setdefault("ask", {})
    ask["eoeb_remainder"] = EOEB_REMAINING
    ask["eoeb_remainder_note"] = (
        "Per-invoice UNALLOCATED (proposed FEE, no matching Cost). "
        "Do not dump onto Consultant or Art Sales. 1/16 Canosan OPEN. 5/1 not Venus."
    )
    ask["l5"] = L5_REMAINING
    ask["l5_after_belt"] = L5_REMAINING
    ask["fortuna_remainder_out"] = 0.0
    ask["fortuna_named_inventory"] = 0.0
    ask["erdal_in"] = ERDAL_REMAINING
    ask["erdal_in_note"] = "11/7 $6,000 FORTUNA PAYMENT still ASK; 7/24 Venus LOCKED; 10/9 August joint LOCKED"
    m["epgc_art_sales_2025_monthly"] = [
        14000.0, 136000.0, 0.0, 0.0, 0.0, 0.0, 76000.0, 0.0, 0.0, 25000.0, 0.0, 0.0
    ]
    packet["deal_classifier"] = {
        "classes": ["SALE", "FEE", "PASS", "UNALLOCATED"],
        "belt": "SALE",
        "venus": "SALE (joint, EPGC share)",
        "august_fortuna": "SALE (joint, EPGC share)",
        "david_aaron": "FEE",
        "koziol_ariadne": "PASS",
        "eoeb_mosaics": "SALE",
        "eoeb_12_2_reimburse": "LOCKED not income",
        "eoeb_remainder": "UNALLOCATED (proposed FEE per invoice; not dumped)",
        "l5_remainder": "UNALLOCATED",
        "erdal_11_7": "ASK",
        "csv": ".cursor/scratch/cc_fill_csv/DEAL_CLASSIFIER.csv",
        "disclaimer": "Not tax advice. Organizational packet only.",
    }
    PACKET.write_text(json.dumps(packet, indent=2) + "\n", encoding="utf-8")

    mercury = json.loads(MERCURY_JSON.read_text(encoding="utf-8"))
    mercury["updated"] = "2026-09-21"
    mercury.setdefault("locked", {}).update(locked)
    mercury.setdefault("ask", {}).update(ask)
    mercury["epgc_art_sales_2025_monthly"] = packet["mercury_8291"]["epgc_art_sales_2025_monthly"]
    mercury["deal_classifier"] = packet["deal_classifier"]
    MERCURY_JSON.write_text(json.dumps(mercury, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    wb = load_workbook(XLSX)
    patch_art(wb["Art Sales and Purchases"])
    patch_income(wb["Income"])
    patch_epgc(wb["EPGC LLC"])
    patch_data_sources(wb["2025 Data sources"])
    try:
        wb.defined_names.clear()
    except Exception:
        pass
    wb._external_links = []
    wb.save(XLSX)
    sanitize_xlsx(XLSX)
    DELIVERABLE.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(XLSX, DELIVERABLE)

    sheet_to_csv(wb["Art Sales and Purchases"], CSV_DIR / "Art_Sales.csv")
    sheet_to_csv(wb["Income"], CSV_DIR / "Income.csv")
    sheet_to_csv(wb["EPGC LLC"], CSV_DIR / "EPGC_LLC.csv")

    patch_mercury_ledger()
    write_ask_mercury()
    write_fortuna()
    write_deal_classifier()
    write_lock_q()
    write_start_here()
    patch_json()

    print("locked august+eoeb classifier")
    print("xlsx", XLSX, XLSX.stat().st_size)
    print("I8", I8, "July art", EPGC_JULY_GROSS, "Oct art", EPGC_OCT_GROSS, "gross", ART_GROSS)
    print("I10 expected", I10_EXPECTED)
    print("L5 remaining", L5_REMAINING, "Erdal remaining", ERDAL_REMAINING)
    print("EOEB remainder UNALLOCATED", EOEB_REMAINING)


if __name__ == "__main__":
    main()

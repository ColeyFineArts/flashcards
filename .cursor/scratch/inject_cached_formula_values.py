#!/usr/bin/env python3
"""Write numeric cached <v> next to each simple <f>SUM/minus formula.

Uses ElementTree so we never swallow later rows (a regex did that).
Google Sheets then shows the number immediately while keeping the formula.

Not tax advice.
"""
from __future__ import annotations

import io
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

ROOT = Path("/workspace/.cursor/scratch")
DEFAULT = ROOT / "tax_turbo_parts" / "Personal Income 2025.drive.xlsx"
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
SUM_RE = re.compile(r"^SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)$")
MINUS_RE = re.compile(r"^([A-Z]+)(\d+)-([A-Z]+)(\d+)$")


def _num(val) -> float:
    if val is None or val == "":
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    return 0.0


def _fmt(n: float) -> str:
    if abs(n - round(n)) < 1e-9:
        return str(int(round(n)))
    return f"{n:.12g}"


def cached_map(path: Path) -> dict[tuple[str, str], str]:
    wb = load_workbook(path, data_only=False)
    out: dict[tuple[str, str], str] = {}
    for idx, name in enumerate(wb.sheetnames, start=1):
        ws = wb[name]
        xml_name = f"xl/worksheets/sheet{idx}.xml"
        values: dict[tuple[int, int], float] = {}
        formulas: list[tuple[int, int, str]] = []
        for r in range(1, (ws.max_row or 1) + 1):
            for c in range(1, (ws.max_column or 1) + 1):
                cell = ws.cell(r, c)
                if cell.data_type == "f" and isinstance(cell.value, str):
                    formulas.append(
                        (r, c, cell.value[1:] if cell.value.startswith("=") else cell.value)
                    )
                else:
                    values[(r, c)] = _num(cell.value)

        def lookup(rr: int, cc: int) -> float:
            return values.get((rr, cc), 0.0)

        for r, c, fml in formulas:
            m = SUM_RE.fullmatch(fml)
            if not m:
                continue
            c1, r1, c2, r2 = (
                column_index_from_string(m.group(1)),
                int(m.group(2)),
                column_index_from_string(m.group(3)),
                int(m.group(4)),
            )
            total = 0.0
            for rr in range(r1, r2 + 1):
                for cc in range(c1, c2 + 1):
                    total += lookup(rr, cc)
            values[(r, c)] = total
            out[(xml_name, f"{get_column_letter(c)}{r}")] = _fmt(total)

        for r, c, fml in formulas:
            m = MINUS_RE.fullmatch(fml)
            if not m:
                continue
            left = lookup(int(m.group(2)), column_index_from_string(m.group(1)))
            right = lookup(int(m.group(4)), column_index_from_string(m.group(3)))
            total = left - right
            values[(r, c)] = total
            out[(xml_name, f"{get_column_letter(c)}{r}")] = _fmt(total)
    return out


def _patch_sheet(xml: str, sheet_path: str, caches: dict[tuple[str, str], str]) -> tuple[str, int]:
    ET.register_namespace("", NS)
    root = ET.fromstring(xml)
    patched = 0
    for cell in root.iter(f"{{{NS}}}c"):
        f = cell.find(f"{{{NS}}}f")
        if f is None:
            continue
        ref = cell.get("r")
        val = caches.get((sheet_path, ref or ""))
        if val is None:
            continue
        v = cell.find(f"{{{NS}}}v")
        if v is None:
            v = ET.SubElement(cell, f"{{{NS}}}v")
        v.text = val
        patched += 1
    body = ET.tostring(root, encoding="unicode")
    if body.startswith("<?xml"):
        return body, patched
    return '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n' + body, patched


def inject(path: Path) -> int:
    caches = cached_map(path)
    src = zipfile.ZipFile(path)
    buf = io.BytesIO()
    patched = 0
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet") and item.filename.endswith(".xml"):
                xml, n = _patch_sheet(data.decode("utf-8"), item.filename, caches)
                patched += n
                data = xml.encode("utf-8")
            out.writestr(item, data)
    path.write_bytes(buf.getvalue())
    return patched


if __name__ == "__main__":
    import sys

    target = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT
    n = inject(target)
    print(f"cached {n} formula values in {target} ({target.stat().st_size} bytes)")

#!/usr/bin/env python3
"""Lock Roman Gold Belt + Venus (joint EPGC share) as 2025 art SALES.

User 2026-09-21:
  Roman Gold Belt: Fortuna Cost 7/21 −$45,000; L5 7/15 +$50,000 (only L5
  payment large enough). Sold to L5. Net $5,000. Not unsold inventory.
  Do NOT attach remaining L5 invoices.
  Venus: Erdal and Jacob bought together. Fortuna 5/9 −$20,000 EPGC stake.
  Erdal sold it and remitted 7/24 +$26,000 = $20,000 + $6,000 earnings.
  Art Sales Sale Price = EPGC proceeds, not the object's full hammer.

Does NOT dump EOEB remainder, other L5, other Erdal IN, Aysel, or Fortuna Aug.
Does NOT re-run apply_checking_answers_2025.py or apply_mercury_2025.py.
Not tax advice. Organizational packet only.
"""
from __future__ import annotations

import csv
import json
import shutil
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import sys

sys.path.insert(0, str(Path("/workspace/.cursor/scratch")))
from sanitize_xlsx_for_sheets import sanitize_xlsx

ROOT = Path("/workspace/.cursor/scratch")
PARTS = ROOT / "tax_turbo_parts"
CSV_DIR = ROOT / "cc_fill_csv"
XLSX = PARTS / "Personal Income 2025.xlsx"
DELIVERABLE = ROOT / "tax_turbo_deliverable" / "Personal Income 2025.xlsx"
PACKET = ROOT / "cpa_packet_v1.json"
MERCURY_JSON = ROOT / "mercury_2025.json"
DRIVE_MAP = ROOT / "drive_folder_map.json"

GREEN = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
PEACH = PatternFill("solid", fgColor="F7CAAC")
BLUE = PatternFill("solid", fgColor="DDEBF7")
CG = Font(name="Century Gothic", size=10)
CG_B = Font(name="Century Gothic", size=10, bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
ACCT = '_("$"* #,##0.00_);_("$"* \\(#,##0.00\\);_("$"* "-"??_);_(@_)'

I8 = 34055.06  # 23055.06 + belt 5000 + venus 6000
EPGC_JULY_GROSS = 76000.0  # belt 50000 + venus 26000
ART_GROSS = 226000.0  # 150000 + 76000
L5_REMAINING = 37396.32
ERDAL_REMAINING = 31000.0
EOEB_REMAINING = 130270.0

BELT_NOTE = (
    "LOCKED 2026-09-21 SALE. Bought Fortuna/Erdal Dere 7/21 −$45,000 "
    "(Mercury 8291 native memo: Roman Gold Belt). Sold to L5. Working sale: "
    "L5 7/15 +$50,000 (only L5 payment large enough; L5 paid 6 days before "
    "Fortuna Cost — back-to-back dealer cash). Net $5,000. Do NOT attach "
    "remaining L5 invoices. Classifier: SALE. Not unsold inventory. Not consultant. "
    "Not pass-through. Not tax advice."
)
VENUS_NOTE = (
    "LOCKED 2026-09-21 joint SALE (EPGC share). Erdal and Jacob bought together. "
    "EPGC Cost: Fortuna 5/9 −$20,000 (native memo: bronze head of a goddess, "
    "likely Venus). Erdal sold it and paid back Jacob’s investment plus earnings. "
    "Working remittance: Erdal Dere IN 7/24 +$26,000 = $20,000 investment + $6,000 "
    "earnings. Sale Price on this register is EPGC proceeds ($26,000), not the "
    "object’s full hammer. Classifier: SALE (joint, EPGC share). Not unsold "
    "inventory. Not consultant. Not tax advice."
)
ART_FOOTER = (
    "Mercury 8291 MATCHED/LOCKED: Berk $14,000 (1/10 seals) + $30,000 (2/7 lots). "
    "EOEB $105,000 (2/21 mosaics). Fortuna $13,000 (1/8 seals Cost). "
    "Plutus $50,000 of mosaics $90,000 Cost — $40,000 of Cost not on 8291 (ASK). "
    "Roman Gold Belt SALE to L5: Cost $45,000 / Sale $50,000 / net $5,000. "
    "Venus joint SALE EPGC share: Cost $20,000 / proceeds $26,000 / net $6,000 "
    "(Sale Price = remittance, not full hammer). "
    "Newstar $23,581 jewelry fabrication COGS parked below. "
    "David Aaron $13,595 LOCKED consultant (EPGC Consultant June — not a sale). "
    "Aquinas Hobor $1,000 LOCKED book sale (Feb cash; Cost TBD, not in I8). "
    "Do not dump EOEB remainder $130,270 or other L5 $37,396.32. Not tax advice."
)
INCOME_NOTE = (
    "2025 Actual: I4 Monarch paycheck cash ≠ W-2 Box 1. I5 216 LTR cash $19,500. "
    "I6 524 #2 STR platform net LOCKED. I7 GCM 1099-NEC $21,500. "
    "I8 Art net $34,055.06 = prior MATCHED nets $23,055.06 (Berk lots + seals "
    "$1,000 + mosaics $15,000) + Belt net $5,000 + Venus EPGC-share net $6,000. "
    "Excludes Sale 6428 $11,000 until Cost. Excludes Aquinas books $1,000 until Cost. "
    "David Aaron $13,595 is EPGC Consultant June (not I8). "
    "EOEB remainder $130,270 / other L5 $37,396.32 / other Erdal IN $31,000 / "
    "Aysel $50,000 / Fortuna Aug $20,000 still ASK. Not tax advice."
)
EPGC_NOTE = (
    "2025 Mercury 2026-09-21: Art Sales gross cash by month = Jan $14,000 + Feb "
    "$136,000 MATCHED + July $76,000 (Belt L5 $50,000 + Venus remittance $26,000) "
    "= $226,000. Consultant June $13,595 = David Aaron Limited LOCKED (not a sale). "
    "Consultant Fees July $334.17 + October $658.63 = Wise expertise write-ups "
    "LOCKED ($992.80). EOEB remainder $130,270 and other L5 $37,396.32 MIXED "
    "PATTERN CONFIRMED — invoice-level split still ASK; not dumped. Other Erdal IN "
    "10/9 $25,000 and 11/7 $6,000 FORTUNA PAYMENT still ASK. Aysel Monarch +$50k vs "
    "native Failed −$50k still ASK (not the belt). Fortuna Aug unnamed $20,000 still "
    "ASK (not Venus, not belt). Koziol/Ariadne $150,000 LOCKED pass-through — not P&L. "
    "Coinbase is Investments. Newstar $23,581 is Art Sales COGS. Dec Wise $565.75 "
    "reimbursed. Aquinas books Cost TBD — cash on Art Sales, not in I8. "
    "Venus Sale Price = EPGC proceeds, not full hammer. GCM 1099 is personal. "
    "Not tax advice."
)


def style_like(src, dest) -> None:
    dest.font = copy(src.font) if src.has_style else CG
    dest.fill = copy(src.fill) if src.has_style else GREEN
    dest.alignment = copy(src.alignment) if src.has_style else Alignment()
    dest.number_format = src.number_format
    dest.border = copy(src.border) if src.has_style else dest.border


def paint(cell, value, fill=GREEN, num=None) -> None:
    cell.value = value
    cell.fill = fill
    cell.font = CG
    if num:
        cell.number_format = num


def patch_art(ws) -> None:
    ws["A1"] = "Art Sales and Purchases"
    ws["A1"].font = CG_B

    # Classifier header on year blocks that already have Object headers.
    for header_row in (3, 16, 39, 50):
        if str(ws.cell(header_row, 1).value or "") == "Object":
            ws.cell(header_row, 11, "Classifier").font = CG_B

    # Unmerge the footer note so insert_rows does not split it.
    to_unmerge = [str(rng) for rng in ws.merged_cells.ranges if rng.min_row == 79]
    for rng in to_unmerge:
        ws.unmerge_cells(rng)

    if str(ws["A78"].value or "") != "Roman Gold Belt":
        if str(ws["A78"].value or "") == "Total":
            ws.insert_rows(78, 2)

    src = ws["A73"]

    # Row 78 — Roman Gold Belt SALE
    paint(ws["A78"], "Roman Gold Belt")
    paint(ws["B78"], "Fortuna / Erdal Dere")
    paint(ws["C78"], 45000, num=ACCT)
    paint(ws["D78"], datetime(2025, 7, 21))
    ws["D78"].number_format = "YYYY-MM-DD"
    paint(ws["E78"], "L5")
    paint(ws["F78"], 50000, num=ACCT)
    paint(ws["G78"], datetime(2025, 7, 15))
    ws["G78"].number_format = "YYYY-MM-DD"
    paint(ws["H78"], "=F78-C78")
    ws["H78"].number_format = ACCT
    paint(ws["I78"], BELT_NOTE)
    ws["I78"].alignment = WRAP
    paint(ws["K78"], "SALE")
    ws.row_dimensions[78].height = 48
    for col in range(1, 12):
        ws.cell(78, col).font = CG
        if ws.cell(78, col).fill.fill_type in (None, "none"):
            ws.cell(78, col).fill = GREEN

    # Row 79 — Venus joint SALE (EPGC share)
    paint(ws["A79"], "Bronze head of a goddess, likely Venus (joint; EPGC share)")
    paint(ws["B79"], "Fortuna / Erdal Dere")
    paint(ws["C79"], 20000, num=ACCT)
    paint(ws["D79"], datetime(2025, 5, 9))
    ws["D79"].number_format = "YYYY-MM-DD"
    paint(ws["E79"], "Erdal Dere")
    paint(ws["F79"], 26000, num=ACCT)
    paint(ws["G79"], datetime(2025, 7, 24))
    ws["G79"].number_format = "YYYY-MM-DD"
    paint(ws["H79"], "=F79-C79")
    ws["H79"].number_format = ACCT
    paint(ws["I79"], VENUS_NOTE)
    ws["I79"].alignment = WRAP
    paint(ws["K79"], "SALE (joint, EPGC share)")
    ws.row_dimensions[79].height = 60
    for col in range(1, 12):
        ws.cell(79, col).font = CG
        if ws.cell(79, col).fill.fill_type in (None, "none"):
            ws.cell(79, col).fill = GREEN

    # Total row should now be 80.
    total_row = None
    for r in range(78, 90):
        if str(ws.cell(r, 1).value or "") == "Total":
            total_row = r
            break
    if total_row is None:
        total_row = 80
        ws.cell(80, 1, "Total")
    paint(ws.cell(total_row, 1), "Total", fill=PEACH)
    ws.cell(total_row, 8).value = "=SUM(H51:H75)+H78+H79"
    ws.cell(total_row, 8).number_format = ACCT
    ws.cell(total_row, 8).fill = PEACH
    ws.cell(total_row, 8).font = CG_B

    note_row = total_row + 1
    for rng in [str(r) for r in ws.merged_cells.ranges if r.min_row == note_row]:
        try:
            ws.unmerge_cells(rng)
        except Exception:
            pass
    ws.cell(note_row, 1, ART_FOOTER)
    ws.cell(note_row, 1).fill = GREEN
    ws.cell(note_row, 1).font = CG
    ws.cell(note_row, 1).alignment = WRAP
    try:
        ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=9)
    except Exception:
        pass
    ws.row_dimensions[note_row].height = 64

    # Remove Venus + Belt from unsold parked inventory.
    for r in range(1, (ws.max_row or 1) + 1):
        a = str(ws.cell(r, 1).value or "")
        if a.startswith("2025 Mercury Fortuna inventory") or a.startswith("2025 Mercury Fortuna —"):
            paint(
                ws.cell(r, 1),
                "2025 Mercury Fortuna — Venus + Belt SOLD 2025 (see Sales 2025 rows 78–79); Aug $20k unnamed still ASK",
                fill=GREEN,
            )
            ws.cell(r, 1).font = CG_B
            ws.cell(r, 1).alignment = WRAP
        if "likely Venus — Fortuna inventory" in a or a.startswith("Bronze head of a goddess, likely Venus — Fortuna"):
            paint(ws.cell(r, 1), "SOLD 2025 — Venus joint (see Sales 2025 row 79). Not unsold inventory.")
            paint(ws.cell(r, 3), None)
            paint(ws.cell(r, 9), "Moved to sold lots. EPGC proceeds $26,000 / Cost $20,000 / net $6,000.")
        if a.startswith("Roman Gold Belt — Fortuna inventory"):
            paint(ws.cell(r, 1), "SOLD 2025 — Roman Gold Belt to L5 (see Sales 2025 row 78). Not unsold inventory.")
            paint(ws.cell(r, 3), None)
            paint(ws.cell(r, 9), "Moved to sold lots. L5 $50,000 / Cost $45,000 / net $5,000. Remaining L5 not attached.")
        if "object still unnamed" in a:
            for rng in [str(x) for x in ws.merged_cells.ranges if x.min_row == r]:
                try:
                    ws.unmerge_cells(rng)
                except Exception:
                    pass
            paint(ws.cell(r, 1), a, fill=YELLOW)
            paint(ws.cell(r, 2), "Mercury 8291", fill=YELLOW)
            paint(ws.cell(r, 3), 20000, fill=YELLOW, num=ACCT)
            paint(ws.cell(r, 4), "2025-08-15", fill=YELLOW)
            paint(
                ws.cell(r, 9),
                "ASK: same $20,000 as Mercury native 2025-08-18. No object name. Not Venus. Not belt. Do not add to sold Cost.",
                fill=YELLOW,
            )
        if a.startswith("Fortuna unnamed remaining 2025"):
            paint(
                ws.cell(r, 1),
                "Fortuna unnamed remaining 2025 (excl. seals $13,000 sold + Venus $20,000 sold + Belt $45,000 sold)",
                fill=YELLOW,
            )
            paint(ws.cell(r, 3), 20000, fill=YELLOW, num=ACCT)
        if a == "Newstar 2025 total":
            ws.cell(r, 3).value = f"=SUM(C{r-3}:C{r-1})"
            ws.cell(r, 3).number_format = ACCT
            ws.cell(r, 3).fill = GREEN


def patch_income(ws) -> None:
    ws["I8"] = I8
    ws["I8"].number_format = ACCT
    ws["I8"].fill = GREEN
    ws["I8"].font = CG
    ws["A12"] = INCOME_NOTE
    ws["A12"].alignment = WRAP
    ws["A12"].font = CG


def patch_epgc(ws) -> None:
    # H = July = column 8, row 53 Art Sales
    ws["H53"] = EPGC_JULY_GROSS
    ws["H53"].number_format = ACCT
    ws["H53"].fill = GREEN
    ws["H53"].font = CG
    # Keep Jan/Feb matched peach; do not zero them.
    ws["B53"].fill = PEACH
    ws["C53"].fill = PEACH
    note_row = None
    for r in range(70, 80):
        v = str(ws.cell(r, 1).value or "")
        if v.startswith("2025 Mercury"):
            note_row = r
            break
    if note_row is None:
        note_row = 75
    ws.cell(note_row, 1, EPGC_NOTE)
    ws.cell(note_row, 1).alignment = WRAP
    ws.cell(note_row, 1).font = CG
    ws.row_dimensions[note_row].height = 72


def patch_data_sources(ws) -> None:
    for r in range(1, (ws.max_row or 1) + 1):
        a = str(ws.cell(r, 1).value or "")
        b = str(ws.cell(r, 2).value or "")
        if a == "Income" and "I8" in b:
            ws.cell(r, 3, "34,055.06")
            ws.cell(
                r,
                4,
                "Prior MATCHED $23,055.06 + Belt net $5,000 + Venus EPGC-share net $6,000. "
                "Sale 6428 and Aquinas wait on Cost.",
            )
            ws.cell(r, 5, "LOCKED green")
            for c in range(1, 6):
                ws.cell(r, c).fill = GREEN
        if a == "EPGC LLC" and b == "Art Sales":
            ws.cell(r, 3, "14,000 Jan + 136,000 Feb + 76,000 July = 226,000")
            ws.cell(
                r,
                4,
                "Gross sale cash. July = Belt L5 $50,000 + Venus remittance $26,000. "
                "Do not dump EOEB $130,270 or other L5 $37,396.32.",
            )
            ws.cell(r, 5, "LOCKED")
            for c in range(1, 6):
                ws.cell(r, c).fill = GREEN
        if a == "Art Sales tab":
            ws.cell(r, 2, "Seals / lots / mosaics / Belt / Venus")
            ws.cell(r, 3, "rows 70–79")
            ws.cell(
                r,
                4,
                "2022–2024 lots stay. 2025 adds Roman Gold Belt SALE and Venus joint SALE (EPGC share).",
            )
            ws.cell(r, 5, "LOCKED")
        if a.startswith("Dump EOEB"):
            ws.cell(
                r,
                1,
                "Dump EOEB $130,270 or other L5 $37,396.32 onto P&L",
            )
            ws.cell(
                r,
                2,
                "Mixed art + consultant. Invoice split still ASK. Belt $50k is locked SALE; remaining L5 is not.",
            )
        if a.startswith("EPGC Art Sales vs"):
            ws.cell(r, 1, "EPGC Art Sales vs $226,000")
            ws.cell(r, 3, 226000)

    # Deal Classifier block (no extra tab — last-year 11 tabs + Data sources only).
    start = (ws.max_row or 1) + 2
    ws.cell(start, 1, "DEAL CLASSIFIER (conceptual)").font = CG_B
    ws.cell(start, 1).fill = BLUE
    headers = ["Classifier", "Meaning", "2025 locked example", "Still ASK / unallocated", ""]
    for i, h in enumerate(headers, 1):
        ws.cell(start + 1, i, h).font = CG_B
        ws.cell(start + 1, i).fill = BLUE
    rows = [
        ("SALE", "Art object sold; Art Sales register + EPGC Art Sales gross cash; I8 gets NET if Cost known", "Roman Gold Belt L5 $50,000 / Cost $45,000", ""),
        ("SALE (joint, EPGC share)", "Joint purchase; register Sale Price = EPGC proceeds, not full hammer", "Venus Erdal remittance $26,000 / EPGC Cost $20,000", ""),
        ("FEE", "Advisory / consultant income or expertise expense — never I8", "David Aaron June $13,595 Consultant; Wise write-ups $992.80 Fees", ""),
        ("PASS", "In and out, not P&L", "Koziol $150,000 10/20 → Ariadne $150,000 10/21", ""),
        ("INVENTORY", "Named or unnamed purchase still unsold", "Fortuna Aug unnamed $20,000", ""),
        ("UNALLOCATED", "Mixed pattern confirmed; invoice-level split still ASK; do not dump", f"EOEB remainder ${EOEB_REMAINING:,.2f}; other L5 ${L5_REMAINING:,.2f}; other Erdal IN ${ERDAL_REMAINING:,.0f}; Aysel $50,000", ""),
    ]
    for i, row in enumerate(rows):
        fill = GREEN if i < 4 else YELLOW
        for c, val in enumerate(row, 1):
            ws.cell(start + 2 + i, c, val).fill = fill
            ws.cell(start + 2 + i, c).font = CG
            ws.cell(start + 2 + i, c).alignment = WRAP
        ws.row_dimensions[start + 2 + i].height = 32


def sheet_to_csv(ws, path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        max_c = min(ws.max_column or 1, 14)
        for r in range(1, (ws.max_row or 1) + 1):
            w.writerow([("" if ws.cell(r, c).value is None else ws.cell(r, c).value) for c in range(1, max_c + 1)])


def classify_ledger_row(date, counterparty, amount, bucket, status, deal) -> tuple[str, str, str, str, str]:
    """Return classifier, new_bucket, new_status, new_deal, extra_note."""
    amt = float(amount)
    if date == "2025-07-15" and abs(amt - 50000) < 0.01:
        return (
            "SALE",
            "ART_SALE",
            "LOCKED",
            "Roman Gold Belt — Sale Price (L5)",
            "User 2026-09-21: sold to L5. Only L5 payment large enough. "
            "Hits EPGC Art Sales July +$50,000. Do NOT attach remaining L5 invoices. "
            "L5 paid 6 days before Fortuna Cost — back-to-back dealer cash. Not tax advice.",
        )
    if date == "2025-07-21" and abs(amt + 45000) < 0.01:
        return (
            "SALE",
            "ART_PURCHASE",
            "LOCKED",
            "Roman Gold Belt — Cost (Fortuna)",
            "User 2026-09-21: Cost of SALE to L5, not unsold inventory. Native memo Roman Gold Belt. "
            "Paired with L5 7/15 +$50,000. Net $5,000 on I8. Not tax advice.",
        )
    if date == "2025-05-09" and abs(amt + 20000) < 0.01:
        return (
            "SALE (joint, EPGC share)",
            "ART_PURCHASE",
            "LOCKED",
            "Venus — EPGC Cost (joint with Erdal)",
            "User 2026-09-21: Erdal and Jacob bought together. Not unsold inventory. "
            "Paired with Erdal 7/24 +$26,000 remittance. I8 net $6,000. Not tax advice.",
        )
    if date == "2025-07-24" and abs(amt - 26000) < 0.01:
        return (
            "SALE (joint, EPGC share)",
            "ART_SALE",
            "LOCKED",
            "Venus — EPGC proceeds (Erdal remittance)",
            "User 2026-09-21: Erdal sold it and paid back Jacob’s $20,000 + $6,000 earnings. "
            "Art Sales Sale Price = EPGC proceeds $26,000, not full hammer. "
            "Hits EPGC Art Sales July. Other Erdal IN still ASK. Not tax advice.",
        )
    if "David Aaron" in deal or (date == "2025-06-24" and abs(amt - 13595) < 0.01):
        return ("FEE", bucket, status, deal, "")
    if "expertise" in deal.lower() or (bucket == "EPGC_EXPENSE" and "Wise" in counterparty):
        return ("FEE", bucket, status, deal, "")
    if bucket == "PASS_THROUGH":
        return ("PASS", bucket, status, deal, "")
    if date == "2025-08-15" and abs(amt + 20000) < 0.01:
        return ("INVENTORY", bucket, "ASK", deal, "Still unnamed. Not Venus, not belt.")
    if counterparty == "L5":
        return (
            "UNALLOCATED",
            "ADVISORY",
            "ASK",
            "L5 — mixed unallocated after Belt $50k SALE removed",
            f"Remaining L5 ${L5_REMAINING:,.2f} still ASK (7/11 $8,534.79; 7/15 $2,500; "
            "9/2 $5,042; 9/23 $9,119.27; 10/21 $7,500; 12/4 $4,700.26). Do not dump. Not tax advice.",
        )
    if "EOEB" in counterparty and status == "ASK":
        return ("UNALLOCATED", bucket, status, deal, "")
    if date in {"2025-10-09", "2025-11-07"} and "ERDAL" in counterparty.upper():
        return ("UNALLOCATED", bucket, "ASK", deal, "Still ASK. Not the Venus remittance.")
    if "Aysel" in counterparty:
        return ("UNALLOCATED", bucket, "ASK", deal, "Not the belt. Monarch IN vs native Failed OUT still ASK.")
    if bucket in {"ART_SALE", "ART_PURCHASE"} and status in {"LOCKED", "MATCHED"}:
        return ("SALE", bucket, status, deal, "")
    if bucket in {"TRANSFER", "TEST", "INVESTMENT", "COGS_JEWELRY", "REIMBURSE"}:
        return ("n/a", bucket, status, deal, "")
    return ("UNALLOCATED" if status == "ASK" else "n/a", bucket, status, deal, "")


def patch_mercury_ledger() -> None:
    path = CSV_DIR / "MERCURY_LEDGER.csv"
    rows = list(csv.reader(path.open(encoding="utf-8")))
    banner = (
        "Mercury Checking 8291 — EPGC LLC 2025. 64 unique cash txns. "
        "Deal Classifier: SALE vs FEE vs PASS vs INVENTORY vs UNALLOCATED. "
        "Roman Gold Belt = SALE (Fortuna Cost $45,000 / L5 Sale $50,000). "
        "Venus = joint SALE EPGC share (Cost $20,000 / proceeds $26,000). "
        "Remaining EOEB $130,270 and other L5 $37,396.32 = UNALLOCATED. "
        "David Aaron = FEE. Koziol/Ariadne = PASS. Fortuna Aug $20k = INVENTORY ASK. "
        "Not tax advice."
    )
    rows[0] = [banner] + [""] * 12
    header = rows[2]
    # Insert Classifier after Bucket if missing.
    if "Classifier" not in header:
        header.insert(5, "Classifier")
        for r in rows[3:]:
            if len(r) >= 5:
                r.insert(5, "")
    idx = {name: i for i, name in enumerate(header)}
    cls_i = idx["Classifier"]
    bkt_i = idx["Bucket"]
    st_i = idx["Status"]
    deal_i = idx["Deal / object"]
    hits_i = idx["Hits EPGC Art Sales?"]
    notes_i = idx["Notes"]
    date_i = idx["Date"]
    who_i = idx["Counterparty"]
    amt_i = idx["Amount"]
    for r in rows[3:]:
        if len(r) < notes_i + 1 or not r[date_i]:
            continue
        cls, bkt, st, deal, extra = classify_ledger_row(
            r[date_i], r[who_i], r[amt_i], r[bkt_i], r[st_i], r[deal_i]
        )
        r[cls_i] = cls
        r[bkt_i] = bkt
        r[st_i] = st
        r[deal_i] = deal
        if extra:
            r[notes_i] = extra
        if cls.startswith("SALE") and float(r[amt_i]) > 0 and bkt == "ART_SALE" and st == "LOCKED":
            r[hits_i] = "YES"
        if "Roman Gold Belt — Cost" in deal or "Venus — EPGC Cost" in deal:
            r[hits_i] = "no (Cost of SALE)"
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def write_ask_mercury() -> None:
    path = CSV_DIR / "ASK_Mercury.csv"
    banner = (
        "ASK — remaining Mercury 8291 questions after Belt + Venus SALE lock 2026-09-21. "
        "Q6 consultant, Q7 Wise expertise, Q8 pass-through, Q9 Aquinas books LOCKED (green). "
        "Belt = SALE to L5. Venus = joint SALE (EPGC share). "
        "Q1 EOEB remainder and remaining L5 MIXED PATTERN CONFIRMED (blue) — dollars unallocated. "
        "Not tax advice."
    )
    rows = [
        [banner, "", ""],
        ["", "", ""],
        ["Question", "What I need", "Classifier"],
        [
            "1. EOEB LLC remainder after mosaics — MIXED PATTERN CONFIRMED, dollars unallocated",
            f"${EOEB_REMAINING:,.2f} of EOEB IN is not mosaics $105,000 and not Wise reimburse $565.75. "
            "User: EOEB is mixed art sales AND advisory/consultant fees. Do NOT dump this remainder "
            "onto Consultant or Art Sales. Mark each invoice sale vs advisory: 1/16 $1,000 "
            "(Canosan-horse STILL OPEN — Art Sales horse is to Erdal, this cash is EOEB); "
            "3/10 $14,000; 5/1 $21,000; 5/13 $16,200; 6/12 $25,000 + $3,300; 7/8 $2,000; "
            "8/29 $10,000; 9/19 $5,000; 9/25 $3,100; 10/1 $10,000; 10/15 $16,000; "
            "10/24 $1,500; 10/31 $2,170.",
            "UNALLOCATED",
        ],
        [
            "2. L5 remainder after Belt $50,000 SALE — MIXED PATTERN CONFIRMED, dollars unallocated",
            f"${L5_REMAINING:,.2f} remaining after locking 7/15 $50,000 as Roman Gold Belt SALE. "
            "Do NOT attach these invoices to the belt. Do NOT dump onto Consultant or Art Sales. "
            "Mark each remaining invoice sale vs advisory: 7/11 $8,534.79; 7/15 $2,500; "
            "9/2 $5,042; 9/23 $9,119.27; 10/21 $7,500; 12/4 $4,700.26.",
            "UNALLOCATED",
        ],
        [
            "3. Fortuna / Erdal $20,000 Aug — object still unnamed",
            "$20,000.00. Monarch posted 8/15 = Mercury native initiated 8/18 (same cash, not a second unique txn). "
            "Native memo has no object name. Not Venus (5/9, now SOLD). Not Roman Gold Belt (7/21, now SOLD).",
            "INVENTORY",
        ],
        [
            "4. Other Erdal Dere IN after Venus remittance",
            f"${ERDAL_REMAINING:,.2f} remaining (10/9 $25,000 and 11/7 $6,000 ‘FORTUNA PAYMENT’). "
            "7/24 $26,000 is LOCKED Venus remittance (investment + earnings). These two still ASK.",
            "UNALLOCATED",
        ],
        [
            "5. Aysel Dere $50,000 — Monarch IN vs native Failed OUT",
            "Monarch 7/18 +$50,000 IN is in the 64. Mercury native CSV has NO Sent Aysel cash — only Failed "
            "$50,000 OUT 7/17 'Recipient account does not exist'. Is the Monarch IN a phantom of the failed send, "
            "or a real incoming the export omitted? Not the belt. Do not dump until confirmed.",
            "UNALLOCATED",
        ],
        [
            "6. David Aaron Limited $13,595 IN (6/24) — LOCKED consultant fee",
            "CONFIRMED: consultant fee, not a sale. Booked EPGC Consultant June $13,595. Not Art Sales. Not I8.",
            "FEE",
        ],
        [
            "7. Wise $334.17 (7/9) and $658.63 (10/30) — LOCKED expertise write-ups",
            "CONFIRMED: business expenses (expertise write-ups). EPGC Consultant Fees July $334.17 + October "
            "$658.63 = $992.80. Distinct from Dec Wise $565.75 reimbursed by EOEB (still not P&L).",
            "FEE",
        ],
        [
            "8. Jack Koziol & Tracy Hoffman $150,000 IN (10/20) / Ariadne Demirjian $150,000 OUT (10/21) — LOCKED pass-through",
            "CONFIRMED: pass-through, not P&L. Not a sale, not a purchase, not COGS, not Consultant.",
            "PASS",
        ],
        [
            "9. Aquinas Hobor $1,000 IN (2/4) — LOCKED book sale",
            "CONFIRMED: book SALE. Hits EPGC Art Sales February. Cost / titles TBD (not in I8 until Cost).",
            "SALE",
        ],
        [
            "10. Mosaics Cost $90,000 vs Mercury $50,000",
            "Plutus & Mnemosyne paid $50,000 in February. Where is the other $40,000 of Cost "
            "(other account / 2024 / still owed)?",
            "ASK",
        ],
        ["", "", ""],
        ["LOCKED / MATCHED this pass (do not recast without saying so)", "", "Classifier"],
        ["Roman Gold Belt", "Fortuna 7/21 −$45,000 Cost / L5 7/15 +$50,000 Sale. Net $5,000. Remaining L5 not attached.", "SALE"],
        ["Venus (bronze head, joint)", "Fortuna 5/9 −$20,000 EPGC Cost / Erdal 7/24 +$26,000 proceeds (investment + $6,000 earnings). Sale Price = EPGC proceeds, not full hammer.", "SALE (joint, EPGC share)"],
        ["Coinbase ACH net", "$-51,000.00 (OUT $-59,000.00 / IN $8,000.00) — Investments tab, not EPGC", "n/a"],
        ["Newstar Jewelers", "$23,581.00 jewelry fabrication COGS", "n/a"],
        ["Berk sales", "$14,000 seals (1/10) + $30,000 lots (2/7) = $44,000 Art Sales", "SALE"],
        ["Mosaics", "EOEB $105,000 sale (2/21) / Plutus $50,000 of $90,000 Cost", "SALE"],
        ["BoA 9922 draws", "$261,172.00 owner transfer — not P&L", "n/a"],
        ["Wise/EOEB 12/2", "$565.75 reimbursed — not income, not expense", "n/a"],
        ["David Aaron Limited", "$13,595 (6/24) consultant fee → EPGC Consultant June. Not a sale.", "FEE"],
        ["Koziol / Ariadne", "$150,000 IN 10/20 + $150,000 OUT 10/21 LOCKED pass-through — not P&L", "PASS"],
        ["Wise expertise write-ups", "$334.17 (7/9) + $658.63 (10/30) = $992.80 → EPGC Consultant Fees", "FEE"],
        ["Aquinas Hobor books", "$1,000 IN 2/4 LOCKED book sale → EPGC Art Sales February. Cost TBD (not in I8).", "SALE"],
        ["EOEB / remaining L5 mixed pattern", f"CONFIRMED mixed art + advisory. Remainder ${EOEB_REMAINING:,.2f} / ${L5_REMAINING:,.2f} unallocated.", "UNALLOCATED"],
        ["Native CSV 2026-09-20", "epgc-llc-transactions-2025-jan-01-to-2025-dec-31.csv is account 8291. 0 new unique Sent cash vs the 64.", "n/a"],
    ]
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def write_fortuna() -> None:
    path = CSV_DIR / "FORTUNA_INVENTORY.csv"
    rows = [
        ["Object", "Source", "Cost", "Date", "Status", "Classifier", "Notes"],
        [
            "Bronze head of a goddess likely Venus",
            "Mercury 8291 Fortuna",
            "20000.00",
            "2025-05-09",
            "LOCKED SOLD 2025",
            "SALE (joint, EPGC share)",
            "Removed from unsold parked. Erdal and Jacob bought together. Erdal sold it and remitted 7/24 +$26,000 = $20,000 + $6,000 earnings. Sale Price = EPGC proceeds, not full hammer.",
        ],
        [
            "Roman Gold Belt",
            "Mercury 8291 Fortuna",
            "45000.00",
            "2025-07-21",
            "LOCKED SOLD 2025",
            "SALE",
            "Removed from unsold parked. Sold to L5 7/15 +$50,000 (paid 6 days before Cost). Net $5,000. Do not attach remaining L5 invoices.",
        ],
        [
            "Fortuna unnamed",
            "Mercury 8291 Fortuna",
            "20000.00",
            "2025-08-15",
            "ASK",
            "INVENTORY",
            "Monarch 8/15 = native 8/18. No object name. Same cash not a second unique txn. Not Venus. Not belt.",
        ],
    ]
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def write_deal_classifier() -> None:
    path = CSV_DIR / "DEAL_CLASSIFIER.csv"
    rows = [
        [
            "Deal Classifier — Mercury 8291 / EPGC 2025. Classes: SALE, FEE, PASS, INVENTORY, UNALLOCATED. "
            "Last-year model: Art Sales = gross sale cash by month; Income I8 = art NET for lots with Cost; "
            "Consultant never goes to I8. Not tax advice.",
            "",
            "",
            "",
            "",
            "",
        ],
        ["Classifier", "Counterparty / object", "Date", "Amount", "Hits", "Notes"],
        ["SALE", "Roman Gold Belt / L5", "2025-07-15", "50000.00", "EPGC Art Sales July; I8 net $5,000", "Only L5 payment attached to the belt."],
        ["SALE", "Roman Gold Belt / Fortuna", "2025-07-21", "-45000.00", "Art Sales Cost", "Bought from Fortuna, sold to L5. Not unsold inventory."],
        ["SALE (joint, EPGC share)", "Venus / Erdal remittance", "2025-07-24", "26000.00", "EPGC Art Sales July; I8 net $6,000", "Sale Price = EPGC proceeds, not full hammer."],
        ["SALE (joint, EPGC share)", "Venus / Fortuna", "2025-05-09", "-20000.00", "Art Sales Cost (EPGC stake)", "Erdal and Jacob bought together."],
        ["SALE", "Berk seals", "2025-01-10", "14000.00", "EPGC Art Sales January; I8 net $1,000", "Cost Fortuna 1/8 $13,000."],
        ["SALE", "Berk lots lump", "2025-02-07", "30000.00", "EPGC Art Sales February; I8 includes lots net", "F71 lump."],
        ["SALE", "Three Ancient Mosaics / EOEB", "2025-02-21", "105000.00", "EPGC Art Sales February; I8 net $15,000", "Cost $90,000; $40,000 Cost missing on 8291."],
        ["SALE", "Aquinas Hobor books", "2025-02-04", "1000.00", "EPGC Art Sales February; not I8 until Cost", "Cost / titles TBD."],
        ["FEE", "David Aaron Limited", "2025-06-24", "13595.00", "EPGC Consultant June only", "Not Art Sales. Not I8."],
        ["FEE", "Wise expertise write-ups", "2025-07-09 / 2025-10-30", "-992.80", "EPGC Consultant Fees", "Dec $565.75 is reimbursed, not this."],
        ["PASS", "Koziol → Ariadne", "2025-10-20 / 2025-10-21", "150000.00", "Not P&L", "Locked pass-through."],
        ["INVENTORY", "Fortuna unnamed", "2025-08-15", "-20000.00", "Parked Cost ASK", "Not Venus. Not belt."],
        ["UNALLOCATED", "EOEB remainder", "various", f"{EOEB_REMAINING:.2f}", "none until invoice split", "Mixed art + advisory. Do not dump."],
        ["UNALLOCATED", "L5 remainder after belt", "various", f"{L5_REMAINING:.2f}", "none until invoice split", "7/11 $8,534.79; 7/15 $2,500; 9/2 $5,042; 9/23 $9,119.27; 10/21 $7,500; 12/4 $4,700.26."],
        ["UNALLOCATED", "Other Erdal IN", "2025-10-09 / 2025-11-07", "31000.00", "none", "10/9 $25,000 and 11/7 $6,000 FORTUNA PAYMENT. Not Venus."],
        ["UNALLOCATED", "Aysel Dere Monarch IN", "2025-07-18", "50000.00", "none", "Native Failed OUT 7/17. Not the belt."],
    ]
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def write_lock_q() -> None:
    path = CSV_DIR / "LOCK_Mercury_Q6_Q9.csv"
    rows = [
        ["Q", "Status", "What", "Where it hits"],
        ["6", "LOCKED", "David Aaron Limited $13,595 IN (6/24) consultant fee", "EPGC Consultant June $13,595. Not Art Sales. Not I8. Classifier: FEE."],
        ["7", "LOCKED", "Wise $334.17 (7/9) + $658.63 (10/30) expertise write-ups", "EPGC Consultant Fees $992.80. Classifier: FEE."],
        ["8", "LOCKED", "Koziol $150,000 IN (10/20) / Ariadne $150,000 OUT (10/21)", "Pass-through — not P&L. Classifier: PASS."],
        ["9", "LOCKED", "Aquinas Hobor $1,000 IN (2/4) book SALE", "EPGC Art Sales February. Cost TBD. Classifier: SALE."],
        ["Belt", "LOCKED", "Roman Gold Belt Fortuna $45,000 / L5 $50,000", "EPGC Art Sales July +$50,000. I8 net +$5,000. Classifier: SALE. Remaining L5 not attached."],
        ["Venus", "LOCKED", "Venus Fortuna $20,000 / Erdal remittance $26,000", "EPGC Art Sales July +$26,000. I8 net +$6,000. Sale Price = EPGC proceeds, not full hammer. Classifier: SALE (joint, EPGC share)."],
        ["1", "PATTERN CONFIRMED / dollars unallocated", f"EOEB remainder ${EOEB_REMAINING:,.2f} mixed art sales + advisory/consultant", "Do NOT dump. Classifier: UNALLOCATED. Jan 16 $1,000 Canosan-horse still open."],
        ["2", "PATTERN CONFIRMED / dollars unallocated", f"L5 remainder ${L5_REMAINING:,.2f} mixed after Belt $50k SALE", "Do NOT dump. Classifier: UNALLOCATED."],
        ["ASK remaining", "ASK", f"Fortuna unnamed $20,000 / other Erdal IN ${ERDAL_REMAINING:,.0f} / Aysel Monarch IN vs native Failed OUT / mosaics Cost $40,000", "Aug $20k object name (INVENTORY). Other Erdal not Venus. Aysel not the belt."],
        ["LOCKED totals", "LOCKED", f"Art Sales gross $226,000 (Jan 14k + Feb 136k + July 76k); I8 $34,055.06; Consultant June $13,595", "Dec Wise $565.75 REIMBURSE. EOEB remainder / other L5 not on P&L."],
    ]
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


START_HERE_MERCURY = """START HERE — Mercury Bank 8291 (2026-09-21 Belt + Venus SALE lock)

Not tax advice. Same Personal Income.xlsx tabs as last year. Organizational packet only.

Mercury Checking 8291 is the EPGC LLC operating account (Choice Financial).
64 unique 2025 cash transactions from Monarch IDs.

User 2026-09-21 LOCKED two deals that were parked as unsold inventory:

1. Roman Gold Belt — SALE to L5 (not unsold, not consultant, not pass-through).
   Cost: Mercury 8291 Fortuna/Erdal Dere 2025-07-21 −$45,000 (native memo: Roman Gold Belt).
   Working sale: L5 2025-07-15 +$50,000 (only L5 payment large enough). Net $5,000.
   L5 paid 6 days before Fortuna Cost — back-to-back dealer cash.
   Do NOT attach remaining L5 invoices to the belt.

2. Venus (bronze head of a goddess) — joint SALE, EPGC share (not unsold, not consultant).
   Erdal and Jacob bought together. Erdal sold it and paid back Jacob’s investment plus earnings.
   EPGC Cost: Fortuna 2025-05-09 −$20,000.
   Working remittance: Erdal Dere IN 2025-07-24 +$26,000 = $20,000 + $6,000 earnings.
   Art Sales Sale Price = EPGC proceeds ($26,000), not the object’s full hammer.

Open ASK (remainder still unallocated):
https://docs.google.com/spreadsheets/d/ASK_PLACEHOLDER/edit

EPGC LLC 2025 (Art Sales cash $226,000 = Jan $14,000 / Feb $136,000 / July $76,000; Consultant June $13,595; Consultant Fees $992.80):
https://docs.google.com/spreadsheets/d/EPGC_PLACEHOLDER/edit

One spreadsheet like last year (11 tabs, 2023|2024|2025, calculating formulas, Art Sales includes 2022–2024 lots + Belt + Venus):
https://docs.google.com/spreadsheets/d/XLSX_PLACEHOLDER/edit

Deal Classifier: SALE vs FEE vs PASS vs INVENTORY vs UNALLOCATED.
Belt = SALE. Venus = SALE (joint, EPGC share). Remaining EOEB/L5 = UNALLOCATED.
David Aaron = FEE. Koziol/Ariadne = PASS. Fortuna Aug unnamed = INVENTORY ASK.

LOCKED this pass
- Roman Gold Belt SALE: Fortuna Cost $45,000 / L5 Sale $50,000 / net $5,000 → EPGC Art Sales July + I8.
- Venus joint SALE: Fortuna Cost $20,000 / Erdal proceeds $26,000 / net $6,000 → EPGC Art Sales July + I8.
- Coinbase ACH net $-51,000.00 → Investments. Not EPGC.
- Newstar Jewelers $23,581.00 → jewelry fabrication COGS.
- Art sales matched: Berk seals $14,000 (1/10) + Berk lots $30,000 (2/7) + mosaics $105,000 (EOEB 2/21) + Aquinas books $1,000 (2/4).
- Seals Cost $13,000 = Fortuna 1/8. Mosaics Cost $90,000 of which Plutus paid $50,000 ($40,000 still ASK).
- BoA 9922 transfers $261,172.00 = owner draws, not P&L.
- David Aaron Limited $13,595 (6/24) → EPGC Consultant June. Consultant fee, not a sale.
- Koziol $150,000 (10/20) / Ariadne $150,000 (10/21) → LOCKED pass-through. Not P&L.
- Wise $334.17 (7/9) + $658.63 (10/30) → EPGC Consultant Fees (expertise write-ups).
- Aquinas Hobor $1,000 (2/4) → books sold. EPGC Art Sales February. Cost TBD (not in I8 until Cost).

EPGC 2025 Art Sales MATCHED/LOCKED cash $226,000 (Jan $14,000 / Feb $136,000 / July $76,000).
EPGC Consultant June is $13,595 (David Aaron only).
Income I8 Art net is $34,055.06 (prior $23,055.06 + Belt $5,000 + Venus $6,000).

ASK remaining — do not dump
1. EOEB remainder $130,270.00 — MIXED PATTERN CONFIRMED; invoice-level split ASK. Jan 16 $1,000 Canosan-horse still open.
2. Other L5 $37,396.32 after removing $50k belt — MIXED PATTERN CONFIRMED. 7/11 $8,534.79; 7/15 $2,500; 9/2 $5,042; 9/23 $9,119.27; 10/21 $7,500; 12/4 $4,700.26.
3. Fortuna unnamed $20,000.00 — Aug $20,000 (Monarch 8/15 = native 8/18). Not Venus, not belt.
4. Other Erdal IN $31,000.00 — 10/9 $25,000 and 11/7 $6,000 “FORTUNA PAYMENT”. Not the Venus remittance.
5. Aysel Dere — Monarch 7/18 +$50,000 IN vs native Failed 7/17 −$50,000 OUT. Not the belt.
10. Mosaics Cost $40,000 missing on Mercury 8291.

Sale 6428 $11,000 and Aquinas books $1,000 wait on Cost (not in I8).
"""

START_HERE_ALL = """START HERE — last year's spreadsheet, 2025 numbers filled
Not tax advice. Organizational packet only.

Open this ONE spreadsheet (same 11 tab names as last year, 2023 | 2024 | 2025 year blocks, 2025 numbers, formulas that calculate, Art Sales has 2022–2024 lots + 2025 Belt and Venus):
https://docs.google.com/spreadsheets/d/XLSX_PLACEHOLDER/edit

START HERE doc:
https://docs.google.com/document/d/DOC_PLACEHOLDER/edit

Copies
07 Income folder:
https://docs.google.com/spreadsheets/d/XLSX07_PLACEHOLDER/edit
My Drive / 2025 Taxes:
https://docs.google.com/spreadsheets/d/XLSXROOT_PLACEHOLDER/edit

Tabs in last year's order (all inside that one file)
1. Income — 2022–2025. 2025 Actual: Hindman $70,618 / 216 $19,500 / 524 #2 $17,086.94 / GCM $21,500 / Art $34,055.06. I10 EBITDA =SUM(I4:I9) → $162,760.
2. 524 Ferdinand Ave, Unit 2 — 2023 YR TOTAL | 2024 YR TOTAL | 2025 YR TOTAL. 2025 STR $13,879.55. Grove Collaborative $791.21.
3. 216 N. Oak Park Ave — 2023 | 2024 | 2025. Lemonade $42.84/mo = $514.08. Rent $19,500. Mortgage $14,617.54. HOA $5,058.36. Brennan $150 / Ace $11 / Joan $675 / Imelda $300.
4. EPGC LLC — 2023, then 2024, then 2025 (Art Sales $14,000 + $136,000 + July $76,000 = $226,000; Consultant June $13,595). 2025 Total is a live SUM.
5. Refrence Library
6. Art Sales and Purchases — 2022–2024 lots, then Sales 2025 including Roman Gold Belt SALE to L5 and Venus joint SALE (EPGC share). Do not dump EOEB $130,270 or other L5 $37,396.32.
7. GCM — 2022–2025 stacked like last year. 2025 board $4,100 / $9,200 / $4,100 / $4,100 = $21,500.
8. Hindman W2
9. Megan T4
10. Investments — Coinbase ACH net -$51,000.
11. 524 Ferdinand Ave, Unit 1 — 2023 YR TOTAL | 2025 YR TOTAL (same as last year; no 2024 rental block).

Deal Classifier (on 2025 Data sources + DEAL_CLASSIFIER.csv): SALE vs FEE vs PASS vs INVENTORY vs UNALLOCATED.
Belt = SALE. Venus = SALE (joint, EPGC share). Remaining EOEB/L5 = UNALLOCATED.

Last year's file (2023/2024, do not edit):
https://drive.google.com/file/d/1XqZyLXxYiyrVWvt4TXGd3kxh7G6uAyW2/view

How this stays correct
1. Copy last year's Personal Income.xlsx. Never rebuild from CSVs.
2. Fill the 2025 year block only. 2023 and 2024 stay as they were.
3. python3 .cursor/scratch/assert_personal_income_2025.py must print LOCK OK + SHEETS-SAFE OK.
4. Drive upload is a compact xlsx Google converts (cell refs like AO5 survive). Include Art Sales row 1 and 2022–2024 lots — sparse tiny OOXML dropped Art cells.
5. CSV builders (.cursor/scratch/build_drive_all_tabs.py) now refuse to run.

216 N. Oak Park #1Z (LTR) 2025 — same rows as last year
RENTER'S INSURANCE: Lemonade $42.84/mo × 12. Not an October lump $514.
RENTAL INCOME: $1,950 × 6, then $1,450 / $2,450 / vacant Sep–Nov / $3,900 Dec = $19,500.
MORTGAGE: $1,226.25 Jan–Oct / $1,177.52 Nov–Dec (Jan–Jun yellow WAIT 8507).
HOA: $421.53 × 12 (Jan–May yellow WAIT 8507).
INTERIOR MAINTENANCE: $150 Chris Brennan August; $11 Ace keys December.
INTERIOR REPAIR: $675 Joan December.
MOVE OUT FEE: $300 Imelda October.

524 Ferdinand
Travelers was prior home insurance. Geico was the cars. 2025 State Farm home share only on Units 1 & 2. Auto + Geico credit are personal. Grove Collaborative ≠ 827 N Grove.

Do not open anything titled _SUPERSEDED. Those files stored =SUM as text, split this into 11 separate files, or put 2025 numbers into the 2023 columns.

Not tax advice.
"""


def write_start_here() -> None:
    (CSV_DIR / "START_HERE_mercury.txt").write_text(START_HERE_MERCURY, encoding="utf-8")
    (CSV_DIR / "START_HERE_all_tabs.txt").write_text(START_HERE_ALL, encoding="utf-8")


def patch_json() -> None:
    packet = json.loads(PACKET.read_text(encoding="utf-8"))
    packet["updated"] = "2026-09-21"
    packet["xlsx_note"] = (
        "2026-09-21: Live file is ONE Google Sheet with last year's 11 tab names, "
        "2023|2024|2025 year blocks, 2025 numbers, calculating SUM formulas. "
        "Art Sales includes 2022–2024 lots plus Roman Gold Belt SALE and Venus joint SALE. "
        f"Income I8 $34,055.06. I10 $162,760. EPGC Art Sales July $76,000 (year $226,000). "
        "Not tax advice."
    )
    m = packet.setdefault("mercury_8291", {})
    m["updated"] = "2026-09-21"
    locked = m.setdefault("locked", {})
    locked["art_sale_locked_gross"] = ART_GROSS
    locked["art_net_locked"] = I8
    locked["roman_gold_belt_sale"] = 50000.0
    locked["roman_gold_belt_cost"] = 45000.0
    locked["roman_gold_belt_net"] = 5000.0
    locked["venus_epgc_proceeds"] = 26000.0
    locked["venus_epgc_cost"] = 20000.0
    locked["venus_epgc_net"] = 6000.0
    locked.pop("fortuna_venus_inventory", None)
    locked.pop("fortuna_roman_gold_belt", None)
    ask = m.setdefault("ask", {})
    ask["eoeb_remainder"] = EOEB_REMAINING
    ask["l5"] = L5_REMAINING
    ask["l5_after_belt"] = L5_REMAINING
    ask["l5_note"] = (
        "7/15 $50,000 locked as Roman Gold Belt SALE; remaining "
        "7/11 $8,534.79; 7/15 $2,500; 9/2 $5,042; 9/23 $9,119.27; "
        "10/21 $7,500; 12/4 $4,700.26 still UNALLOCATED"
    )
    ask["fortuna_remainder_out"] = 20000.0
    ask["fortuna_named_inventory"] = 0.0
    ask["erdal_in"] = ERDAL_REMAINING
    ask["erdal_in_note"] = "10/9 $25,000 and 11/7 $6,000 FORTUNA PAYMENT still ASK; 7/24 $26,000 is Venus remittance LOCKED"
    m["epgc_art_sales_2025_monthly"] = [
        14000.0, 136000.0, 0.0, 0.0, 0.0, 0.0, 76000.0, 0.0, 0.0, 0.0, 0.0, 0.0
    ]
    packet["deal_classifier"] = {
        "classes": ["SALE", "FEE", "PASS", "INVENTORY", "UNALLOCATED"],
        "belt": "SALE",
        "venus": "SALE (joint, EPGC share)",
        "david_aaron": "FEE",
        "koziol_ariadne": "PASS",
        "fortuna_aug_unnamed": "INVENTORY",
        "eoeb_remainder": "UNALLOCATED",
        "l5_remainder": "UNALLOCATED",
        "disclaimer": "Not tax advice. Organizational packet only.",
    }
    PACKET.write_text(json.dumps(packet, indent=2) + "\n", encoding="utf-8")

    mercury = json.loads(MERCURY_JSON.read_text(encoding="utf-8"))
    mercury["updated"] = "2026-09-21"
    mercury.setdefault("locked", {}).update(locked)
    mercury.setdefault("ask", {}).update(ask)
    mercury["epgc_art_sales_2025_monthly"] = packet["mercury_8291"]["epgc_art_sales_2025_monthly"]
    MERCURY_JSON.write_text(json.dumps(mercury, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    wb = load_workbook(XLSX)
    patch_art(wb["Art Sales and Purchases"])
    patch_income(wb["Income"])
    patch_epgc(wb["EPGC LLC"])
    patch_data_sources(wb["2025 Data sources"])
    try:
        wb.defined_names.clear()
    except Exception:
        pass
    wb._external_links = []
    wb.save(XLSX)
    sanitize_xlsx(XLSX)
    DELIVERABLE.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(XLSX, DELIVERABLE)

    sheet_to_csv(wb["Art Sales and Purchases"], CSV_DIR / "Art_Sales.csv")
    sheet_to_csv(wb["Income"], CSV_DIR / "Income.csv")
    sheet_to_csv(wb["EPGC LLC"], CSV_DIR / "EPGC_LLC.csv")

    patch_mercury_ledger()
    write_ask_mercury()
    write_fortuna()
    write_deal_classifier()
    write_lock_q()
    write_start_here()
    patch_json()

    print("locked belt+venus")
    print("xlsx", XLSX, XLSX.stat().st_size)
    print("I8", I8, "July art", EPGC_JULY_GROSS, "gross", ART_GROSS)
    print("I10 expected", 70618 + 19500 + 17086.94 + 21500 + I8)
    print("L5 remaining", L5_REMAINING, "Erdal remaining", ERDAL_REMAINING)


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""Apply 2026-09-18 expense answers onto the already-filled Tax Turbo workbook.

User answers (not tax advice):
  1–4  Nicor / Prime AT&T / Village water / TruGreen-Alsip-Good Earth CONFIRMED
  5    March Home Depot → 524 repairs (treated “534” as 524 Ferdinand). Remaining
       Prime HD / Lowe’s / IKEA split evenly by transaction 1/3 524 repairs,
       1/3 personal, 1/3 827 Grove. Extra Space stays HOLD.
  6    ComEd paid checking/ACH or check — amounts kept, yellow TBD; remind Megan
  7    Canva CONFIRMED EPGC
  8    Park Chicago CONFIRMED EPGC Travel
  9    Great Frame Up = inventory framing (Art Sales, not EPGC P&L)
  10   LAMA / Hindman = inventory purchases; Field Museum = EPGC LLC expense

Third pass (State Farm 2026-09-18) is apply_state_farm_2025.py — REPLACES 524
INSURANCE (does not add). Do not re-run this file on an already-filled workbook.
"""
from __future__ import annotations

import csv
import json
import shutil
import sys
from copy import copy
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path("/workspace/.cursor/scratch")
sys.path.insert(0, str(ROOT))
import apply_cc_2025 as base  # noqa: E402

XLSX = ROOT / "Personal_Income_2025_Tax_Turbo.xlsx"
DELIVERABLE = ROOT / "tax_turbo_deliverable" / "Personal_Income_2025_Tax_Turbo.xlsx"
PARTS = ROOT / "tax_turbo_parts"
CSV_DIR = ROOT / "cc_fill_csv"
JSON_PATH = ROOT / "cc_starting_point_2025.json"
PARTS.mkdir(exist_ok=True)
CSV_DIR.mkdir(exist_ok=True)

MONTHS = base.MONTHS
PROP_COLS = base.PROP_COLS
EPGC_COLS = base.EPGC_COLS
GREEN = base.GREEN
YELLOW = base.YELLOW
BLUE = base.BLUE
GRAY = base.GRAY
HEADER_FILL = base.HEADER_FILL
HEADER_FONT = base.HEADER_FONT
THIN = base.THIN
ORANGE = PatternFill("solid", fgColor="FCE4D6")
LIGHT_GREEN = PatternFill("solid", fgColor="E2EFDA")

D = base.D
money = base.money
split_half = base.split_half
zeros = base.zeros
add = base.add


def split_n(amount, n: int) -> list[Decimal]:
    """Split a dollar amount into n parts; leftover pennies go to the first buckets."""
    cents = int((D(amount) * Decimal("100")).to_integral_value(rounding=ROUND_HALF_UP))
    base_c, rem = divmod(cents, n)
    return [Decimal(base_c + (1 if i < rem else 0)) / Decimal("100") for i in range(n)]


# Prime HOME HD / Lowe's / IKEA (Schauer / Good Earth already on Unit 2 — not here)
MARCH_HD = [
    ("2025-03-15", "THE HOME DEPOT #1901", D("28.14")),
    ("2025-03-15", "THE HOME DEPOT #1901", D("50.75")),
    ("2025-03-22", "THE HOME DEPOT #1901", D("93.96")),
    ("2025-03-26", "HOMEDEPOT.COM", D("1185.80")),
]
REMAINING = [
    ("2025-02-16", "THE HOME DEPOT #1901", D("34.01")),
    ("2025-02-17", "IKEA BOLINGBROOK", D("299.37")),
    ("2025-02-18", "IKEA 471245314", D("65.98")),
    ("2025-03-23", "LOWES #01845", D("55.69")),
    ("2025-04-20", "THE HOME DEPOT #1901", D("236.36")),
    ("2025-04-20", "THE HOME DEPOT #1901", D("1.07")),
    ("2025-07-06", "THE HOME DEPOT #1901", D("305.94")),
    ("2025-09-01", "LOWES #01845", D("209.19")),
    ("2025-10-11", "THE HOME DEPOT #1901", D("410.88")),
    ("2025-11-02", "THE HOME DEPOT #1901", D("193.97")),
]

u2_hd, u1_hd = zeros(), zeros()
grove_hd, personal_hd = zeros(), zeros()
split_rows = []  # audit rows for HOME_SPLIT


def month_idx(iso: str) -> int:
    return int(iso[5:7]) - 1


for date, merch, amt in MARCH_HD:
    u2_p, u1_p = split_half(amt)
    add(u2_hd, month_idx(date), u2_p)
    add(u1_hd, month_idx(date), u1_p)
    split_rows.append(
        {
            "date": date,
            "merchant": merch,
            "amount": amt,
            "rule": "March HD 100% 524 repairs, then 50/50 U1/U2",
            "to_524": amt,
            "to_personal": D("0"),
            "to_grove": D("0"),
            "u2": u2_p,
            "u1": u1_p,
        }
    )

for date, merch, amt in REMAINING:
    to_524, to_personal, to_grove = split_n(amt, 3)
    u2_p, u1_p = split_half(to_524)
    add(u2_hd, month_idx(date), u2_p)
    add(u1_hd, month_idx(date), u1_p)
    add(personal_hd, month_idx(date), to_personal)
    add(grove_hd, month_idx(date), to_grove)
    split_rows.append(
        {
            "date": date,
            "merchant": merch,
            "amount": amt,
            "rule": "1/3 524 repairs / 1/3 personal / 1/3 827 Grove (by transaction)",
            "to_524": to_524,
            "to_personal": to_personal,
            "to_grove": to_grove,
            "u2": u2_p,
            "u1": u1_p,
        }
    )

epgc_travel = zeros()
for m in (1, 4, 6, 7, 9):  # Park Chicago $20 Feb/May/Jul/Aug/Oct
    add(epgc_travel, m, "20.00")
add(epgc_travel, 10, "53.00")  # Field Museum Nov 1

MARCH_HD_TOTAL = sum((r[2] for r in MARCH_HD), D("0"))
REMAINING_TOTAL = sum((r[2] for r in REMAINING), D("0"))
assert sum(u2_hd) + sum(u1_hd) + sum(personal_hd) + sum(grove_hd) == MARCH_HD_TOTAL + REMAINING_TOTAL


def add_month_row(ws: Worksheet, row: int, values, colmap: dict, fill):
    for i, m in enumerate(MONTHS):
        if values[i] == 0:
            continue
        cell = ws.cell(row, colmap[m])
        cur = D(cell.value or 0)
        cell.value = money(cur + values[i])
        cell.fill = fill
        cell.number_format = "0.00"


def yellow_nonzero(ws: Worksheet, row: int, colmap: dict):
    for m in MONTHS:
        cell = ws.cell(row, colmap[m])
        if D(cell.value or 0) != 0:
            cell.fill = YELLOW
            cell.number_format = "0.00"


def label_map(ws: Worksheet, col: int, last: int = 40) -> dict:
    out = {}
    for r in range(1, last + 1):
        v = ws.cell(r, col).value
        if v:
            out[str(v).strip()] = r
    return out


def write_ledger_row(led, i, row, fills):
    for c, val in enumerate(row, 1):
        cell = led.cell(i, c, val)
        cell.border = THIN
        cell.fill = fills.get(row[6], BLUE)
        if c == 4 and isinstance(val, (int, float, Decimal)) and val:
            cell.number_format = "0.00"


def save_sheet_part(ws: Worksheet, path: Path):
    nb = Workbook()
    ns = nb.active
    ns.title = (ws.title or "Sheet")[:31]
    for row in ws.iter_rows():
        for cell in row:
            dst = ns.cell(cell.row, cell.column, cell.value)
            if cell.has_style:
                dst.font = copy(cell.font)
                dst.fill = copy(cell.fill)
                dst.number_format = cell.number_format
                dst.alignment = copy(cell.alignment)
    for i, col in ws.column_dimensions.items():
        if col.width:
            ns.column_dimensions[i].width = col.width
    nb.save(path)


def sheet_to_csv(ws: Worksheet, path: Path):
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        for row in ws.iter_rows(max_row=ws.max_row, max_col=max(ws.max_column, 1)):
            w.writerow([("" if c.value is None else c.value) for c in row])


def build_ledger() -> list[tuple]:
    rows = []
    for row in base.LEDGER:
        card, date, merch, amt, bucket, tab, status, note = row
        if "PARK CHICAGO" in merch:
            rows.append(
                (
                    card,
                    "2025-02/05/07/08/10",
                    "PARK CHICAGO MOBILE x5 $20",
                    100.00,
                    "Travel, Meals, and Entertainment",
                    "EPGC LLC",
                    "APPLIED",
                    "User 2026-09-18: yes EPGC. Cash months Feb 7, May 2, Jul 31, Aug 29, Oct 2.",
                )
            )
        elif "Field Museum" in merch:
            rows.append(
                (
                    "Coinbase 6108",
                    "2025-11-01",
                    "Field Museum",
                    53.00,
                    "Travel, Meals, and Entertainment",
                    "EPGC LLC",
                    "APPLIED",
                    "User 2026-09-18: Field Museum is EPGC LLC expense.",
                )
            )
            rows.append(
                (
                    "Coinbase 6108",
                    "2025-11-01",
                    "Soldier Field East Parking",
                    57.00,
                    "Travel?",
                    "EPGC / personal",
                    "HOLD",
                    "Same day as Field Museum. Not confirmed — left HOLD.",
                )
            )
        elif merch == "The Great Frame Up":
            rows.append(
                (
                    card,
                    date,
                    merch,
                    amt,
                    "Inventory framing",
                    "Art Sales",
                    "INVENTORY",
                    "User 2026-09-18: inventory framing. Parked on Art Sales; not EPGC P&L; not in Cost total until tied to an object.",
                )
            )
        elif "Los Angeles Modern" in merch:
            rows.append(
                (
                    card,
                    date,
                    merch,
                    amt,
                    "Inventory purchases",
                    "Art Sales",
                    "INVENTORY",
                    "User 2026-09-18: inventory purchase. Identify object before Cost (COGS) total — avoid Mercury double count.",
                )
            )
        elif merch == "Hindman LLC":
            rows.append(
                (
                    card,
                    date,
                    merch,
                    amt,
                    "Inventory purchases",
                    "Art Sales",
                    "INVENTORY",
                    "User 2026-09-18: inventory purchase. Identify object before Cost total.",
                )
            )
        elif date == "HOME" or merch.startswith("Home Depot"):
            continue
        elif "Canva" in merch:
            rows.append(
                (
                    card,
                    date,
                    merch,
                    amt,
                    "Software Fees",
                    "EPGC LLC",
                    "APPLIED",
                    "User 2026-09-18: yes EPGC.",
                )
            )
        elif "ComEd Jacob" in merch:
            rows.append(
                (
                    card,
                    date,
                    merch,
                    amt,
                    "ELECTRIC",
                    "524 Unit 2",
                    "TBD-MEGAN",
                    "REMIND JACOB: double-check ComEd with Megan. Monarch Adv Plus ACH (not CC). Keep amounts until she confirms meters / any extra check.",
                )
            )
        elif "ComEd Megan" in merch:
            rows.append(
                (
                    card,
                    date,
                    merch,
                    amt,
                    "ELECTRIC",
                    "524 Unit 1",
                    "TBD-MEGAN",
                    "REMIND JACOB: double-check ComEd with Megan. INDN Megan Gerrard on Jake checking. Keep amounts until confirmed.",
                )
            )
        else:
            rows.append(row)

    for rec in split_rows:
        if rec["to_524"]:
            rows.append(
                (
                    "Prime 2351",
                    rec["date"],
                    rec["merchant"],
                    float(rec["to_524"]),
                    "EXTERIOR REPAIRS",
                    "U1 50% / U2 50%",
                    "APPLIED",
                    rec["rule"] + f" → U2 ${rec['u2']} / U1 ${rec['u1']}",
                )
            )
        if rec["to_personal"]:
            rows.append(
                (
                    "Prime 2351",
                    rec["date"],
                    rec["merchant"],
                    float(rec["to_personal"]),
                    "personal",
                    "—",
                    "EXCLUDED",
                    "User 2026-09-18: remaining HD/Lowe's/IKEA 1/3 personal (by transaction).",
                )
            )
        if rec["to_grove"]:
            rows.append(
                (
                    "Prime 2351",
                    rec["date"],
                    rec["merchant"],
                    float(rec["to_grove"]),
                    "827 Grove materials",
                    "827 Grove CapEx",
                    "GROVE-SPLIT",
                    "1/3 of txn. Pre-close materials — NOT folded into locked TY2025 CapEx $53,660.",
                )
            )
    rows.append(
        (
            "Prime 2351",
            "2025-12-10",
            "EXTRA SPACE 6720",
            360.00,
            "storage?",
            "HOLD",
            "HOLD",
            "Was in Prime HOME total. Not HD/Lowe's/IKEA — still HOLD.",
        )
    )
    return rows


CONFIRMED = [
    ("Nicor meters", "CONFIRMED. Smaller Nicor → Unit 2 gas; larger seasonal → Unit 1."),
    ("Prime AT&T $45–50", "CONFIRMED. 100% Unit 2 INTERNET."),
    ("Village water", "CONFIRMED. Forest Park bill 50/50 U1/U2."),
    ("TruGreen / Alsip / Good Earth", "CONFIRMED. Unit 2 exterior/yard (not EPGC)."),
    ("Canva $15 Sep–Dec", "CONFIRMED. EPGC Software Fees."),
    ("Park Chicago $100", "CONFIRMED. EPGC Travel (five $20 BoA charges)."),
    ("Great Frame Up $536.41", "CONFIRMED inventory framing. Art Sales tab; not EPGC P&L."),
    (
        "LAMA $6,821.06 + Hindman $685.75",
        "CONFIRMED inventory purchases. Art Sales tab; identify objects before Cost total.",
    ),
    (
        "Field Museum $53",
        "CONFIRMED EPGC Travel (Coinbase 2025-11-01). Soldier Field East Parking $57 same day still HOLD.",
    ),
    (
        "March Home Depot → 524 repairs",
        f"CONFIRMED (read “534” as 524 Ferdinand). March HD ${MARCH_HD_TOTAL} → EXTERIOR REPAIRS 50/50 U1/U2.",
    ),
    (
        "Remaining HD / Lowe’s / IKEA",
        f"CONFIRMED split by transaction 1/3 524 / 1/3 personal / 1/3 Grove. Remaining ${REMAINING_TOTAL}. Extra Space $360 HOLD.",
    ),
    (
        "State Farm bundled home + auto",
        "CONFIRMED. Previously home=Travelers, car=Geico. Only the HOME share of Prime State Farm is on 524 INSURANCE (50/50 U1/U2, yellow). Auto share and Geico $540.89 credit are personal. Split uses Feb 11 unbundled $157.53 home / $145.25 auto as the ratio.",
    ),
]

STILL_OPEN = [
    (
        "REMIND JACOB — ComEd with Megan",
        "ComEd is not on these cards. Monarch Adv Plus shows ACH: Jacob Coley $829.15 → Unit 2 ELECTRIC; Megan Gerrard $1,565.26 → Unit 1 ELECTRIC (yellow). User: must have been paid directly or with a check. Double-check with Megan: meters, whether Jake checking ACH is the full picture, and any extra check/direct payments.",
    ),
    ("Soldier Field East Parking $57", "Coinbase 2025-11-01 same day as Field Museum. Not confirmed — HOLD."),
    ("Extra Space $360", "Prime 2025-12-10. Not part of the HD/Lowe’s/IKEA split — HOLD (storage vs moving vs personal)."),
    (
        "State Farm declarations / Sep–Dec",
        "Home vs auto $ is a Feb-charge proxy ($157.53 home / $145.25 auto) until declarations pages. None on Prime Sep–Dec 2025 (524 sold 12/18 — those months may be another account). 2023–24 Excel booked all homeowners on U1; this packet 50/50’s the home share. No 2025 Travelers found on Prime/Sapphire/BoA/Monarch.",
    ),
    (
        "Inventory object IDs",
        "Frame Up $536.41 + LAMA $6,821.06 + Hindman $282.24 + $403.51 parked off Cost total until objects named (avoid Mercury double count).",
    ),
    (
        "827 Grove 1/3 hardware vs locked CapEx",
        f"Grove share ${sum(grove_hd)} is pre-close materials on the Grove tab. Do NOT fold into locked TY2025 CapEx $53,660 until CPA says so.",
    ),
]


def rebuild_ledger(wb):
    if "CC_LEDGER" in wb.sheetnames:
        del wb["CC_LEDGER"]
    led = wb.create_sheet("CC_LEDGER", 1)
    led["A1"] = "2025 credit-card ledger after 2026-09-18 answers — not tax advice"
    led["A1"].font = Font(bold=True, size=14)
    led.merge_cells("A1:H1")
    headers = ["Card", "Date", "Merchant / item", "Amount", "Prior-Excel category", "Tab", "Status", "Note"]
    for i, h in enumerate(headers, 1):
        cell = led.cell(3, i, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN
    status_fill = {
        "APPLIED": GREEN,
        "EXCLUDED": GRAY,
        "HOLD": YELLOW,
        "HOLD-COGS": YELLOW,
        "INVENTORY": LIGHT_GREEN,
        "TBD-MEGAN": ORANGE,
        "GROVE-SPLIT": BLUE,
    }
    ledger = build_ledger()
    for i, row in enumerate(ledger, 4):
        write_ledger_row(led, i, row, status_fill)
    last_led = 3 + len(ledger)
    led.cell(
        last_led + 2,
        1,
        "STATUS: APPLIED = monthly P&L. EXCLUDED = personal. HOLD = still confirm. "
        "INVENTORY = Art Sales (not P&L, not Cost total yet). TBD-MEGAN = ComEd — remind Jacob to ask Megan. "
        "GROVE-SPLIT = 827 materials, not locked $53,660.",
    )
    led.merge_cells(start_row=last_led + 2, start_column=1, end_row=last_led + 2, end_column=8)
    for i, w in enumerate([18, 16, 48, 12, 32, 18, 14, 78], 1):
        led.column_dimensions[get_column_letter(i)].width = w
    led.auto_filter.ref = f"A3:H{last_led}"
    led.freeze_panes = "A4"
    return led


def rebuild_ask(wb):
    if "ASK" in wb.sheetnames:
        del wb["ASK"]
    ask = wb.create_sheet("ASK", 2)
    ask["A1"] = "REMIND JACOB: Double-check ComEd with Megan (checking ACH and/or a check; confirm meters)."
    ask["A1"].font = Font(bold=True, size=14, color="9C5700")
    ask["A1"].fill = ORANGE
    ask.merge_cells("A1:B1")
    ask.row_dimensions[1].height = 28
    ask["A2"] = (
        "Monarch Adv Plus already has ComEd ACH: Jacob Coley $829.15 (Unit 2, yellow) and Megan Gerrard $1,565.26 "
        "(Unit 1, yellow). User said it must have been paid directly or with a check — confirm with Megan whether "
        "that ACH is complete or if extra check/direct payments exist. Not tax advice."
    )
    ask["A2"].alignment = Alignment(wrap_text=True)
    ask["A2"].fill = ORANGE
    ask.merge_cells("A2:B2")
    ask.row_dimensions[2].height = 48

    ask["A4"] = "CONFIRMED 2026-09-18"
    ask["B4"] = "What was applied"
    ask["A4"].fill = HEADER_FILL
    ask["B4"].fill = HEADER_FILL
    ask["A4"].font = HEADER_FONT
    ask["B4"].font = HEADER_FONT
    r = 5
    for item, detail in CONFIRMED:
        ask.cell(r, 1, item).fill = GREEN
        ask.cell(r, 2, detail).fill = GREEN
        ask.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ask.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ask.cell(r, 1).border = THIN
        ask.cell(r, 2).border = THIN
        ask.row_dimensions[r].height = 36
        r += 1

    r += 1
    ask.cell(r, 1, "STILL OPEN").fill = HEADER_FILL
    ask.cell(r, 2, "What to confirm").fill = HEADER_FILL
    ask.cell(r, 1).font = HEADER_FONT
    ask.cell(r, 2).font = HEADER_FONT
    r += 1
    for item, detail in STILL_OPEN:
        fill = ORANGE if item.startswith("REMIND") else YELLOW
        ask.cell(r, 1, item).fill = fill
        ask.cell(r, 2, detail).fill = fill
        ask.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ask.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ask.cell(r, 1).border = THIN
        ask.cell(r, 2).border = THIN
        ask.row_dimensions[r].height = 52
        r += 1

    r += 1
    ask.cell(r, 1, "Applied / split totals").font = Font(bold=True)
    ask.cell(r, 1).fill = HEADER_FILL
    ask.cell(r, 1).font = HEADER_FONT
    ask.cell(r, 2).fill = HEADER_FILL
    summary = [
        ("EPGC Travel (Park Chicago $100 + Field Museum $53)", sum(epgc_travel)),
        ("March HD 100% 524 repairs", MARCH_HD_TOTAL),
        ("Remaining HD/Lowe’s/IKEA (split pool)", REMAINING_TOTAL),
        ("524 repairs from remaining 1/3 (then 50/50 U1/U2)", sum(u2_hd) + sum(u1_hd) - MARCH_HD_TOTAL),
        ("Unit 2 EXTERIOR add (March 50% + remaining 524-share 50%)", sum(u2_hd)),
        ("Unit 1 EXTERIOR add (March 50% + remaining 524-share 50%)", sum(u1_hd)),
        ("Personal 1/3 (excluded from P&L)", sum(personal_hd)),
        ("827 Grove 1/3 (not in locked $53,660)", sum(grove_hd)),
        ("Unit 2 ELECTRIC ComEd Jacob (YELLOW — ask Megan)", sum(base.u2_elec)),
        ("Unit 1 ELECTRIC ComEd Megan (YELLOW — ask Megan)", sum(base.u1_elec)),
        ("524 INSURANCE home share (Unit 2 50%, yellow proxy)", sum(base.u2_ins)),
        ("524 INSURANCE home share (Unit 1 50%, yellow proxy)", sum(base.u1_ins)),
        ("State Farm auto share (personal, excluded)", sum(base.sf_auto)),
        ("Geico auto credit (personal, excluded)", -base.GEICO_AUTO_CREDIT),
        ("Inventory parked (Frame Up + LAMA + Hindman)", D("8043.22")),
        ("Extra Space HOLD", D("360")),
        ("Soldier Field parking HOLD", D("57")),
    ]
    for i, (lab, val) in enumerate(summary):
        ask.cell(r + 1 + i, 1, lab)
        cell = ask.cell(r + 1 + i, 2, money(val))
        cell.number_format = '"$"#,##0.00'
        fill = ORANGE if "Megan" in lab or "YELLOW" in lab else GREEN
        if "HOLD" in lab or "Personal" in lab or "personal" in lab or "Geico" in lab or "auto share" in lab:
            fill = YELLOW if "HOLD" in lab else GRAY
        if "Grove" in lab:
            fill = BLUE
        if "INSURANCE" in lab:
            fill = YELLOW
        ask.cell(r + 1 + i, 1).fill = fill
        cell.fill = fill
    ask.column_dimensions["A"].width = 64
    ask.column_dimensions["B"].width = 110
    return ask


def main():
    wb = load_workbook(XLSX)

    # ----- 524 Unit 2 -----
    u2 = wb["524 Ferdinand Ave, Unit 2"]
    u2_map = label_map(u2, 2)
    add_month_row(u2, u2_map["EXTERIOR REPAIRS"], u2_hd, PROP_COLS, GREEN)
    yellow_nonzero(u2, u2_map["ELECTRIC"], PROP_COLS)
    u2["B30"] = (
        "CC answers 2026-09-18: Nicor small / Prime ATT / Grove / Schauer / yard CONFIRMED. "
        f"March HD ${MARCH_HD_TOTAL} → 524 EXTERIOR REPAIRS 50/50 with Unit 1. "
        "Remaining Prime HD/Lowe’s/IKEA 1/3 524 (50/50), 1/3 personal, 1/3 Grove. "
        "ELECTRIC yellow — REMIND JACOB: double-check ComEd with Megan (checking ACH Jacob $829.15; any extra check?). "
        "Home Depot March read as 524 (not 534). Extra Space $360 HOLD. Not tax advice."
    )
    u2["B30"].alignment = Alignment(wrap_text=True)
    u2["B30"].fill = YELLOW
    u2.row_dimensions[30].height = 78

    # ----- 524 Unit 1 -----
    u1 = wb["524 Ferdinand Ave, Unit 1"]
    u1_map = label_map(u1, 2)
    add_month_row(u1, u1_map["EXTERIOR REPAIRS"], u1_hd, PROP_COLS, GREEN)
    yellow_nonzero(u1, u1_map["ELECTRIC"], PROP_COLS)
    u1["B30"] = (
        "§121 residence through sale 2025-12-18. Nicor large / water 50/50 CONFIRMED. "
        f"March HD + remaining HD/Lowe’s/IKEA 524-share → EXTERIOR REPAIRS (U1 ${sum(u1_hd)}). "
        "ELECTRIC yellow — REMIND JACOB: double-check ComEd with Megan (checking ACH INDN Megan Gerrard $1,565.26; paid ACH and/or check?). "
        "Internet still $0 on this tab (Prime ATT on Unit 2). Not tax advice."
    )
    u1["B30"].alignment = Alignment(wrap_text=True)
    u1["B30"].fill = YELLOW
    u1.row_dimensions[30].height = 78

    # ----- EPGC travel -----
    epgc = wb["EPGC LLC"]
    epgc_map = label_map(epgc, 1)
    travel_key = "Travel, Meals, and Entertainment"
    base.write_month_row(epgc, epgc_map[travel_key], epgc_travel, EPGC_COLS, GREEN)
    last = epgc.max_row
    # replace old starting-point note if present
    for r in range(1, last + 1):
        v = epgc.cell(r, 1).value
        if v and "CC starting point" in str(v):
            epgc.cell(r, 1).value = (
                "CC answers 2026-09-18: Canva CONFIRMED Software Fees; Park Chicago $100 + Field Museum $53 → Travel. "
                "Coinbase golf/voids still excluded; Soldier Field parking $57 HOLD. "
                "Frame Up / LAMA / Hindman are Art Sales inventory, not this P&L. TruGreen is 524 yard. Not tax advice."
            )
            epgc.cell(r, 1).alignment = Alignment(wrap_text=True)
            break

    # ----- Art Sales inventory -----
    art = wb["Art Sales and Purchases"]
    art["A28"] = "2025 CC inventory (CONFIRMED category — NOT in Cost total until object is identified)"
    art["A28"].fill = YELLOW
    art["A29"] = "Los Angeles Modern Auctions (inventory purchase)"
    art["B29"] = "Sapphire 5423"
    art["C29"] = 6821.06
    art["D29"] = "2025-07-24"
    art["C29"].number_format = "0.00"
    art["C29"].fill = YELLOW
    art["A30"] = "Hindman LLC (inventory purchase)"
    art["B30"] = "Sapphire 5423"
    art["C30"] = 282.24
    art["D30"] = "2025-04-03"
    art["C30"].number_format = "0.00"
    art["C30"].fill = YELLOW
    art["A31"] = "Hindman LLC (inventory purchase)"
    art["B31"] = "Sapphire 5423"
    art["C31"] = 403.51
    art["D31"] = "2025-07-29"
    art["C31"].number_format = "0.00"
    art["C31"].fill = YELLOW
    art["A32"] = "The Great Frame Up (inventory framing)"
    art["B32"] = "Sapphire 5423"
    art["C32"] = 536.41
    art["D32"] = "2025-02-16"
    art["C32"].number_format = "0.00"
    art["C32"].fill = YELLOW
    art["A33"] = (
        "User 2026-09-18: these are inventory, not EPGC operating expense. "
        "Do not add into Cost (COGS) total until the object is named — avoid double count vs Mercury."
    )
    art["A33"].fill = YELLOW
    art.merge_cells("A33:I33")

    # ----- 827 Grove CapEx materials (do not touch locked $53,660) -----
    grove = wb["827 Grove CapEx"]
    grove["A11"] = "Prime HD / Lowe’s / IKEA — Grove 1/3 (user split 2026-09-18)"
    grove["A11"].font = Font(bold=True)
    grove["A11"].fill = BLUE
    grove.merge_cells("A11:D11")
    headers = ["Date", "Merchant", "Full txn", "Grove 1/3"]
    for i, h in enumerate(headers, 1):
        cell = grove.cell(12, i, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN
    r = 13
    for rec in split_rows:
        if rec["to_grove"] == 0:
            continue
        grove.cell(r, 1, rec["date"]).border = THIN
        grove.cell(r, 2, rec["merchant"]).border = THIN
        c3 = grove.cell(r, 3, money(rec["amount"]))
        c4 = grove.cell(r, 4, money(rec["to_grove"]))
        c3.number_format = "0.00"
        c4.number_format = "0.00"
        c3.border = THIN
        c4.border = THIN
        c4.fill = BLUE
        r += 1
    grove.cell(r, 2, "Grove 1/3 total").font = Font(bold=True)
    tot = grove.cell(r, 4, money(sum(grove_hd)))
    tot.number_format = '"$"#,##0.00'
    tot.font = Font(bold=True)
    tot.fill = BLUE
    grove.cell(r + 2, 1, (
        "CPA: this 1/3 is pre-close hardware/furniture (close 2025-12-18). "
        "Do NOT fold into locked TY2025 CapEx $53,660 (the 2025-12-19 contractor payments) until you say so. "
        "May be additional 827 basis / CapEx / personal — FLAG."
    ))
    grove.cell(r + 2, 1).alignment = Alignment(wrap_text=True)
    grove.merge_cells(start_row=r + 2, start_column=1, end_row=r + 3, end_column=4)
    grove.column_dimensions["A"].width = 14
    grove.column_dimensions["B"].width = 28
    grove.column_dimensions["C"].width = 12
    grove.column_dimensions["D"].width = 12

    # ----- SOURCE -----
    src = wb["_SOURCE_2025"]
    src["A16"] = "CC answers 2026-09-18"
    src["B16"] = (
        "Q1–4 CONFIRMED. March HD → 524 repairs 50/50. Remaining HD/Lowe’s/IKEA 1/3 524 / 1/3 personal / 1/3 Grove. "
        "ComEd yellow — REMIND JACOB: double-check with Megan (checking ACH or extra check). "
        "Canva + Park Chicago + Field Museum $53 EPGC. Frame Up framing + LAMA/Hindman inventory. "
        "Soldier Field parking $57 and Extra Space $360 still HOLD. Not tax advice."
    )
    src["A19"] = "REMIND JACOB"
    src["B19"] = "Double-check ComEd with Megan — paid from checking ACH (and/or a check); confirm meters and whether any extra direct/check payments exist."
    src["A19"].fill = YELLOW
    src["B19"].fill = YELLOW

    # ----- HOME_SPLIT audit -----
    if "HOME_SPLIT" in wb.sheetnames:
        del wb["HOME_SPLIT"]
    hs = wb.create_sheet("HOME_SPLIT", 3)
    hs["A1"] = "Prime HOME HD / Lowe’s / IKEA split (user 2026-09-18) — Extra Space $360 excluded / HOLD"
    hs["A1"].font = Font(bold=True, size=14)
    hs.merge_cells("A1:I1")
    hs["A2"] = (
        "March Home Depot 100% → 524 repairs, then 50/50 Unit 1 / Unit 2. "
        "Each remaining txn split 1/3 524 / 1/3 personal / 1/3 827 Grove; leftover pennies to 524 then personal. "
        "524 third then 50/50 (extra penny to Unit 2). Not tax advice."
    )
    hs["A2"].alignment = Alignment(wrap_text=True)
    hs.merge_cells("A2:I2")
    hs.row_dimensions[2].height = 36
    hh = ["Date", "Merchant", "Txn $", "Rule", "524 $", "Personal $", "Grove $", "U2 50%", "U1 50%"]
    for i, h in enumerate(hh, 1):
        cell = hs.cell(4, i, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN
    for i, rec in enumerate(split_rows, 5):
        vals = [
            rec["date"],
            rec["merchant"],
            money(rec["amount"]),
            rec["rule"],
            money(rec["to_524"]),
            money(rec["to_personal"]),
            money(rec["to_grove"]),
            money(rec["u2"]),
            money(rec["u1"]),
        ]
        for c, val in enumerate(vals, 1):
            cell = hs.cell(i, c, val)
            cell.border = THIN
            if c in (3, 5, 6, 7, 8, 9):
                cell.number_format = "0.00"
            if rec["to_personal"] == 0 and rec["to_grove"] == 0:
                cell.fill = GREEN
            else:
                if c == 6:
                    cell.fill = GRAY
                elif c == 7:
                    cell.fill = BLUE
                elif c in (5, 8, 9):
                    cell.fill = GREEN
    tot_r = 5 + len(split_rows)
    hs.cell(tot_r, 2, "TOTAL").font = Font(bold=True)
    for col, val in (
        (3, MARCH_HD_TOTAL + REMAINING_TOTAL),
        (5, sum(u2_hd) + sum(u1_hd)),
        (6, sum(personal_hd)),
        (7, sum(grove_hd)),
        (8, sum(u2_hd)),
        (9, sum(u1_hd)),
    ):
        cell = hs.cell(tot_r, col, money(val))
        cell.number_format = '"$"#,##0.00'
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    hs.cell(tot_r + 2, 1, "Extra Space 2025-12-10 $360.00 — HOLD, not in this split.")
    hs.cell(tot_r + 2, 1).fill = YELLOW
    widths = [12, 28, 12, 55, 12, 12, 12, 12, 12]
    for i, w in enumerate(widths, 1):
        hs.column_dimensions[get_column_letter(i)].width = w
    hs.freeze_panes = "A5"

    rebuild_ledger(wb)
    rebuild_ask(wb)

    wb.save(XLSX)
    shutil.copy2(XLSX, DELIVERABLE)

    export_names = [
        ("EPGC LLC", "EPGC_LLC.csv"),
        ("524 Ferdinand Ave, Unit 2", "524_Unit_2.csv"),
        ("524 Ferdinand Ave, Unit 1", "524_Unit_1.csv"),
        ("CC_LEDGER", "CC_LEDGER.csv"),
        ("ASK", "ASK.csv"),
        ("HOME_SPLIT", "HOME_SPLIT.csv"),
        ("Art Sales and Purchases", "Art_Sales.csv"),
        ("827 Grove CapEx", "Grove_CapEx.csv"),
        ("_SOURCE_2025", "SOURCE_2025.csv"),
    ]
    for name, fname in export_names:
        sheet_to_csv(wb[name], CSV_DIR / fname)
        safe = name.replace(",", "").replace(".", "")[:40]
        save_sheet_part(wb[name], PARTS / f"{safe}.xlsx")

    payload = {
        "updated": "2026-09-18",
        "answers": "user 2026-09-18 Q1-10",
        "method": "cash basis, transaction date month; March HD 100% 524 then 50/50; remaining HD/Lowes/IKEA 1/3 each by txn",
        "remind_jacob": "Double-check ComEd with Megan — checking ACH and/or extra check; confirm meters.",
        "march_hd_524": float(MARCH_HD_TOTAL),
        "remaining_pool": float(REMAINING_TOTAL),
        "unit2_exterior_add": float(sum(u2_hd)),
        "unit1_exterior_add": float(sum(u1_hd)),
        "personal_third": float(sum(personal_hd)),
        "grove_third": float(sum(grove_hd)),
        "epgc_travel": float(sum(epgc_travel)),
        "field_museum": 53.0,
        "park_chicago": 100.0,
        "inventory_parked": 8043.22,
        "comed": {
            "unit2_jacob": float(sum(base.u2_elec)),
            "unit1_megan": float(sum(base.u1_elec)),
            "status": "TBD-MEGAN",
        },
        "hold": {"extra_space": 360.0, "soldier_field_parking": 57.0},
        "confirmed": [a[0] for a in CONFIRMED],
        "still_open": [a[0] for a in STILL_OPEN],
    }
    JSON_PATH.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    print("March HD", MARCH_HD_TOTAL)
    print("Remaining pool", REMAINING_TOTAL)
    print("U2 HD add", sum(u2_hd), "U1 HD add", sum(u1_hd))
    print("Personal 1/3", sum(personal_hd), "Grove 1/3", sum(grove_hd))
    print("EPGC travel", sum(epgc_travel))
    print("conservation", sum(u2_hd) + sum(u1_hd) + sum(personal_hd) + sum(grove_hd))
    print("saved", XLSX, "bytes", XLSX.stat().st_size)


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""Clarify “grove orders” = Grove Collaborative (524 U2 SUPPLIES), not 827 N Grove.

Does not change dollar amounts. Do not re-run apply_cc_answers_2025.py (would
double HD). Not tax advice.
"""
from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

ROOT = Path("/workspace/.cursor/scratch")
sys.path.insert(0, str(ROOT))
import apply_cc_2025 as base  # noqa: E402
import apply_cc_answers_2025 as ans  # noqa: E402
import apply_state_farm_2025 as sf  # noqa: E402

XLSX = ans.XLSX
DELIVERABLE = ans.DELIVERABLE


def main():
    assert XLSX.exists(), XLSX
    wb = load_workbook(XLSX)

    u2 = wb["524 Ferdinand Ave, Unit 2"]
    u2["B30"] = sf.U2_NOTE
    u2["B30"].alignment = Alignment(wrap_text=True)
    u2["B30"].fill = base.YELLOW
    u2.row_dimensions[30].height = 108

    u1 = wb["524 Ferdinand Ave, Unit 1"]
    u1["B30"] = sf.U1_NOTE
    u1["B30"].alignment = Alignment(wrap_text=True)
    u1["B30"].fill = base.YELLOW

    grove = wb["827 Grove CapEx"]
    grove["A11"] = "Prime HD / Lowe’s / IKEA — 827 N Grove the house 1/3 (NOT Grove Collaborative supplies)"
    grove["A11"].font = Font(bold=True)
    grove["A11"].fill = base.BLUE
    # CPA note sits two rows after the total; find the FLAG cell.
    for r in range(1, grove.max_row + 1):
        v = grove.cell(r, 1).value
        if v and "CPA:" in str(v):
            grove.cell(r, 1).value = (
                "CPA: this 1/3 is pre-close hardware/furniture for 827 N Grove the house (close 2025-12-18). "
                "This is NOT Grove Collaborative. "
                "Do NOT fold into locked TY2025 CapEx $53,660 (the 2025-12-19 contractor payments) until you say so. "
                "May be additional 827 basis / CapEx / personal — FLAG."
            )
            grove.cell(r, 1).alignment = Alignment(wrap_text=True)
            break

    if "HOME_SPLIT" in wb.sheetnames:
        hs = wb["HOME_SPLIT"]
        hs["A2"] = (
            "March Home Depot 100% → 524 repairs, then 50/50 Unit 1 / Unit 2. "
            "Each remaining txn split 1/3 524 / 1/3 personal / 1/3 827 N Grove the house "
            "(NOT Grove Collaborative); leftover pennies to 524 then personal. "
            "524 third then 50/50 (extra penny to Unit 2). Not tax advice."
        )
        hs["A2"].alignment = Alignment(wrap_text=True)

    src = wb["_SOURCE_2025"]
    src["A21"] = "Grove Collaborative 2026-09-18"
    src["B21"] = (
        f"User: “grove orders” = Grove Collaborative, not 827 N Grove the house. "
        f"Prime ${sum(base.u2_sup)} already on 524 Unit 2 SUPPLIES (8 charges). "
        f"HD/Lowe’s/IKEA 1/3 ${sum(ans.grove_hd)} stays on 827 N Grove the house (Q5), "
        "not mixed into Grove Collaborative. Betty’s Pizza / KS Grove is personal dining. Not tax advice."
    )
    src["A21"].fill = base.GREEN
    src["B21"].fill = base.GREEN

    ans.rebuild_ledger(wb)
    ans.rebuild_ask(wb)

    wb.save(XLSX)
    shutil.copy2(XLSX, DELIVERABLE)

    export_names = [
        ("EPGC LLC", "EPGC_LLC.csv"),
        ("524 Ferdinand Ave, Unit 2", "524_Unit_2.csv"),
        ("524 Ferdinand Ave, Unit 1", "524_Unit_1.csv"),
        ("CC_LEDGER", "CC_LEDGER.csv"),
        ("ASK", "ASK.csv"),
        ("HOME_SPLIT", "HOME_SPLIT.csv"),
        ("SF_SPLIT", "SF_SPLIT.csv"),
        ("Art Sales and Purchases", "Art_Sales.csv"),
        ("827 Grove CapEx", "Grove_CapEx.csv"),
        ("_SOURCE_2025", "SOURCE_2025.csv"),
    ]
    for name, fname in export_names:
        ans.sheet_to_csv(wb[name], ans.CSV_DIR / fname)
        safe = name.replace(",", "").replace(".", "")[:40]
        ans.save_sheet_part(wb[name], ans.PARTS / f"{safe}.xlsx")

    payload = json.loads(ans.JSON_PATH.read_text(encoding="utf-8")) if ans.JSON_PATH.exists() else {}
    payload["grove_collaborative"] = {
        "user": "grove orders = Grove Collaborative, not 827 N Grove the house",
        "unit2_supplies": float(sum(base.u2_sup)),
        "charges": 8,
        "house_hd_third": float(sum(ans.grove_hd)),
        "house_hd_third_tab": "827 Grove CapEx (Q5 split, unchanged dollars)",
    }
    payload["confirmed"] = [a[0] for a in ans.CONFIRMED]
    payload["still_open"] = [a[0] for a in ans.STILL_OPEN]
    ans.JSON_PATH.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    print("Grove Collaborative U2 SUPPLIES", sum(base.u2_sup))
    print("827 N Grove house HD 1/3 unchanged", sum(ans.grove_hd))
    print("saved", XLSX, "bytes", XLSX.stat().st_size)


if __name__ == "__main__":
    main()

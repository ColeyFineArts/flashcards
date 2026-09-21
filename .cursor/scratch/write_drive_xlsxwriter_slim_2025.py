#!/usr/bin/env python3
"""Slim xlsxwriter Personal Income 2025 for Drive MCP convert.

openpyxl theme/styles and custom no-theme OOXML both fail Drive convert
or exceed the create_file payload. xlsxwriter writes a real theme and
small XML. Keep Art Sales rows 1–83 (2022–2024 lots + Belt/Venus/August).

Not tax advice.
"""
from __future__ import annotations

import sys
from pathlib import Path

import xlsxwriter
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from export_ssml_personal_income_2025 import (
    PACKET_TABS,
    keep_row,
    keep_value,
    max_col_for,
)

ROOT = Path("/workspace/.cursor/scratch")
SRC = ROOT / "tax_turbo_parts" / "Personal Income 2025.xlsx"
OUT = ROOT / "tax_turbo_parts" / "Personal Income 2025.drive.xlsx"


def write_book(src: Path, dest: Path) -> None:
    wb_src = load_workbook(src, data_only=False)
    wb = xlsxwriter.Workbook(str(dest), {"constant_memory": False})
    f_n = wb.add_format({"font_name": "Century Gothic", "font_size": 10})
    f_b = wb.add_format({"font_name": "Century Gothic", "font_size": 11, "bold": True})
    f_g = wb.add_format(
        {
            "font_name": "Century Gothic",
            "font_size": 10,
            "bg_color": "C6EFCE",
            "num_format": '$#,##0.00',
        }
    )
    f_y = wb.add_format(
        {
            "font_name": "Century Gothic",
            "font_size": 10,
            "bg_color": "FFF2CC",
            "num_format": '$#,##0.00',
        }
    )
    f_m = wb.add_format(
        {"font_name": "Century Gothic", "font_size": 10, "num_format": '$#,##0.00'}
    )
    for name in PACKET_TABS:
        ws_s = wb_src[name]
        ws = wb.add_worksheet(name)
        ws.set_column(1, 1, 36)
        max_r = ws_s.max_row or 1
        max_c = max_col_for(name, ws_s)
        for r in range(1, max_r + 1):
            if not keep_row(name, r, ws_s):
                continue
            for c in range(1, max_c + 1):
                cell = ws_s.cell(r, c)
                val = cell.value
                is_f = cell.data_type == "f" and isinstance(val, str)
                if is_f and "[" in val:
                    is_f = False
                if not is_f and not keep_value(val, cell, ws_s):
                    continue
                bg = None
                fill = cell.fill
                if fill and fill.fill_type not in (None, "none"):
                    rgb = getattr(getattr(fill, "fgColor", None), "rgb", None)
                    if rgb:
                        s = str(rgb)
                        if s.startswith("FF") and len(s) == 8:
                            s = s[2:]
                        bg = s.upper()
                money = "$" in (cell.number_format or "") or "0.00" in (cell.number_format or "")
                if bg == "C6EFCE":
                    fmt = f_g
                elif bg == "FFF2CC":
                    fmt = f_y
                elif money or is_f:
                    fmt = f_m
                elif cell.font and cell.font.bold:
                    fmt = f_b
                else:
                    fmt = f_n
                row, col = r - 1, c - 1
                if is_f:
                    fml = val if val.startswith("=") else f"={val}"
                    ws.write_formula(row, col, fml, fmt)
                elif isinstance(val, (int, float)) and not isinstance(val, bool):
                    ws.write_number(row, col, float(val), fmt)
                elif isinstance(val, str):
                    text = val if len(val) <= 120 else val[:117] + "..."
                    ws.write_string(row, col, text, fmt)
                elif val is not None:
                    ws.write(row, col, str(val), fmt)
        for rng in ws_s.merged_cells.ranges:
            if keep_row(name, rng.min_row, ws_s) and rng.max_col <= max_c:
                anchor = ws_s.cell(rng.min_row, rng.min_col)
                ws.merge_range(str(rng), anchor.value or "", f_b)
    wb.close()


if __name__ == "__main__":
    write_book(SRC, OUT)
    sys.path.insert(0, str(ROOT))
    from assert_personal_income_2025 import assert_workbook
    from inject_cached_formula_values import inject
    from sanitize_xlsx_for_sheets import assert_formulas_are_real, sanitize_xlsx

    sanitize_xlsx(OUT)
    n = inject(OUT)
    sanitize_xlsx(OUT)
    try:
        assert_workbook(OUT)
    except SystemExit as e:
        print("assert", e)
    assert_formulas_are_real(OUT)
    print("slim drive", OUT, OUT.stat().st_size, "bytes", f"cached={n}")

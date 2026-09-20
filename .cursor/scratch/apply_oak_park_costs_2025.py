#!/usr/bin/env python3
"""216 N. Oak Park 2025 costs — user list 2026-09-20, last-year Excel pattern.

Does NOT re-run apply_checking_answers_2025.py or apply_ltr_insurance_2025.py.
Does NOT touch 524, EPGC, Mercury P&L, EOEB/L5.

User list (accounted on the 216 tab like 2023|2024):
  * Monthly mortgage $1,226.25
  * HOA $421.53
  * Insurance $42.84 per month  (Lemonade $514 cash 10/3 amortized, same as 2023/2024)
  * $150 Chris Brennan Exterminator 8/7/25
  * Smoke detector Amazon          (ASK — no SKU on Monarch)
  * $300 cleaning Imelda 9/30/25 empty apartment
  * $675 painting and repairs to Joan
  * $11 key cutting

Not tax advice.
"""
from __future__ import annotations

import csv
import json
import shutil
import sys
from copy import copy
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path("/workspace/.cursor/scratch")
sys.path.insert(0, str(ROOT))
import apply_cc_answers_2025 as ans  # noqa: E402

XLSX = ans.XLSX
DELIVERABLE = ans.DELIVERABLE
CSV_DIR = ans.CSV_DIR
PARTS = ans.PARTS
PACKET = ROOT / "cpa_packet_v1.json"

GREEN = PatternFill("solid", fgColor="C6EFCE")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
ACCT = '_("$"* #,##0.00_);_("$"* \\(#,##0.00\\);_("$"* "-"??_);_(@_)'
CG = "Century Gothic"
MONTHS = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEPT", "OCT", "NOV", "DEC"]

INS_MO = Decimal("42.84")  # user; 12 × 42.84 = 514.08 vs Lemonade cash $514.00
OAK_INS = [float(INS_MO)] * 12

# Cash on 8507: Jul–Oct $1,226.25, Nov–Dec $1,177.52. Jan–Jun missing (WAIT 8507).
# User: $1,226.25 monthly — yellow Jan–Jun at that amount like last year's even HOA/mortgage.
OAK_MORTGAGE = [
    1226.25, 1226.25, 1226.25, 1226.25, 1226.25, 1226.25,
    1226.25, 1226.25, 1226.25, 1226.25, 1177.52, 1177.52,
]
OAK_MORTGAGE_WAIT = [0, 1, 2, 3, 4, 5]
OAK_HOA = [421.53] * 12
OAK_HOA_WAIT = [0, 1, 2, 3, 4]

# INTERIOR MAINTENANCE: Brennan $150 Aug (8507 check posted 8/18) + Ace keys $11 Dec.
OAK_INTERIOR = [0, 0, 0, 0, 0, 0, 0, 150.0, 0, 0, 0, 11.0]
# INTERIOR REPAIR: Joan $675 painting + repairs 12/1 (was on INTERIOR MAINTENANCE).
OAK_REPAIR = [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 675.0]
# MOVE OUT FEE: Imelda $300 empty-apt. User 9/30; Monarch cash 10/1 on 0203.
# 2023/2024 booked MOVE OUT FEE in October — same vacancy-month pattern.
OAK_MOVEOUT = [0, 0, 0, 0, 0, 0, 0, 0, 0, 300.0, 0, 0]

# Turbo-only CLEANING row (no MOVE OUT FEE on the 2025-only CSV).
TURBO_CLEANING = [0, 0, 0, 0, 0, 0, 0, 0, 0, 300.0, 0, 0]
TURBO_INTERIOR = [0, 0, 0, 0, 0, 0, 0, 150.0, 0, 0, 0, 686.0]  # Brennan + Joan + keys

NOTE = (
    "LOCKED 2026-09-20 (user cost list, last-year Excel pattern). "
    "RENTER'S INSURANCE $42.84/mo × 12 = $514.08 (Lemonade cash $514.00 on 10/3 Megan Chase — 8¢ rounding, same amortize as 2023 $26.75 and 2024 $35.08). "
    "MORTGAGE: user $1,226.25/mo. Cash 8507 Jul–Oct $1,226.25 (green) / Nov–Dec $1,177.52 (green, escrow drop). Jan–Jun $1,226.25 yellow WAIT 8507 (not on Monarch). "
    "HOA Santa Maria C326 $421.53: Jun–Dec cash 8507 green; Jan–May $421.53 yellow WAIT 8507. "
    "Chris Brennan exterminator $150 8/7 → INTERIOR Aug (8507 check posted 8/18 $150). "
    "Imelda empty-apt cleaning $300 9/30 → MOVE OUT FEE October (cash 10/1 joint 0203; 2023/24 move-out was October). Other Imelda $150s left off 216 (524 STR pattern). "
    "Joan $675 12/1 painting+repairs → INTERIOR REPAIR December (8507 Zelle). "
    "Ace Hardware $11 12/8 Megan Chase → INTERIOR Dec keys. Same-day Ace $17 NOT booked. "
    "Smoke detector Amazon: no SKU on Monarch/Prime/Sapphire/BoA — EQUIPMENT stays $0 ASK. "
    "Sch E: deduct Rocket 1098 interest, not full P+I. Not tax advice."
)


def D(x) -> Decimal:
    return Decimal(str(x)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def write_months(ws: Worksheet, row: int, start_col: int, values, wait_idx=None) -> None:
    wait_idx = set(wait_idx or [])
    for i, v in enumerate(values):
        cell = ws.cell(row, start_col + i)
        cell.value = float(D(v))
        cell.number_format = ACCT
        if i in wait_idx:
            cell.fill = YELLOW
        elif D(v) != 0:
            cell.fill = GREEN


def copy_sheet_light(src: Worksheet, dest: Worksheet, max_col: int | None = None) -> None:
    mc = min(src.max_column or 1, max_col or 42)
    mr = src.max_row or 1
    for r in src.iter_rows(min_row=1, max_row=mr, max_col=mc):
        for cell in r:
            d = dest.cell(cell.row, cell.column, cell.value)
            if cell.has_style:
                d.font = copy(cell.font)
                d.fill = copy(cell.fill)
                d.alignment = copy(cell.alignment)
                d.number_format = cell.number_format
    for col, dim in src.column_dimensions.items():
        dest.column_dimensions[col].width = dim.width
    for rng in src.merged_cells.ranges:
        try:
            dest.merge_cells(str(rng))
        except Exception:
            pass


def fill_xlsx(ws: Worksheet) -> None:
    assert "INSUR" in str(ws["B14"].value).upper(), ws["B14"].value
    assert str(ws["B15"].value).upper().startswith("INTERIOR MAINT")
    assert "REPAIR" in str(ws["B16"].value).upper()
    assert "MORTGAGE" in str(ws["B35"].value).upper()
    assert str(ws["B36"].value).upper() == "HOA"
    assert "MOVE OUT" in str(ws["B38"].value).upper()
    ac = 29
    write_months(ws, 14, ac, OAK_INS)
    write_months(ws, 15, ac, OAK_INTERIOR)
    write_months(ws, 16, ac, OAK_REPAIR)
    write_months(ws, 35, ac, OAK_MORTGAGE, wait_idx=OAK_MORTGAGE_WAIT)
    write_months(ws, 36, ac, OAK_HOA, wait_idx=OAK_HOA_WAIT)
    write_months(ws, 38, ac, OAK_MOVEOUT)
    # EQUIPMENT stays $0 — smoke detector ASK
    for i in range(12):
        cell = ws.cell(29, ac + i)
        if cell.value in (None, ""):
            cell.value = 0
            cell.number_format = ACCT
        cell.fill = YELLOW

    notes = {
        60: "2025 (this column) — user cost list LOCKED 2026-09-20, last-year Excel pattern.",
        61: "Rent — joint 0203 (Delach $1,450 + Ava $500; Joan $1,950 Dec + prepaid Jan). Vacant Oct + half Nov. Half-Nov $975 WAIT 8507.",
        62: "Mortgage — user $1,226.25/mo. Cash 8507 Jul–Oct $1,226.25 / Nov–Dec $1,177.52 (green). Jan–Jun $1,226.25 yellow WAIT 8507 (not on Monarch). Sch E uses Rocket 1098 interest, not full P+I.",
        63: "HOA — user $421.53/mo Santa Maria C326. Jun–Dec cash 8507 green. Jan–May $421.53 yellow WAIT 8507.",
        64: "Insurance — user $42.84/mo × 12 = $514.08 on RENTER'S INSURANCE (amortized like 2023 $26.75 / 2024 $35.08). Lemonade cash $514.00 10/3 Megan Chase (8¢). Not Travelers/Geico/State Farm.",
        65: "Interior — Chris Brennan exterminator $150 8/7 (8507 check posted 8/18) INTERIOR Aug. Ace keys $11 12/8 Megan Chase INTERIOR Dec. Same-day Ace $17 not booked.",
        66: "Interior repair — Joan $675 12/1 8507 Zelle painting + repairs → INTERIOR REPAIR December (2024 used this row for Oak Park painting).",
        67: "Move-out — Imelda $300 empty-apt. User 9/30; cash 10/1 joint 0203 → MOVE OUT FEE October (same month as 2023/24 $200). Other 2025 Imelda $150s not on 216.",
        68: "Smoke detector Amazon — still ASK (checklist TBD). No Kidde/First Alert/detector SKU on Monarch/Prime/Sapphire/BoA. EQUIPMENT yellow $0. Not tax advice.",
    }
    for r, text in notes.items():
        ws.cell(r, 2, text)
        ws.cell(r, 2).alignment = Alignment(wrap_text=True)
        ws.cell(r, 2).font = Font(name=CG, size=10)
        ws.cell(r, 2).fill = GREEN if r < 68 else YELLOW
        ws.row_dimensions[r].height = 36
    ws.row_dimensions[64].height = 48
    ws.row_dimensions[62].height = 48


def set_csv_months(rows: list[list[str]], label: str, values: list[float]) -> None:
    found = False
    for row in rows:
        if len(row) > 1 and row[1] == label:
            for i, v in enumerate(values):
                row[2 + i] = str(float(D(v)))
            found = True
            break
    if not found:
        raise SystemExit(f"missing {label} in 216_Oak_Park.csv")


def export_turbo_csv() -> None:
    path = CSV_DIR / "216_Oak_Park.csv"
    rows = list(csv.reader(path.open(encoding="utf-8")))
    set_csv_months(rows, "INSURANCE", OAK_INS)
    set_csv_months(rows, "INTERIOR MAINTENANCE", TURBO_INTERIOR)
    set_csv_months(rows, "CLEANING / TURNOVER", TURBO_CLEANING)
    set_csv_months(rows, "MORTGAGE (cash P+I — memo only)", OAK_MORTGAGE)
    set_csv_months(rows, "HOA", OAK_HOA)
    for row in rows:
        if len(row) > 1 and isinstance(row[1], str) and (
            row[1].startswith("LTR after") or "Oak Park" in row[1] and "Lemonade" in row[1]
        ):
            row[1] = NOTE
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def write_lock_csv() -> None:
    path = CSV_DIR / "LOCK_Oak_Park_2025.csv"
    rows = [
        ["Item (user list)", "Status", "Where on 216 tab", "2025 dollars", "Cash found", "Note"],
        [
            "Monthly mortgage $1,226.25",
            "LOCKED cash Jul–Dec / WAIT Jan–Jun",
            "MORTGAGE (cash P+I — memo). Sch E uses 1098 interest.",
            f"${sum(OAK_MORTGAGE):.2f} sheet / cash Jul–Dec ${1226.25 * 4 + 1177.52 * 2:.2f}",
            "8507 Rocket Jul–Oct $1,226.25 × 4; Nov–Dec $1,177.52 × 2. Jan–Jun none on Monarch.",
            "Nov–Dec booked at cash $1,177.52 (like 2023 Nov–Dec step-change), not forced to $1,226.25. Jan–Jun yellow $1,226.25 WAIT 8507.",
        ],
        [
            "HOA $421.53",
            "LOCKED cash Jun–Dec / WAIT Jan–May",
            "HOA",
            f"${sum(OAK_HOA):.2f} sheet / cash Jun–Dec ${421.53 * 7:.2f}",
            "8507 Santa Maria C326 Jun–Dec $421.53 × 7.",
            "Jan–May yellow $421.53 WAIT 8507 (same even monthly as 2023/2024 Excel).",
        ],
        [
            "Insurance $42.84 per month",
            "LOCKED amortized (last-year pattern)",
            "RENTER'S INSURANCE",
            f"${float(INS_MO * 12):.2f}",
            "Lemonade $514.00 cash 2025-10-03 Megan Chase 1702.",
            "12 × $42.84 = $514.08 vs cash $514.00 (8¢). 2023 Excel $26.75 × 12 = $321; 2024 $35.08 × 12 ≈ $421. Not Oct lump. Not Travelers/Geico/SF.",
        ],
        [
            "$150 Chris Brennan Exterminator 8/7/25",
            "LOCKED (payee confirm on 8507 stmt)",
            "INTERIOR MAINTENANCE August",
            "150.00",
            "8507 Check SERIAL posted 2025-08-18 $150.00. No Brennan/Orkin/Terminix string on Monarch.",
            "User date 8/7; check posted 8/18. Amount+timing lock. Confirm payee Chris Brennan when 8507 PDF arrives.",
        ],
        [
            "Smoke detector Amazon",
            "ASK — no SKU",
            "EQUIPMENT/APPLIANCE (yellow $0)",
            "0.00",
            "None labeled smoke/Kidde/First Alert/detector on Prime/Sapphire/BoA/Monarch 2025.",
            "Checklist was TBD. Do not invent an Amazon marketplace amount.",
        ],
        [
            "$300 cleaning to Imelda 9/30/25 empty apartment",
            "LOCKED",
            "MOVE OUT FEE October (turbo CLEANING October)",
            "300.00",
            "Joint 0203 Zelle IMELDA CORTEZ 2025-10-01 $300.00 Conf# nq057i1yf.",
            "User 9/30; cash 10/1 (1-day). 2023/24 MOVE OUT FEE was October $200. Other 2025 Imelda $150s NOT on 216 (occupied-LTR vs 524 STR).",
        ],
        [
            "$675 painting and repairs to Joan",
            "LOCKED",
            "INTERIOR REPAIR December",
            "675.00",
            "8507 Zelle to JOAN FRIEDBERGER 2025-12-01 $675.00.",
            "Moved off INTERIOR MAINTENANCE onto INTERIOR REPAIR (2024 used that row for Oak Park painting).",
        ],
        [
            "$11 key cutting",
            "LOCKED",
            "INTERIOR MAINTENANCE December",
            "11.00",
            "Megan Chase 1702 Ace Hardware NO. 152 2025-12-08 $11.00.",
            "Same-day Ace $17.00 NOT booked — ASK if also 216.",
        ],
        [
            "Other 2025 Imelda (not on this list)",
            "NOT on 216",
            "—",
            "0.00",
            "Recurring $150 (and some $80) Zelle on 0203/8507 during occupancy.",
            "Left off 216 unless you say those are Oak Park too. Pattern matches 524 STR turnover.",
        ],
    ]
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def write_start_here() -> None:
    (CSV_DIR / "START_HERE_oak_park_costs.txt").write_text(
        "START HERE — 216 N. Oak Park 2025 costs (2026-09-20)\n"
        "\n"
        "Not tax advice. Organizational packet only. Same tab as last year’s Personal Income.xlsx.\n"
        "\n"
        "Your list is on the 216 tab (2025 columns AC–AN):\n"
        "\n"
        "LOCKED\n"
        "1. Mortgage $1,226.25/mo — cash Jul–Oct $1,226.25 and Nov–Dec $1,177.52 on Checking 8507.\n"
        "   Jan–Jun yellow $1,226.25 until 8507 statements show the ACH (not on Monarch).\n"
        "   Nov–Dec kept at cash $1,177.52 (2023 Excel also stepped when the payment changed).\n"
        "   Sch E deducts Rocket 1098 interest, not full P+I. This row is cash like 2023/2024.\n"
        "2. HOA $421.53 — Santa Maria C326 on 8507 Jun–Dec. Jan–May yellow $421.53 WAIT 8507.\n"
        "3. Insurance $42.84/mo × 12 = $514.08 on RENTER'S INSURANCE.\n"
        "   Same amortized Lemonade as 2023 ($26.75) and 2024 ($35.08). Cash $514.00 on 10/3 Megan Chase (8¢).\n"
        "   Travelers/Geico/State Farm are not 216.\n"
        "4. Chris Brennan exterminator $150 8/7 — INTERIOR August.\n"
        "   Cash: 8507 check posted 8/18 $150 (confirm payee on the 8507 PDF).\n"
        "5. Imelda $300 empty-apartment 9/30 — MOVE OUT FEE October.\n"
        "   Cash: joint 0203 Zelle 10/1 $300. 2023/24 move-out fee was October $200.\n"
        "   Other Imelda $150s are not on 216.\n"
        "6. Joan $675 painting + repairs 12/1 — INTERIOR REPAIR December (8507 Zelle).\n"
        "7. Key cutting $11 — Ace Hardware 12/8 Megan Chase → INTERIOR December.\n"
        "\n"
        "ASK\n"
        "- Smoke detector Amazon: no labeled charge. EQUIPMENT stays $0 until you point at the order.\n"
        "- Ace Hardware $17 same day as the $11 keys — 216 or not?\n"
        "- Confirm 8507 check 8/18 $150 payee = Chris Brennan.\n"
        "- WAIT MEGAN: 8507 statements for Jan–Jun Rocket and Jan–May HOA (yellow on the sheet).\n"
        "\n"
        "Do not re-run apply_checking_answers_2025.py.\n",
        encoding="utf-8",
    )


def patch_ask() -> None:
    path = CSV_DIR / "ASK.csv"
    rows = list(csv.reader(path.open(encoding="utf-8")))
    ins_lock = [
        "Lemonade 216 LTR insurance $42.84/mo",
        "LOCKED 2026-09-20 (user cost list). $42.84 × 12 = $514.08 on 216 RENTER'S INSURANCE "
        "(amortized like 2023 $26.75 / 2024 $35.08). Lemonade cash $514.00 10/3 Megan Chase (8¢). "
        "Not October lump. Not Travelers/Geico/State Farm.",
    ]
    joan_lock = [
        "Joan 12/1 Zelle OUT $675 = painting + repairs",
        "LOCKED. Checking 8507 Zelle to Joan 12/1 $675 → 216 INTERIOR REPAIR (December). "
        "2024 used INTERIOR REPAIR for Oak Park painting. 12/1 $1,950 IN stays December occupancy cash.",
    ]
    extra_confirmed = [
        [
            "216 mortgage $1,226.25 / HOA $421.53",
            "LOCKED cash where 8507 shows it: Rocket Jul–Oct $1,226.25, Nov–Dec $1,177.52; "
            "HOA Jun–Dec $421.53. Jan–Jun mortgage and Jan–May HOA yellow at the user amounts WAIT 8507.",
        ],
        [
            "Chris Brennan exterminator $150 8/7",
            "LOCKED INTERIOR August. 8507 check posted 8/18 $150 (user date 8/7). Confirm payee on 8507 PDF.",
        ],
        [
            "Imelda $300 empty-apartment 9/30",
            "LOCKED MOVE OUT FEE October. Joint 0203 Zelle 10/1 $300 Conf# nq057i1yf (user 9/30). "
            "Other 2025 Imelda $150s not on 216.",
        ],
        [
            "Ace Hardware key cutting $11 12/8",
            "LOCKED INTERIOR December. Megan Chase Ace $11. Same-day Ace $17 ASK (not booked).",
        ],
    ]
    extra_open = [
        [
            "216 smoke detector Amazon",
            "ASK. No Kidde/First Alert/detector SKU on Monarch/Prime/Sapphire/BoA 2025. EQUIPMENT $0 until the order is pointed at.",
        ],
        [
            "Ace Hardware $17 12/8 same day as $11 keys",
            "Megan Chase 1702. Not booked on 216. Say if that is also Oak Park.",
        ],
        [
            "WAIT MEGAN — 8507 Jan–Jun Rocket / Jan–May HOA",
            "Yellow on 216 at user amounts ($1,226.25 and $421.53). Not on Monarch 8507 export. Do not treat yellow as cash until the PDF shows the ACH.",
        ],
    ]
    extra_totals = [
        ["216 RENTER'S INSURANCE ($42.84 × 12)", f"{float(INS_MO * 12):.2f}"],
        ["216 MORTGAGE sheet (yellow Jan–Jun + cash Jul–Dec)", f"{sum(OAK_MORTGAGE):.2f}"],
        ["216 HOA sheet (yellow Jan–May + cash Jun–Dec)", f"{sum(OAK_HOA):.2f}"],
        ["216 Brennan exterminator Aug", "150.0"],
        ["216 Imelda empty-apt MOVE OUT Oct", "300.0"],
        ["216 Joan painting/repairs Dec", "675.0"],
        ["216 Ace keys Dec", "11.0"],
        ["216 smoke detector Amazon (ASK)", "0.0"],
    ]

    out = []
    for row in rows:
        if row and row[0].startswith("Lemonade 216 LTR insurance"):
            out.append(ins_lock)
            continue
        if row and row[0].startswith("Joan 12/1 Zelle OUT $675"):
            out.append(joan_lock)
            continue
        if row and row[0].startswith("216 LTR insurance (Lemonade Oct cash"):
            continue  # replaced by extra_totals
        out.append(row)
        if row and row[0].startswith("Checking 8507 statements"):
            out.extend(extra_confirmed)
        if row and row[0].startswith("Joan half-November rent"):
            out.extend(extra_open)
        if row and row[0].startswith("216 INTERIOR MAINTENANCE Dec"):
            continue
    # totals: append before checking inflows block if present
    final = []
    inserted_totals = False
    for row in out:
        if (not inserted_totals) and row and row[0].startswith("Checking inflows after answers"):
            final.extend(extra_totals)
            inserted_totals = True
        final.append(row)
    if not inserted_totals:
        final.extend(extra_totals)
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(final)


def patch_start_here_insurance() -> None:
    path = CSV_DIR / "START_HERE_insurance.txt"
    text = path.read_text(encoding="utf-8")
    old = (
        "5. 216 N. Oak Park LTR insurance = Lemonade $514.00 cash 10/3 Megan Chase.\n"
        "   Booked October on 216 RENTER'S INSURANCE. Not Travelers, not Geico, not State Farm.\n"
    )
    new = (
        "5. 216 N. Oak Park LTR insurance = $42.84/mo × 12 = $514.08 on RENTER'S INSURANCE "
        "(Lemonade cash $514.00 10/3 Megan Chase, amortized like 2023/2024).\n"
    )
    if old in text:
        text = text.replace(old, new)
        path.write_text(text, encoding="utf-8")
    elif "42.84/mo" not in text:
        text = text.replace(
            "4. Grove Collaborative ≠ 827 N Grove the house (already locked).\n",
            "4. Grove Collaborative ≠ 827 N Grove the house (already locked).\n" + new,
        )
        path.write_text(text, encoding="utf-8")


def patch_start_here_checking() -> None:
    path = CSV_DIR / "START_HERE_checking.txt"
    text = path.read_text(encoding="utf-8")
    old = (
        "9. 216 LTR insurance LOCKED — Lemonade $514 cash 2025-10-03 Megan Chase → October "
        "RENTER'S INSURANCE. Travelers/State Farm/Geico are not 216.\n"
    )
    new = (
        "9. 216 costs LOCKED 2026-09-20 (last-year pattern): insurance $42.84/mo; "
        "mortgage $1,226.25 cash Jul–Oct / $1,177.52 Nov–Dec; HOA $421.53 Jun–Dec; "
        "Brennan $150 Aug; Imelda $300 Oct move-out; Joan $675 Dec painting; Ace keys $11 Dec. "
        "Jan–Jun mortgage and Jan–May HOA yellow WAIT 8507. Smoke detector Amazon ASK.\n"
    )
    if old in text:
        text = text.replace(old, new)
        path.write_text(text, encoding="utf-8")
    elif "insurance $42.84/mo" not in text and "216 costs LOCKED" not in text:
        needle = "Cash still on 216 Oak Park STR / RENTAL INCOME: $19,500"
        if needle in text:
            text = text.replace(needle, new + "\n" + needle)
            path.write_text(text, encoding="utf-8")


def patch_rebuild_comments() -> None:
    path = ROOT / "rebuild_like_prior_2025.py"
    text = path.read_text(encoding="utf-8")
    text = text.replace(
        "OAK_MORTGAGE = [0, 0, 0, 0, 0, 0, 1226.25, 1226.25, 1226.25, 1226.25, 1177.52, 1177.52]",
        "OAK_MORTGAGE = [1226.25, 1226.25, 1226.25, 1226.25, 1226.25, 1226.25, 1226.25, 1226.25, 1226.25, 1226.25, 1177.52, 1177.52]",
    )
    text = text.replace(
        "OAK_HOA = [0, 0, 0, 0, 0, 421.53, 421.53, 421.53, 421.53, 421.53, 421.53, 421.53]",
        "OAK_HOA = [421.53, 421.53, 421.53, 421.53, 421.53, 421.53, 421.53, 421.53, 421.53, 421.53, 421.53, 421.53]",
    )
    text = text.replace(
        "OAK_INTERIOR = [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 675]",
        "OAK_INTERIOR = [0, 0, 0, 0, 0, 0, 0, 150.0, 0, 0, 0, 11.0]  # Brennan Aug + Ace keys Dec",
    )
    text = text.replace(
        "OAK_INS = [0, 0, 0, 0, 0, 0, 0, 0, 0, 514, 0, 0]  # Lemonade 10/3 cash LOCKED 216 LTR",
        "OAK_INS = [42.84] * 12  # user $42.84/mo; Lemonade cash $514 10/3 amortized like 2023/2024",
    )
    path.write_text(text, encoding="utf-8")


def export_216_part(wb) -> None:
    out = Workbook()
    out.remove(out.active)
    src = wb["216 N. Oak Park Ave"]
    ws = out.create_sheet("216 N. Oak Park Ave")
    copy_sheet_light(src, ws)
    PARTS.mkdir(exist_ok=True)
    out.save(PARTS / "216 N. Oak Park Ave.xlsx")
    shutil.copy2(PARTS / "216 N. Oak Park Ave.xlsx", PARTS / "216 N Oak Park Ave.xlsx")
    ans.sheet_to_csv(src, CSV_DIR / "216_Oak_Park_like_last_year.csv")


def write_packet(wb) -> None:
    pkt = json.loads(PACKET.read_text(encoding="utf-8")) if PACKET.exists() else {}
    pkt["updated"] = "2026-09-20"
    pkt["xlsx_bytes_local"] = XLSX.stat().st_size
    pkt["ltr_insurance"] = {
        "carrier": "Lemonade",
        "property": "216 N. Oak Park Avenue #1Z",
        "locked": "2026-09-20",
        "method": "amortized $42.84/mo like 2023/2024 Excel (not October lump)",
        "monthly": 42.84,
        "sheet_total": float(INS_MO * 12),
        "cash": 514.0,
        "date": "2025-10-03",
        "account": "Megan Chase 1702",
        "rounding_vs_cash": 0.08,
    }
    pkt["oak_park_2025_costs"] = {
        "mortgage_sheet": sum(OAK_MORTGAGE),
        "mortgage_cash_jul_dec": 1226.25 * 4 + 1177.52 * 2,
        "hoa_sheet": sum(OAK_HOA),
        "hoa_cash_jun_dec": 421.53 * 7,
        "insurance": float(INS_MO * 12),
        "brennan_aug": 150.0,
        "imelda_moveout_oct": 300.0,
        "joan_repair_dec": 675.0,
        "ace_keys_dec": 11.0,
        "smoke_detector_amazon": 0.0,
        "smoke_detector_status": "ASK",
    }
    pkt["xlsx_note"] = (
        "2026-09-20 216 Oak Park costs locked to last-year pattern: insurance $42.84/mo; "
        "mortgage/HOA cash + yellow WAIT 8507; Brennan $150; Imelda $300 move-out; "
        "Joan $675 INTERIOR REPAIR; Ace keys $11. Smoke detector Amazon ASK. Not tax advice."
    )
    PACKET.write_text(json.dumps(pkt, indent=2), encoding="utf-8")


def main() -> None:
    assert XLSX.exists(), XLSX
    wb = load_workbook(XLSX)
    oak = wb["216 N. Oak Park Ave"]
    fill_xlsx(oak)
    wb.save(XLSX)
    if DELIVERABLE.parent.exists():
        shutil.copy2(XLSX, DELIVERABLE)
    export_turbo_csv()
    export_216_part(wb)
    write_lock_csv()
    write_start_here()
    patch_ask()
    patch_start_here_insurance()
    patch_start_here_checking()
    patch_rebuild_comments()
    write_packet(wb)

    ins = [oak.cell(14, 29 + i).value for i in range(12)]
    mort = [oak.cell(35, 29 + i).value for i in range(12)]
    hoa = [oak.cell(36, 29 + i).value for i in range(12)]
    interior = [oak.cell(15, 29 + i).value for i in range(12)]
    repair = [oak.cell(16, 29 + i).value for i in range(12)]
    move = [oak.cell(38, 29 + i).value for i in range(12)]
    print("ins", ins, "sum", round(sum(float(x or 0) for x in ins), 2))
    print("mort", mort, "sum", round(sum(float(x or 0) for x in mort), 2))
    print("hoa", hoa, "sum", round(sum(float(x or 0) for x in hoa), 2))
    print("interior", interior)
    print("repair", repair)
    print("moveout", move)
    print("saved", XLSX, XLSX.stat().st_size)


if __name__ == "__main__":
    main()

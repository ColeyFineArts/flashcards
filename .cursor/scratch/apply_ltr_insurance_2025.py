#!/usr/bin/env python3
"""Book 216 Oak Park LTR Lemonade insurance (user 2026-09-20).

Does NOT touch 524 INSURANCE (Travelers = Ferdinand home / Geico = cars /
State Farm home share only). Does NOT re-run apply_checking_answers_2025.py.

Hunt: Monarch CSV, Megan Chase 1702, Prime/Sapphire/BoA, checking 0203/8507,
216 CSV, HOME_SPLIT, apply_state_farm. Only Oak Park policy cash found is
Lemonade $514 on 2025-10-03 (same Oct 3 annual pattern as 2023 $321 and
2024 $421, which match last year’s 216 RENTER'S INSURANCE monthly totals).

Not tax advice.
"""
from __future__ import annotations

import csv
import json
import shutil
import sys
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path("/workspace/.cursor/scratch")
sys.path.insert(0, str(ROOT))
import apply_cc_2025 as base  # noqa: E402
import apply_cc_answers_2025 as ans  # noqa: E402

XLSX = ans.XLSX
DELIVERABLE = ans.DELIVERABLE
CSV_DIR = ans.CSV_DIR
PARTS = ans.PARTS
PACKET = ROOT / "cpa_packet_v1.json"
FULL_TABS = PARTS / "Personal Income 2025 — like last year (all tabs).xlsx"
LOCK_CSV = CSV_DIR / "LOCK_LTR_Lemonade.csv"
START_LTR = CSV_DIR / "START_HERE_ltr_insurance.txt"

GREEN = base.GREEN
YELLOW = base.YELLOW
PEACH = PatternFill("solid", fgColor="F7CAAC")
ACCT = '_("$"* #,##0.00_);_("$"* \\(#,##0.00\\);_("$"* "-"??_);_(@_)'
CG = "Century Gothic"

# Cash-basis 2025: Lemonade annual premium paid October (not amortized).
OAK_INS = [0, 0, 0, 0, 0, 0, 0, 0, 0, 514.0, 0, 0]
LEMONADE_DATE = "2025-10-03"
LEMONADE_AMT = 514.0
LEMONADE_ACCOUNT = "Megan Chase 1702"
NOTE = (
    "LOCKED 2026-09-20: 216 LTR insurance = Lemonade $514.00 cash 10/3 on Megan Chase (1702). "
    "Same annual Oct 3 Lemonade as 2023 $321 (= Excel $26.75/mo renter's) and 2024 $421 "
    "(≈ Excel $35.08/mo renter's). 2025 books cash in October (not amortized). "
    "Travelers = prior 524 Ferdinand home (not 216). Geico = cars. State Farm 2025 bundle = "
    "Ferdinand car+home — only home share on 524, not 216. No other Oak Park policy cash on "
    "Prime/Sapphire/BoA/0203/8507/Monarch. Not tax advice."
)

FULL_TAB_NAMES = [
    "Income",
    "216 N. Oak Park Ave",
    "524 Ferdinand Ave, Unit 2",
    "524 Ferdinand Ave, Unit 1",
    "EPGC LLC",
    "Art Sales and Purchases",
    "GCM",
    "Investments",
    "ASK Mercury",
    "MERCURY 8291",
]


def copy_sheet_light(src: Worksheet, dest: Worksheet, max_col: int | None = None) -> None:
    mc = min(src.max_column or 1, max_col or 42)
    mr = src.max_row or 1
    for r in src.iter_rows(min_row=1, max_row=mr, max_col=mc):
        for cell in r:
            d = dest.cell(cell.row, cell.column, cell.value)
            if cell.has_style:
                d.font = copy(cell.font)
                d.fill = copy(cell.fill)
                d.alignment = copy(cell.alignment)
                d.number_format = cell.number_format
    for col, dim in src.column_dimensions.items():
        dest.column_dimensions[col].width = dim.width
    for rng in src.merged_cells.ranges:
        try:
            dest.merge_cells(str(rng))
        except Exception:
            pass
    if src.sheet_properties.tabColor:
        dest.sheet_properties.tabColor = src.sheet_properties.tabColor.rgb


def write_2025_ins(ws: Worksheet) -> None:
    """Like-last-year 216: 2025 block starts at column 29 (AC). Row 14 = RENTER'S INSURANCE."""
    assert str(ws["B14"].value).upper().find("INSUR") >= 0, ws["B14"].value
    ac = 29
    for i, v in enumerate(OAK_INS):
        cell = ws.cell(14, ac + i)
        cell.value = float(v)
        cell.number_format = ACCT
        if v:
            cell.fill = GREEN
    ws["B64"] = NOTE
    ws["B64"].alignment = Alignment(wrap_text=True)
    ws["B64"].font = Font(name=CG, size=10)
    ws["B64"].fill = GREEN
    ws.row_dimensions[64].height = 64


def export_216_csv(ws: Worksheet) -> None:
    """Turbo-style 2025-only CSV used by the checking-answers 216 sheet."""
    path = CSV_DIR / "216_Oak_Park.csv"
    rows = list(csv.reader(path.open(encoding="utf-8")))
    for row in rows:
        if len(row) > 1 and row[1] == "INSURANCE":
            # cols: 0 empty, 1 label, 2-13 months
            for i, v in enumerate(OAK_INS):
                row[2 + i] = str(v)
        if len(row) > 1 and isinstance(row[1], str) and row[1].startswith("LTR after answers"):
            row[1] = (
                "LTR after answers 2026-09-18 + insurance lock 2026-09-20: "
                "Ava CONFIRMED Delach co-tenant — keep $500/mo. TY2025 cash $19500.00 "
                "(Delach $11,600 + Ava $4,000 + Joan $3,900). Occupancy vacant October and half of November. "
                "Joan half-Nov $975 NOT booked — WAIT 8507. INTERIOR Dec $675 Joan repair. "
                "INSURANCE LOCKED: Lemonade $514 cash 2025-10-03 Megan Chase → October RENTER'S INSURANCE. "
                "Travelers/State Farm/Geico are not 216. HOA + Rocket cash from Monarch. Not tax advice."
            )
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def patch_ask() -> None:
    path = CSV_DIR / "ASK.csv"
    rows = list(csv.reader(path.open(encoding="utf-8")))
    lock_row = [
        "Lemonade 216 LTR insurance",
        "LOCKED 2026-09-20. Megan Chase 1702 2025-10-03 $514. Same Oct 3 Lemonade as "
        "2023 $321 (= Excel $26.75/mo renter's) and 2024 $421 (≈ Excel $35.08/mo). "
        "Booked cash October on 216 RENTER'S INSURANCE. Travelers = 524 home, Geico = cars, "
        "State Farm home share = 524 only — none of those are 216. No other Oak Park policy "
        "cash on Prime/Sapphire/BoA/0203/8507/Monarch.",
    ]
    total_row = ["216 LTR insurance (Lemonade Oct cash, locked)", "514.0"]
    # insert lock after Travelers confirmed row; insert total after 524 insurance totals
    out = []
    inserted_lock = False
    inserted_total = False
    for row in rows:
        out.append(row)
        if (
            not inserted_lock
            and row
            and row[0].startswith("Travelers prior 524 home")
        ):
            out.append(lock_row)
            inserted_lock = True
        if (
            not inserted_total
            and row
            and row[0].startswith("524 INSURANCE home share (Unit 1")
        ):
            out.append(total_row)
            inserted_total = True
    if not inserted_lock:
        # put after CONFIRMED header block
        out.append(lock_row)
    if not inserted_total:
        out.append(total_row)
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(out)


def write_lock_csv() -> None:
    rows = [
        ["Item", "Status", "Where", "Dollars / months"],
        [
            "Lemonade = 216 Oak Park LTR renter's / HO-6",
            "LOCKED 2026-09-20",
            "216 N. Oak Park Ave #1Z RENTER'S INSURANCE (2025 cash October)",
            "$514.00 on 2025-10-03 Megan Chase 1702",
        ],
        [
            "2023 Lemonade (identity check)",
            "MATCHES 2023 Excel renter's",
            "Same merchant, same Megan Chase, Oct 3",
            "$321.00 = 12 × $26.75",
        ],
        [
            "2024 Lemonade (identity check)",
            "MATCHES 2024 Excel renter's",
            "Same merchant, same Megan Chase, Oct 3",
            "$421.00 ≈ 12 × $35.08",
        ],
        [
            "Travelers / Geico / State Farm",
            "NOT 216",
            "Travelers leftover = 524 home $0 cash; Geico = cars; State Farm home share = 524 only",
            "Do not book on 216",
        ],
        [
            "Other 2025 Oak Park policy cash",
            "NONE FOUND",
            "Searched Monarch TY2025, Prime 2351, Sapphire 5423, BoA 4469, checking 0203/8507, 216 CSV, HOME_SPLIT, apply_state_farm",
            "$0.00 besides Lemonade $514",
        ],
    ]
    with LOCK_CSV.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)


def write_start_here_ltr() -> None:
    START_LTR.write_text(
        "START HERE — 216 Oak Park LTR insurance (2026-09-20)\n"
        "\n"
        "Not tax advice. Organizational packet only.\n"
        "\n"
        "LOCKED: Lemonade $514.00 cash 2025-10-03 on Megan Chase (1702) is the 216 N. Oak Park "
        "Avenue #1Z (LTR) insurance expense. Booked on 216 RENTER'S INSURANCE in October.\n"
        "\n"
        "Why this is 216, not 524\n"
        "- Last year’s 216 tab called the row RENTER'S INSURANCE and amortized Lemonade: "
        "2023 $26.75/mo × 12 = $321 (Monarch Lemonade 2023-10-03 $321); "
        "2024 $35.08/mo × 12 ≈ $421 (Monarch Lemonade 2024-10-03 $421).\n"
        "- 2025 is the same merchant, same card, same Oct 3 date, $514. This packet books "
        "cash in the month paid (October), not amortized.\n"
        "- Travelers was home for Ferdinand (524). Geico was the cars. State Farm 2025 is "
        "the Ferdinand car+home bundle — only the home share is on 524, not 216.\n"
        "\n"
        "Searched (no other Oak Park policy dollars): Monarch TY2025, Prime/Sapphire/BoA, "
        "checking 0203 and 8507, 216 CSV, HOME_SPLIT, apply_state_farm.\n"
        "\n"
        "524 insurance lock is unchanged (Travelers = 524 home, Geico = cars).\n"
        "Do not re-run apply_checking_answers_2025.py.\n",
        encoding="utf-8",
    )


def patch_start_here_insurance() -> None:
    path = CSV_DIR / "START_HERE_insurance.txt"
    extra = (
        "\n"
        "5. 216 N. Oak Park LTR insurance = Lemonade $514.00 cash 10/3 Megan Chase.\n"
        "   Booked October on 216 RENTER'S INSURANCE. Not Travelers, not Geico, not State Farm.\n"
    )
    text = path.read_text(encoding="utf-8")
    if "216 N. Oak Park LTR insurance" not in text:
        text = text.replace(
            "4. Grove Collaborative ≠ 827 N Grove the house (already locked).\n",
            "4. Grove Collaborative ≠ 827 N Grove the house (already locked)." + extra,
        )
        path.write_text(text, encoding="utf-8")


def patch_start_here_checking() -> None:
    path = CSV_DIR / "START_HERE_checking.txt"
    text = path.read_text(encoding="utf-8")
    needle = "Cash still on 216 Oak Park STR / RENTAL INCOME: $19,500"
    add = (
        "9. 216 LTR insurance LOCKED — Lemonade $514 cash 2025-10-03 Megan Chase → October "
        "RENTER'S INSURANCE. Travelers/State Farm/Geico are not 216.\n"
        "\n"
        "Cash still on 216 Oak Park STR / RENTAL INCOME: $19,500"
    )
    if "216 LTR insurance LOCKED" not in text:
        text = text.replace(needle, add)
        path.write_text(text, encoding="utf-8")


def build_full_tabs(wb) -> Path:
    pack = Workbook()
    pack.remove(pack.active)
    for name in FULL_TAB_NAMES:
        if name not in wb.sheetnames:
            continue
        src = wb[name]
        ws = pack.create_sheet(name)
        copy_sheet_light(src, ws)
    FULL_TABS.parent.mkdir(exist_ok=True)
    pack.save(FULL_TABS)
    return FULL_TABS


def export_216_part(wb) -> None:
    from openpyxl import Workbook as WB

    out = WB()
    out.remove(out.active)
    src = wb["216 N. Oak Park Ave"]
    ws = out.create_sheet("216 N. Oak Park Ave")
    copy_sheet_light(src, ws)
    out.save(PARTS / "216 N. Oak Park Ave.xlsx")
    # also the git-tracked name without the extra dot
    shutil.copy2(PARTS / "216 N. Oak Park Ave.xlsx", PARTS / "216 N Oak Park Ave.xlsx")


def patch_rebuild_comments() -> None:
    path = ROOT / "rebuild_like_prior_2025.py"
    text = path.read_text(encoding="utf-8")
    if "GREEN = PatternFill" not in text:
        text = text.replace(
            'YELLOW = PatternFill("solid", fgColor="FFF2CC")\n',
            'YELLOW = PatternFill("solid", fgColor="FFF2CC")\nGREEN = PatternFill("solid", fgColor="C6EFCE")\n',
        )
    text = text.replace(
        "  RENTER'S INSURANCE Oct Lemonade $514 (confirm policy)",
        "  RENTER'S INSURANCE Oct Lemonade $514 LOCKED 216 LTR (cash 10/3 Megan Chase)",
    )
    text = text.replace(
        "OAK_INS = [0, 0, 0, 0, 0, 0, 0, 0, 0, 514, 0, 0]  # Lemonade 10/3 cash",
        "OAK_INS = [0, 0, 0, 0, 0, 0, 0, 0, 0, 514, 0, 0]  # Lemonade 10/3 cash LOCKED 216 LTR",
    )
    text = text.replace(
        "    write_months(ws, 14, ac, OAK_INS, fill=YELLOW)  # Lemonade — confirm 216 policy",
        "    write_months(ws, 14, ac, OAK_INS, fill=GREEN)  # Lemonade LOCKED 216 LTR",
    )
    text = text.replace(
        '    ws["B64"] = "Insurance — Lemonade $514 cash 10/3 on Megan Chase (yellow: confirm 216 policy). 2024 was $35.08/mo renter\'s."',
        '    ws["B64"] = "LOCKED 2026-09-20: Lemonade $514 cash 10/3 Megan Chase is 216 LTR insurance (Oct cash, not amortized). 2023 $321 / 2024 $421 same Oct 3 pattern. Travelers/SF/Geico are not 216."',
    )
    path.write_text(text, encoding="utf-8")


def main() -> None:
    assert XLSX.exists(), XLSX
    wb = load_workbook(XLSX)
    # Do not clobber 524 Travelers/Geico lock — only 216.
    oak = wb["216 N. Oak Park Ave"]
    write_2025_ins(oak)
    wb.save(XLSX)
    shutil.copy2(XLSX, DELIVERABLE)

    export_216_csv(oak)
    export_216_part(wb)
    full = build_full_tabs(wb)
    patch_ask()
    write_lock_csv()
    write_start_here_ltr()
    patch_start_here_insurance()
    patch_start_here_checking()
    patch_rebuild_comments()

    if PACKET.exists():
        pkt = json.loads(PACKET.read_text(encoding="utf-8"))
        pkt["updated"] = "2026-09-20"
        pkt["ltr_insurance"] = {
            "carrier": "Lemonade",
            "property": "216 N. Oak Park Avenue #1Z",
            "locked": "2026-09-20",
            "cash": LEMONADE_AMT,
            "date": LEMONADE_DATE,
            "account": LEMONADE_ACCOUNT,
            "months": {"OCT": LEMONADE_AMT},
            "method": "2025 cash in October (not amortized). 2023/2024 Excel amortized the same Lemonade annual premium.",
            "not_216": ["Travelers (524 home)", "Geico (cars)", "State Farm home share (524 only)"],
        }
        pkt["xlsx_bytes_local"] = XLSX.stat().st_size
        PACKET.write_text(json.dumps(pkt, indent=2), encoding="utf-8")

    print("xlsx", XLSX, XLSX.stat().st_size)
    print("216 Oct AL14", oak["AL14"].value, "B64 locked")
    print("full tabs", full, full.stat().st_size, "sheets", load_workbook(full).sheetnames)


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""Apply 2025 BoA checking inflows onto Personal Income 2025 Tax Turbo.

Sources:
  folder 07 / bank of america checking_JSC  = Adv Plus 4830 6191 9922 (Jake)
  folder 07 / Bank of America checking_joint = Adv Plus 4830 7359 0203 (Megan+Jacob)
Reconciled to Monarch CSV on those two accounts. Cash-basis by deposit date.

User rules 2026-09-18:
  GCM BOM / board fee → EPGC Consultant (and GCM tab). Watch reimbursements.
  Joan Friedberger + JOHN T DELACH → 216 Oak Park STR / RENTAL INCOME.
  Bear Creek annual gift → not income.
  Emma / Phoebe UTMA ~$10k each → not income.

Do not re-run apply_cc_answers_2025.py (would double HD). Not tax advice.
"""
from __future__ import annotations

import json
import shutil
import sys
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path("/workspace/.cursor/scratch")
sys.path.insert(0, str(ROOT))
import apply_cc_2025 as base  # noqa: E402
import apply_cc_answers_2025 as ans  # noqa: E402

XLSX = ans.XLSX
DELIVERABLE = ans.DELIVERABLE
JSON_PATH = ROOT / "checking_inflows_2025.json"

GREEN = base.GREEN
YELLOW = base.YELLOW
BLUE = base.BLUE
GRAY = base.GRAY
HEADER_FILL = base.HEADER_FILL
HEADER_FONT = base.HEADER_FONT
THIN = base.THIN
ORANGE = PatternFill("solid", fgColor="FCE4D6")
MONTHS = base.MONTHS
PROP_COLS = base.PROP_COLS
EPGC_COLS = base.EPGC_COLS
D = base.D
money = base.money
zeros = base.zeros
add = base.add


# --- 216 N. Oak Park LTR rent (cash-basis, joint 0203) ---
# Delach $1,450 + Ava Hernandez $500 = $1,950 (same as Joan later).
# Jan 2025 occupancy was paid 2024-12-31 — excluded (not TY2025 cash).
oak_rent = zeros()
# month index 0=JAN .. 11=DEC
add(oak_rent, 0, "1450")  # 2025-01-31 Delach February rent
add(oak_rent, 0, "500")   # 2025-01-31 Ava
add(oak_rent, 1, "1450")  # 2025-02-28 Delach March
add(oak_rent, 1, "500")
add(oak_rent, 2, "1450")  # 2025-03-31 Delach April
add(oak_rent, 2, "500")
add(oak_rent, 3, "1450")  # 2025-04-30 Delach May
add(oak_rent, 3, "500")   # 2025-04-28 Ava May
add(oak_rent, 4, "1450")  # 2025-05-30 Delach June
add(oak_rent, 4, "500")
add(oak_rent, 5, "1450")  # 2025-06-30 Delach July
add(oak_rent, 5, "500")   # 2025-06-23 Ava July
add(oak_rent, 6, "1450")  # 2025-07-31 Delach August
add(oak_rent, 7, "500")   # 2025-08-01 Ava August
add(oak_rent, 7, "1450")  # 2025-08-29 Delach September
add(oak_rent, 7, "500")   # 2025-08-29 Ava September
add(oak_rent, 11, "1950")  # 2025-12-01 Joan "rent" (Dec occupancy — ASK)
add(oak_rent, 11, "1950")  # 2025-12-31 Joan "January rent" (2026 prepaid cash)

OAK_RENT_TOTAL = sum(oak_rent, D("0"))
assert OAK_RENT_TOTAL == D("19500.00"), OAK_RENT_TOTAL

# --- GCM board fees (BOM PAYMEN) = 1099-NEC $21,500 ---
gcm_fee = zeros()
add(gcm_fee, 1, "4100")   # 2025-02-12
add(gcm_fee, 3, "9200")   # 2025-04-18
add(gcm_fee, 8, "4100")   # 2025-09-13
add(gcm_fee, 10, "4100")  # 2025-11-13
GCM_FEE_TOTAL = sum(gcm_fee, D("0"))
assert GCM_FEE_TOTAL == D("21500.00"), GCM_FEE_TOTAL

# Reimbursements — EXCLUDED from consultant
GCM_REIMB_9922 = [
    ("2025-02-18", D("1150.99"), "9922", "PMT INFO: BOARD MEETING EXPENSES"),
    ("2025-05-22", D("609.02"), "9922", "JAKE C — no PMT INFO on statement; odd cents, not BOM — treated as reimburse"),
    ("2025-11-05", D("1273.82"), "9922", "Monarch tag Reimburse; statement JAKE C (not BOM)"),
]
GCM_REIMB_8507 = [
    ("2025-07-25", D("462.48"), "8507", "SEFOF WHARTON EXPENSES — not on the two BoA folders"),
    ("2025-07-31", D("1254.93"), "8507", "SEFOF ANNUAL CONFERENCE EXPENSES"),
    ("2025-11-05", D("463.45"), "8507", "SEFOF REIMBURSE EXPENSES FOR EDUCATION COMMITTEE"),
]
REIMB_9922 = sum((a[1] for a in GCM_REIMB_9922), D("0"))
REIMB_8507 = sum((a[1] for a in GCM_REIMB_8507), D("0"))
REIMB_TOTAL = REIMB_9922 + REIMB_8507

BEAR_CREEK = [
    ("2025-01-09", D("9500.00"), "BEAR CREEK INC - CLIENT W / JPMorgan; PMT DET ATS OF 25/01/07"),
    ("2025-04-10", D("9500.00"), "BEAR CREEK INC.-CASH / Northern Trust"),
    ("2025-07-10", D("9500.00"), "BEAR CREEK INC.-CASH / Northern Trust"),
    ("2025-10-09", D("9500.00"), "BEAR CREEK INC.-CASH / Northern Trust"),
]
BEAR_CREEK_TOTAL = sum((a[1] for a in BEAR_CREEK), D("0"))
assert BEAR_CREEK_TOTAL == D("38000.00")

UTMA_GIFTS = [
    ("2025-11-28", D("5000.00"), "CMA-Edge 4729", "WIRE TRF IN ORG THOMAS — Phoebe UTMA (ASK which child)"),
    ("2025-12-02", D("5000.00"), "CMA-Edge 2T48", "WIRE TRF IN ORG THOMAS as-of 11/28 — Emma UTMA 2T48"),
]
UTMA_GIFT_TOTAL = sum((a[1] for a in UTMA_GIFTS), D("0"))

# LTR rent txn list (for ledger)
RENT_TXNS = [
    ("2025-01-31", "JOHN T DELACH", D("1450"), "February rent", "APPLIED", "LTR 216 Oak Park"),
    ("2025-01-31", "AVA L HERNANDEZ", D("500"), "Rent (same day as Delach)", "APPLIED-YELLOW", "Roommate/co-tenant? ASK"),
    ("2025-02-28", "JOHN T DELACH", D("1450"), "March Rent", "APPLIED", "LTR 216 Oak Park"),
    ("2025-02-28", "AVA L HERNANDEZ", D("500"), "Rent", "APPLIED-YELLOW", "Roommate/co-tenant? ASK"),
    ("2025-03-31", "JOHN T DELACH", D("1450"), "April Rent", "APPLIED", "LTR 216 Oak Park"),
    ("2025-03-31", "AVA L HERNANDEZ", D("500"), "April Rent", "APPLIED-YELLOW", "Roommate/co-tenant? ASK"),
    ("2025-04-28", "AVA L HERNANDEZ", D("500"), "May Rent", "APPLIED-YELLOW", "Roommate/co-tenant? ASK"),
    ("2025-04-30", "JOHN T DELACH", D("1450"), "May Rent", "APPLIED", "LTR 216 Oak Park"),
    ("2025-05-30", "JOHN T DELACH", D("1450"), "June rent", "APPLIED", "LTR 216 Oak Park"),
    ("2025-05-30", "AVA L HERNANDEZ", D("500"), "June rent", "APPLIED-YELLOW", "Roommate/co-tenant? ASK"),
    ("2025-06-23", "AVA L HERNANDEZ", D("500"), "july rent", "APPLIED-YELLOW", "Roommate/co-tenant? ASK"),
    ("2025-06-30", "JOHN T DELACH", D("1450"), "july rent", "APPLIED", "LTR 216 Oak Park"),
    ("2025-07-31", "JOHN T DELACH", D("1450"), "august rent", "APPLIED", "LTR 216 Oak Park"),
    ("2025-08-01", "AVA L HERNANDEZ", D("500"), "August rent", "APPLIED-YELLOW", "Roommate/co-tenant? ASK"),
    ("2025-08-29", "JOHN T DELACH", D("1450"), "September Rent", "APPLIED", "LTR 216 Oak Park"),
    ("2025-08-29", "AVA L HERNANDEZ", D("500"), "september rent", "APPLIED-YELLOW", "Roommate/co-tenant? ASK"),
    ("2025-12-01", "JOAN FRIEDBERGER", D("1950"), "rent", "APPLIED-YELLOW", "Dec occupancy vs prepaid? ASK"),
    ("2025-12-31", "JOAN FRIEDBERGER", D("1950"), "January rent", "APPLIED-YELLOW", "Jan 2026 prepaid cash in TY2025. ASK"),
]
assert sum((t[2] for t in RENT_TXNS), D("0")) == OAK_RENT_TOTAL

CHECKING_CONFIRMED = [
    (
        "BoA checking folders = 9922 + 0203",
        "JSC Adv Plus 4830 6191 9922 (Jake) and joint Adv Plus 4830 7359 0203 (Megan+Jacob). "
        "12 statements each for 2025. Monarch credits on those two accounts match the statement deposit sections sampled.",
    ),
    (
        "GCM board fee → Consultant",
        f"Gates Capital BOM PAYMEN ${GCM_FEE_TOTAL}: 2/12 $4,100 + 4/18 $9,200 + 9/13 $4,100 + 11/13 $4,100. "
        "Equals Form 1099-NEC $21,500. Booked EPGC Consultant (cash-basis month) and GCM tab. "
        "Income tab GCM actual was already $21,500.",
    ),
    (
        "GCM reimbursements EXCLUDED",
        f"JAKE C / SEFOF on 9922 ${REIMB_9922} (board meeting $1,150.99; 5/22 $609.02 unlabeled; 11/05 $1,273.82). "
        f"Plus Checking 8507 SEFOF ${REIMB_8507} (Wharton / annual conference / education committee) — that account is not in the two folders. "
        f"Total reimbursement ${REIMB_TOTAL} kept off Consultant.",
    ),
    (
        "LTR rent → 216 Oak Park STR / RENTAL INCOME",
        f"JOHN T DELACH 8×$1,450 (${D('11600')}) + Joan Friedberger 2×$1,950 (${D('3900')}) on joint 0203. "
        f"Ava L Hernandez 8×$500 (${D('4000')}) same rent memo / $1,950 combined — included yellow (ASK). "
        f"TY2025 cash total ${OAK_RENT_TOTAL}. Oct–Nov 2025: no rent on 0203.",
    ),
    (
        "Bear Creek not income",
        f"Four wires $9,500 on Jake 9922 (1/9, 4/10, 7/10, 10/9) = ${BEAR_CREEK_TOTAL}. "
        "Annual gift — excluded from Consultant / rental / Income tab.",
    ),
    (
        "UTMA not income",
        f"Emma CMA 2T48 + sibling CMA 4729: Thomas wires $5,000 each Nov/Dec = ${UTMA_GIFT_TOTAL} (not ~$10k each). "
        "Treasury-note maturities $10,000 each 12/31 are principal, not gifts. Excluded from income.",
    ),
]

CHECKING_STILL_OPEN = [
    (
        "Ava L Hernandez $500/mo rent",
        "Eight Zelle credits on joint 0203, same days as Delach, memo “rent”. "
        "Delach $1,450 + Ava $500 = $1,950, which is Joan’s later rent. "
        "Is Ava Delach’s roommate/co-tenant at 216 Oak Park? I booked her on STR / RENTAL INCOME (yellow) so the $1,950 unit rent is complete. "
        "Say if she should be excluded.",
    ),
    (
        "Oak Park vacant Oct–Nov 2025?",
        "No Delach / Ava / Joan rent on 0203 in October or November. Last Delach/Ava is 8/29 “September Rent”. "
        "Joan starts 12/1. Confirm those two months were vacant (and held for rent) for Sch E. "
        "Jan 2025 occupancy cash hit 12/31/2024 ($1,450+$500) — left out of TY2025 cash. Correct?",
    ),
    (
        "Joan Friedberger Dec cash",
        "12/1 $1,950 memo “rent” and 12/31 $1,950 memo “January rent” both in TY2025 cash. "
        "Is 12/1 December occupancy and 12/31 January 2026 prepaid? "
        "Also: Checking 8507 Zelle OUT $675 to Joan on 12/1 — deposit return / repair? Not on these two BoA folders.",
    ),
    (
        "GCM on EPGC Consultant vs personal 1099",
        "You said board fee → consultant line, so I put BOM $21,500 on EPGC LLC Consultant. "
        "Historically GCM lived on its own tab + Income “GCM Boards”, and EPGC Consultant was empty. "
        "1099-NEC is in Jacob’s name from Gates Capital. Confirm: EPGC LLC (this booking) vs personal Sch C / GCM tab only "
        "(then I will zero EPGC Consultant so LLC P&L is not mixed).",
    ),
    (
        "May 22 Gates $609.02",
        "Statement says GATES CAPITAL MA DES:JAKE C $609.02 — no PMT INFO. Same DES as reimbursements, not BOM PAYMEN. "
        "I excluded it from Consultant. Confirm it is reimbursement (or board fee if you say so).",
    ),
    (
        "Bear Creek $38,000 vs kids’ UTMA ~$10k each",
        "Bear Creek to Jake 9922 is $9,500×4 = $38,000 (not ~$10k). Thomas wires into the two CMA-Edge UTMAs are $5,000 each in late Nov. "
        "Are Bear Creek gifts to you (parents), and Thomas the kids’ UTMA? Is more UTMA cash coming that never hit these checkings "
        "(e.g. directly at Merrill)? 2/25 $36,000 9922→National Financial was reversed 2/27.",
    ),
    (
        "Freeman’s / CIBC $11,000 wire 12/24 on joint 0203",
        "ORIG:FREEMAN'S LLC via CIBC Bank USA; PMT DET Sale 6428 Contract 303468. Not W-2 payroll (those are DIR DEP on 9922). "
        "Art sale? Consignment? Other? HOLD off income until you say.",
    ),
    (
        "Checking 8507 (NOW) 2025 statements?",
        "Third checking in Monarch — Airbnb/VRBO STR, Megan Zelle, and the three GCM SEFOF reimbursements. "
        "Not in the two folders you just added. Upload 2025 statements if you want that account reconciled the same way.",
    ),
]


def month_idx(iso: str) -> int:
    return int(iso[5:7]) - 1


def rebuild_gcm_tab(wb) -> None:
    if "GCM" in wb.sheetnames:
        idx = wb.sheetnames.index("GCM")
        del wb["GCM"]
        ws = wb.create_sheet("GCM", idx)
    else:
        ws = wb.create_sheet("GCM")
    ws["A1"] = "GCM / Gates Capital 2025 — board fees vs reimbursements (not tax advice)"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:N1")
    ws["A2"] = (
        f"BOM PAYMEN ${GCM_FEE_TOTAL} = Form 1099-NEC $21,500 → EPGC Consultant + this tab. "
        f"JAKE C / SEFOF reimbursements ${REIMB_TOTAL} EXCLUDED. "
        "CPA: confirm EPGC LLC vs personal 1099/Sch C."
    )
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:N2")
    ws.row_dimensions[2].height = 36

    headers = ["LINE", *MONTHS, "YR TOTAL"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(4, i, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN
    ws.cell(5, 1, "Board fee (BOM PAYMEN → Consultant)")
    for i, m in enumerate(MONTHS):
        cell = ws.cell(5, i + 2, money(gcm_fee[i]))
        cell.number_format = "0.00"
        cell.border = THIN
        if gcm_fee[i]:
            cell.fill = GREEN
    tot = ws.cell(5, 14, "=SUM(B5:M5)")
    tot.number_format = "0.00"
    tot.fill = GREEN
    tot.border = THIN
    ws.cell(6, 1, "Reimbursements (EXCLUDED from income)")
    for i in range(12):
        cell = ws.cell(6, i + 2, 0)
        cell.number_format = "0.00"
        cell.border = THIN
        cell.fill = GRAY
    # put reimburse in cash months for visibility but gray
    reimb_m = zeros()
    for dt, amt, *_ in GCM_REIMB_9922 + GCM_REIMB_8507:
        add(reimb_m, month_idx(dt), amt)
    for i, m in enumerate(MONTHS):
        cell = ws.cell(6, i + 2, money(reimb_m[i]))
        cell.number_format = "0.00"
        cell.border = THIN
        cell.fill = GRAY
    cell = ws.cell(6, 14, "=SUM(B6:M6)")
    cell.number_format = "0.00"
    cell.fill = GRAY
    cell.border = THIN

    ws.cell(8, 1, "Transaction register")
    ws.cell(8, 1).font = Font(bold=True)
    hdr = ["Date", "Account", "Amount", "Bucket", "Status", "Memo"]
    for i, h in enumerate(hdr, 1):
        cell = ws.cell(9, i, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN
    rows = [
        ("2025-02-12", "9922", D("4100"), "Consultant / 1099", "APPLIED", "BOM PAYMEN"),
        ("2025-04-18", "9922", D("9200"), "Consultant / 1099", "APPLIED", "BOM PAYMEN"),
        ("2025-09-13", "9922", D("4100"), "Consultant / 1099", "APPLIED", "BOM PAYMEN"),
        ("2025-11-13", "9922", D("4100"), "Consultant / 1099", "APPLIED", "BOM PAYMEN"),
    ]
    for dt, amt, acct, memo in GCM_REIMB_9922 + GCM_REIMB_8507:
        rows.append((dt, acct, amt, "Reimbursement", "EXCLUDED", memo))
    fills = {"APPLIED": GREEN, "EXCLUDED": GRAY}
    for i, row in enumerate(rows, 10):
        for c, val in enumerate(row, 1):
            cell = ws.cell(i, c, money(val) if isinstance(val, Decimal) else val)
            cell.border = THIN
            cell.fill = fills[row[4]]
            if c == 3:
                cell.number_format = "0.00"
    last = 9 + len(rows)
    ws.cell(last + 2, 1, "1099-NEC Gates Capital Management LLC TIN 84-1331261 box 1 = $21,500. BOM sum ties. Reimbursements are not on the 1099.")
    ws.cell(last + 2, 1).fill = BLUE
    ws.merge_cells(start_row=last + 2, start_column=1, end_row=last + 2, end_column=6)
    widths = [22, 14, 14, 22, 14, 80]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for i in range(7, 15):
        ws.column_dimensions[get_column_letter(i)].width = 12


def rebuild_inflow_ledger(wb) -> None:
    if "INFLOW_LEDGER" in wb.sheetnames:
        del wb["INFLOW_LEDGER"]
    # place after ASK
    idx = wb.sheetnames.index("ASK") + 1 if "ASK" in wb.sheetnames else 2
    led = wb.create_sheet("INFLOW_LEDGER", idx)
    led["A1"] = "2025 checking inflows — JSC 9922 + joint 0203 (folder 07) — not tax advice"
    led["A1"].font = Font(bold=True, size=14)
    led.merge_cells("A1:H1")
    led["A2"] = (
        "Cash-basis by deposit date. BOM → Consultant. JAKE C / SEFOF → reimbursement EXCLUDED. "
        "Delach / Ava / Joan → 216 Oak Park STR / RENTAL INCOME. Bear Creek + UTMA Thomas wires → not income. "
        "Hindman/Freeman payroll stays on the W-2 tab. EPGC Mercury ACH into 9922 is a transfer, not income."
    )
    led["A2"].alignment = Alignment(wrap_text=True)
    led.merge_cells("A2:H2")
    led.row_dimensions[2].height = 40
    headers = ["Account", "Date", "Counterparty", "Amount", "Bucket", "Tab", "Status", "Note"]
    for i, h in enumerate(headers, 1):
        cell = led.cell(4, i, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = THIN
    status_fill = {
        "APPLIED": GREEN,
        "APPLIED-YELLOW": YELLOW,
        "EXCLUDED": GRAY,
        "HOLD": YELLOW,
        "TRANSFER": BLUE,
    }
    rows = []
    for dt, who, amt, memo, status, note in RENT_TXNS:
        rows.append(("0203 joint", dt, who, float(amt), "STR / RENTAL INCOME", "216 N. Oak Park Ave", status, f"{memo}. {note}"))
    for dt, amt, acct, memo in [
        ("2025-02-12", D("4100"), "9922 JSC", "BOM PAYMEN"),
        ("2025-04-18", D("9200"), "9922 JSC", "BOM PAYMEN"),
        ("2025-09-13", D("4100"), "9922 JSC", "BOM PAYMEN"),
        ("2025-11-13", D("4100"), "9922 JSC", "BOM PAYMEN"),
    ]:
        rows.append((acct, dt, "Gates Capital", float(amt), "Consultant", "EPGC LLC + GCM", "APPLIED", memo))
    for dt, amt, acct, memo in GCM_REIMB_9922 + GCM_REIMB_8507:
        label = "9922 JSC" if acct == "9922" else "8507 NOW (not in folders)"
        rows.append((label, dt, "Gates Capital", float(amt), "Reimbursement", "—", "EXCLUDED", memo))
    for dt, amt, memo in BEAR_CREEK:
        rows.append(("9922 JSC", dt, "Bear Creek Inc.", float(amt), "Annual gift", "—", "EXCLUDED", memo))
    for dt, amt, acct, memo in UTMA_GIFTS:
        rows.append((acct, dt, "THOMAS (wire)", float(amt), "UTMA gift", "—", "EXCLUDED", memo))
    rows.append(("0203 joint", "2025-12-24", "Freeman's LLC / CIBC", 11000.00, "Unknown — Sale 6428 Contract", "—", "HOLD", "Not payroll. ASK: art sale vs other."))
    rows.append(("8507 NOW", "2025-12-01", "JOAN FRIEDBERGER", -675.00, "Zelle OUT", "—", "HOLD", "Same day as Joan rent in. Deposit/repair? ASK."))
    for i, row in enumerate(rows, 5):
        for c, val in enumerate(row, 1):
            cell = led.cell(i, c, val)
            cell.border = THIN
            cell.fill = status_fill.get(row[6], BLUE)
            if c == 4:
                cell.number_format = "0.00"
    last = 4 + len(rows)
    led.cell(last + 2, 1, "STATUS: APPLIED = in monthly P&L (green). APPLIED-YELLOW = in P&L pending your yes/no. EXCLUDED = not income. HOLD = parked.")
    led.merge_cells(start_row=last + 2, start_column=1, end_row=last + 2, end_column=8)
    widths = [28, 14, 28, 12, 32, 22, 16, 70]
    for i, w in enumerate(widths, 1):
        led.column_dimensions[get_column_letter(i)].width = w
    led.auto_filter.ref = f"A4:H{last}"
    led.freeze_panes = "A5"


def append_ask_inflow_totals(ask: Worksheet) -> None:
    r = ask.max_row + 2
    ask.cell(r, 1, "Checking inflows 2026-09-18").font = Font(bold=True, color="FFFFFF")
    ask.cell(r, 1).fill = HEADER_FILL
    ask.cell(r, 2).fill = HEADER_FILL
    summary = [
        ("216 Oak Park STR / RENTAL INCOME (cash, incl. Ava yellow)", OAK_RENT_TOTAL, YELLOW),
        ("  Delach 8 × $1,450", D("11600"), GREEN),
        ("  Ava Hernandez 8 × $500 (ASK roommate)", D("4000"), YELLOW),
        ("  Joan Friedberger 2 × $1,950 (ASK Dec vs Jan prepaid)", D("3900"), YELLOW),
        ("EPGC Consultant = GCM BOM = 1099-NEC", GCM_FEE_TOTAL, GREEN),
        ("GCM reimbursements EXCLUDED (9922 + 8507)", REIMB_TOTAL, GRAY),
        ("Bear Creek gift EXCLUDED (4 × $9,500)", BEAR_CREEK_TOTAL, GRAY),
        ("UTMA Thomas wires EXCLUDED (2 × $5,000)", UTMA_GIFT_TOTAL, GRAY),
        ("Freeman’s CIBC $11,000 12/24 HOLD", D("11000"), YELLOW),
    ]
    for i, (lab, val, fill) in enumerate(summary):
        ask.cell(r + 1 + i, 1, lab).fill = fill
        cell = ask.cell(r + 1 + i, 2, money(val))
        cell.number_format = '"$"#,##0.00'
        cell.fill = fill


def main() -> None:
    assert XLSX.exists(), XLSX
    wb = load_workbook(XLSX)

    # 216 Oak Park rent
    oak = wb["216 N. Oak Park Ave"]
    oak_map = ans.label_map(oak, 2)
    base.write_month_row(oak, oak_map["STR / RENTAL INCOME"], oak_rent, PROP_COLS, YELLOW)
    oak["B30"] = (
        "LTR rent 2026-09-18: joint 0203 Zelle — Delach $1,450 + Ava $500 (yellow, ASK roommate) Feb–Sep cash, "
        "then Joan $1,950 on 12/1 and 12/31 (yellow, ASK Dec vs Jan-2026 prepaid). "
        f"TY2025 cash ${OAK_RENT_TOTAL}. Oct–Nov no rent on 0203 (ASK vacancy). "
        "Jan 2025 occupancy was paid 12/31/2024 — not in this year. "
        "HOA + Rocket payment cash from Monarch (green). Deduct interest from Rocket 1098 only — not full P+I. "
        "Not tax advice."
    )
    oak["B30"].alignment = Alignment(wrap_text=True)
    oak["B30"].fill = YELLOW
    oak.row_dimensions[30].height = 90

    # EPGC Consultant
    epgc = wb["EPGC LLC"]
    epgc_map = {}
    for r in range(1, 40):
        v = epgc.cell(r, 1).value
        if v:
            epgc_map[str(v).strip()] = r
    base.write_month_row(epgc, epgc_map["Consultant"], gcm_fee, EPGC_COLS, GREEN)
    note_r = 43
    epgc.cell(
        note_r,
        1,
        (
            f"Checking inflows 2026-09-18: GCM BOM PAYMEN ${GCM_FEE_TOTAL} → Consultant (ties 1099-NEC). "
            f"JAKE C / SEFOF reimbursements ${REIMB_TOTAL} EXCLUDED. "
            "ASK: EPGC LLC vs personal 1099. Bear Creek / UTMA not on this tab. Not tax advice."
        ),
    )
    epgc.cell(note_r, 1).alignment = Alignment(wrap_text=True)
    epgc.merge_cells(start_row=note_r, start_column=1, end_row=note_r + 1, end_column=14)

    # Income summary 216 actual
    inc = wb["Income"]
    inc["E5"] = money(OAK_RENT_TOTAL)
    inc["E5"].number_format = "0.00"
    inc["E5"].fill = YELLOW
    inc["A12"] = (
        f"Notes: E4 Monarch paycheck cash ≠ W-2 Box 1. E5 216 Oak Park LTR cash ${OAK_RENT_TOTAL} "
        "(Delach+Ava+Joan; Ava/Joan yellow ASK; Oct–Nov vacant ASK). "
        "E6 PLATFORM STR net LOCKED. E7 GCM 1099-NEC $21,500 = BOM on Consultant (reimbursements excluded). "
        "E8 Art net yellow until Sale Prices filled (COGS on Art Sales). Not tax advice."
    )
    inc["A12"].alignment = Alignment(wrap_text=True)
    inc.merge_cells("A12:E13")
    inc.row_dimensions[12].height = 48

    rebuild_gcm_tab(wb)

    src = wb["_SOURCE_2025"]
    src["A22"] = "Checking inflows 2026-09-18"
    src["B22"] = (
        "JSC 9922 + joint 0203 full-year eStmts. "
        f"GCM BOM ${GCM_FEE_TOTAL} → EPGC Consultant. Reimburse ${REIMB_TOTAL} excluded. "
        f"216 Oak Park rent cash ${OAK_RENT_TOTAL} (Delach / Ava yellow / Joan yellow). "
        f"Bear Creek ${BEAR_CREEK_TOTAL} + UTMA Thomas ${UTMA_GIFT_TOTAL} not income. "
        "Freeman’s CIBC $11,000 HOLD. Not tax advice."
    )
    src["A22"].fill = GREEN
    src["B22"].fill = YELLOW
    src["B22"].alignment = Alignment(wrap_text=True)

    # ASK: keep prior CC Qs, add checking
    ans.CONFIRMED.extend(CHECKING_CONFIRMED)
    # refresh ComEd line with statement evidence
    new_open = []
    for item, detail in ans.STILL_OPEN:
        if item.startswith("REMIND JACOB"):
            detail = (
                "Joint 0203 statements show ComEd ACH (e.g. 1/21 Jacob $59.04 + Megan $132.94). "
                + detail
            )
        new_open.append((item, detail))
    ans.STILL_OPEN[:] = new_open + CHECKING_STILL_OPEN
    ans.rebuild_ask(wb)
    append_ask_inflow_totals(wb["ASK"])
    rebuild_inflow_ledger(wb)

    wb.save(XLSX)
    shutil.copy2(XLSX, DELIVERABLE)

    export_names = [
        ("EPGC LLC", "EPGC_LLC.csv"),
        ("216 N. Oak Park Ave", "216_Oak_Park.csv"),
        ("GCM", "GCM.csv"),
        ("Income", "Income.csv"),
        ("INFLOW_LEDGER", "INFLOW_LEDGER.csv"),
        ("ASK", "ASK.csv"),
        ("_SOURCE_2025", "SOURCE_2025.csv"),
        ("524 Ferdinand Ave, Unit 2", "524_Unit_2.csv"),
        ("524 Ferdinand Ave, Unit 1", "524_Unit_1.csv"),
    ]
    for name, fname in export_names:
        ans.sheet_to_csv(wb[name], ans.CSV_DIR / fname)
        safe = name.replace(",", "").replace(".", "")[:40]
        ans.save_sheet_part(wb[name], ans.PARTS / f"{safe}.xlsx")

    payload = {
        "updated": "2026-09-18",
        "disclaimer": "Not tax advice. Organizational packet only.",
        "accounts": {
            "jsc": {"folder": "bank of america checking_JSC", "id": "1zsgbJxKVdpA4a0vHEIix8wUxnzhNYZdZ", "acct": "4830 6191 9922"},
            "joint": {"folder": "Bank of America checking_joint", "id": "1dIxS-VnqWTSNgvNXt47W_GYx88z2FdCC", "acct": "4830 7359 0203"},
        },
        "gcm_board_1099": float(GCM_FEE_TOTAL),
        "gcm_reimburse_excluded": float(REIMB_TOTAL),
        "oak_park_ltr_cash": float(OAK_RENT_TOTAL),
        "oak_park_by_month": [float(x) for x in oak_rent],
        "bear_creek_gift_excluded": float(BEAR_CREEK_TOTAL),
        "utma_thomas_excluded": float(UTMA_GIFT_TOTAL),
        "hold_freemans_cibc": 11000.00,
        "questions": [a[0] for a in CHECKING_STILL_OPEN],
    }
    JSON_PATH.write_text(json.dumps(payload, indent=2), encoding="utf-8")

    pkt = json.loads(ans.JSON_PATH.read_text(encoding="utf-8")) if ans.JSON_PATH.exists() else {}
    pkt["checking_inflows"] = payload
    ans.JSON_PATH.write_text(json.dumps(pkt, indent=2), encoding="utf-8")

    print("Oak Park rent", OAK_RENT_TOTAL)
    print("GCM consultant", GCM_FEE_TOTAL)
    print("reimburse excluded", REIMB_TOTAL)
    print("Bear Creek excluded", BEAR_CREEK_TOTAL)
    print("UTMA excluded", UTMA_GIFT_TOTAL)
    print("saved", XLSX, "bytes", XLSX.stat().st_size)


if __name__ == "__main__":
    main()
